#!/usr/bin/env python3
"""Round-trip test for tools/blind_review.py with a HAND-MADE fake judge (2026-10-15). No model is called.

Takes 3 real questions from a real quiz page, exports them, writes fake verdicts (one with a wrong quote,
one with a too-long justification), and checks: import accepts exactly the good ones and rejects the bad,
the audit export is REFUSED before any blind record exists, the audit export carries the key, and a wrong
audit quote is rejected. Also (2026-10-15 v2): --ids filter (list file and JSON list, unknown ids reported),
no manifest or key in a judge-visible --out dir (manifest in ~/.cache/pa-quizzes/review-manifests/), the legacy
DIR/manifest.json fallback, picture extraction on a synthetic deck (tiny image skipped, TIFF -> PNG, cached),
and the "[image] <deck>|<slide>" quote accepted only for a slide with listed images. Uses a temp records file and a temp export dir OUTSIDE the repo; the real
tools/blind_review.json is never touched. Needs the Semester 2 inbox (decks) on this machine; skips otherwise.

    python3 tools/test_blind_review.py
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import blind_review as R

FOLDER, FILES = "Clinical Medicine and Surgery I Exam 4", "heart-failure-quiz.html"
MICRO, MICRO_FILE = "Microbiology Exam 2", "cocci-of-medical-importance-quiz.html"


def image_unit_tests(tmp):
    """Synthetic deck: slide 1 text only, slide 2 one big PNG + one tiny PNG, slide 3 a big TIFF."""
    import io
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    def noise(w, h, fmt):
        im = Image.frombytes("RGB", (w, h), os.urandom(w * h * 3))
        b = io.BytesIO()
        im.save(b, fmt)
        b.seek(0)
        return b
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Text only slide with no pictures at all"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.add_picture(noise(120, 120, "PNG"), Inches(1), Inches(1))
    s2.shapes.add_picture(noise(8, 8, "PNG"), Inches(5), Inches(1))               # tiny: decorative, skipped
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.add_picture(noise(100, 100, "TIFF"), Inches(1), Inches(1))
    # slide 4: the whole content is a picture inserted into a picture PLACEHOLDER (no picture shape, no text)
    from pptx.enum.shapes import PP_PLACEHOLDER
    s4 = prs.slides.add_slide(prs.slide_layouts[8])
    ph = [p for p in s4.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE][0]
    ph.insert_picture(noise(160, 120, "PNG"))
    deck = os.path.join(tmp, "synthetic deck.pptx")
    prs.save(deck)
    old = R.IMG_CACHE
    R.IMG_CACHE = os.path.join(tmp, "slides-img")
    try:
        got = R.deck_images(deck)
        assert sorted(got) == [2, 3, 4] and len(got[2]) == 1 and len(got[3]) == 1 and len(got[4]) == 1, got
        txt = R.extract_deck(deck)
        assert R.has_pictures(txt[4]), "a filled picture placeholder must be flagged in the slide text: %r" % txt[4]
        assert R.has_pictures(txt[2]) and not R.has_pictures(txt[1])
        assert got[2][0].endswith(".png") and got[3][0].endswith(".png"), got              # TIFF converted by sips
        assert all(os.path.getsize(x) > 0 and x.startswith(R.IMG_CACHE) for v in got.values() for x in v)
        assert not any(x.startswith(R.ROOT) for v in got.values() for x in v)               # never in the repo
        real = R.extract_images
        R.extract_images = lambda *a, **k: (_ for _ in ()).throw(AssertionError("extracted twice"))
        try:
            assert R.deck_images(deck) == got                                              # second call: cached
        finally:
            R.extract_images = real
        assert R.has_pictures("x\n[this slide has 2 picture(s) and 0 chart(s); their content is NOT in this text]")
        assert not R.has_pictures("=== SLIDE 1 ===\nplain text")
    finally:
        R.IMG_CACHE = old
    # the "[image]" quote: accepted only when that slide has listed images
    ev = {"D.pptx|2": "=== SLIDE 2 ===\nthe words on the slide itself are these ones here"}
    m = {"evidence": ev, "letters": {"A": "a", "B": "b", "C": "c", "D": "d"}, "item_text": ["stem?"], "key": "A", "ah": "x",
         "key_text": "a", "evidence_images": {"D.pptx|2": ["/x/s002-1.png"]}}
    v = {"pick": "A", "cover_answer": "a", "why": "Plate shows it.", "quote": "[image] D.pptx|2"}
    assert R.check_blind(v, m)[0] is not None
    assert R.check_blind(dict(v, quote="[Image] d.pptx|2"), m)[0] is not None               # case-insensitive
    r, why = R.check_blind(dict(v, quote="[image] D.pptx|3"), m)
    assert r is None and "no listed evidence_images" in why, why
    r, why = R.check_blind(v, dict(m, evidence_images={}))
    assert r is None and "no listed evidence_images" in why, why
    a = {"verdict": "revise", "defects": [{"field": "explanation A", "reason": "r", "quote": "[image] D.pptx|2"}], "fix": "f"}
    assert R.check_audit(a, m)[0]["defects"][0]["src"] == "image"
    assert R.check_audit(a, dict(m, evidence_images={}))[0] is None


def mis_join_unit_test():
    """report/check must join `mis` from the sources like export does, or the audit hash export recorded
    never matches and every question with `mis` shows as STALE (the 2026-10-15 Microbiology Exam 2 case)."""
    page_q = {"q": "Stem?", "opts": [["k", "ek"], ["x", "ex"], ["y", "ey"], ["z", "ez"]], "c": 0, "cite": ""}
    src = dict(page_q, mis={"x": "mx", "y": "my", "z": "mz"})
    saved_src, saved_load = R._SRC, R.load_folder
    try:
        R._SRC = ({R.norm(src["q"]): src}, {})
        R.load_folder = lambda f, g=None, u=None: [R.Q(f, "f.html", 0, page_q)]
        exported = R.attach_sources([R.Q("F", "f.html", 0, page_q)])[0].ah
        reported = R.gather(["F"])[0][0].ah
        assert exported == reported, "report computes a different audit hash than export"
        assert exported != R.Q("F", "f.html", 0, page_q).ah, "the source `mis` was not joined"
    finally:
        R._SRC, R.load_folder = saved_src, saved_load


def run(*args):
    p = subprocess.run([sys.executable, os.path.join(HERE, "blind_review.py"), "--records", REC] + list(args),
                       capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def first_quote(text):
    for line in text.split("\n"):
        if len(line.split()) >= 6 and not line.startswith("==="):
            return line
    raise SystemExit("no quotable line in evidence")


def main():
    if not os.path.isdir(R.INBOX):
        print("test_blind_review: SKIPPED (no Semester 2 inbox on this machine)")
        return 0
    global REC
    tmp = tempfile.mkdtemp(prefix="blind_review_test_", dir=os.path.expanduser("~/.cache"))
    REC = os.path.join(tmp, "records.json")
    d1, d2 = os.path.join(tmp, "blind"), os.path.join(tmp, "audit")
    made = [d1, d2]                                          # export dirs whose cached manifests are removed at the end
    try:
        image_unit_tests(tmp)
        mis_join_unit_test()
        # 1. the audit pass is refused before any blind record exists
        rc, out = run("export", "--pass", "audit", "--folder", FOLDER, "--files", FILES, "--sample", "3", "--out", d2)
        assert rc == 1 and "REFUSED" in out, (rc, out)
        # 2. an export directory inside the repo is refused
        rc, out = run("export", "--pass", "blind", "--folder", FOLDER, "--out", os.path.join(R.ROOT, "tools", "_nope"))
        assert rc != 0 and "refusing --out inside the repo" in out, out
        assert not os.path.exists(os.path.join(R.ROOT, "tools", "_nope"))
        # 3. blind export: 3 items, no key / explanations / mis anywhere in the batch
        rc, out = run("export", "--pass", "blind", "--folder", FOLDER, "--files", FILES, "--sample", "3", "--seed", "4", "--out", d1)
        assert rc == 0 and "exported 3 item(s)" in out, out
        batch = json.load(open(os.path.join(d1, "judge-001.json")))
        blob = json.dumps(batch)
        for forbidden in ('"key"', '"explanations"', '"mis"', "Correct."):
            assert forbidden not in blob, "blind batch leaks " + forbidden
        items = batch["items"]
        assert len(items) == 3 and all(sorted(i["options"]) == list("ABCD") for i in items)
        assert not os.path.exists(os.path.join(d1, "manifest.json")), "manifest (keys) sits in the judge-visible dir"
        assert os.path.exists(R.manifest_path(d1)) and R.manifest_path(d1).startswith(R.MANIFESTS)
        assert sorted(os.listdir(d1)) == ["judge-001.json", "unresolved.json"], os.listdir(d1)
        for u in json.load(open(os.path.join(d1, "unresolved.json"))):
            assert set(u) <= {"id", "file", "stem", "why"}, u
        man = R.load_manifest(d1)
        assert all("key" in v for v in man["items"].values())            # ... the keys are in the cached manifest
        ev = {k: man["items"][k]["evidence"] for k in man["items"]}
        good = lambda i: first_quote(list(ev[items[i]["id"]].values())[1])      # a real line from the cited slide
        # 3b. --ids: a one-id-per-line file (with a comment, a blank and an unknown id) and a JSON list
        idf = os.path.join(tmp, "ids.txt")
        open(idf, "w").write("# picked\n%s\n\ndeadbeefdeadbeef\n" % items[1]["id"])
        d3 = os.path.join(tmp, "ids1")
        os.makedirs(d3)
        open(os.path.join(d3, "manifest.json"), "w").write("{}")           # a legacy manifest left in the dir is removed
        made.append(d3)
        rc, out = run("export", "--pass", "blind", "--folder", FOLDER, "--files", FILES, "--ids", idf, "--out", d3)
        assert rc == 0 and "exported 1 item(s)" in out and "UNKNOWN id(s)" in out and "deadbeefdeadbeef" in out, out
        assert [i["id"] for i in json.load(open(os.path.join(d3, "judge-001.json")))["items"]] == [items[1]["id"]]
        assert not os.path.exists(os.path.join(d3, "manifest.json")) and os.path.exists(R.manifest_path(d3))
        idj = os.path.join(tmp, "ids.json")
        json.dump([items[0]["id"], items[2]["id"]], open(idj, "w"))
        d3b = os.path.join(tmp, "ids2")
        made.append(d3b)
        rc, out = run("export", "--pass", "blind", "--folder", FOLDER, "--files", FILES, "--ids", idj, "--out", d3b)
        assert rc == 0 and "exported 2 item(s)" in out and "UNKNOWN" not in out, out
        assert {i["id"] for i in json.load(open(os.path.join(d3b, "judge-001.json")))["items"]} == {items[0]["id"], items[2]["id"]}
        # 3c. legacy fallback: a directory exported before the move keeps DIR/manifest.json and no cached manifest
        leg = os.path.join(tmp, "legacy")
        shutil.copytree(d1, leg)
        json.dump(man, open(os.path.join(leg, "manifest.json"), "w"))
        assert not os.path.exists(R.manifest_path(leg)) and R.load_manifest(leg) == man
        rc, out = run("import", "--pass", "blind", leg, "--judge", "fake-judge")
        assert rc == 0 and "imported (blind)" in out and "'accepted': 0" in out, out
        # 4. fake judge: item 0 good, item 1 WRONG QUOTE, item 2 good but 'why' too long
        verdicts = [
            {"id": items[0]["id"], "pick": "B", "cover_answer": "monitor electrolytes and renal function", "why": "The slide says so.",
             "quote": good(0), "alt": None, "missing": "a lab value", "conflict": "slide vs fact", "clue": None},
            {"id": items[1]["id"], "pick": "A", "cover_answer": "x", "why": "Because.",
             "quote": "this sentence is certainly not on any slide of this deck", "alt": None, "missing": None, "conflict": None, "clue": None},
            {"id": items[2]["id"], "pick": "C", "cover_answer": "x", "why": " ".join(["word"] * 41),
             "quote": good(2), "alt": None, "missing": None, "conflict": None, "clue": None},
        ]
        json.dump(verdicts, open(os.path.join(d1, "verdict-001.json"), "w"))
        rc, out = run("import", "--pass", "blind", d1, "--judge", "fake-judge")
        assert rc == 1 and "'accepted': 1" in out and "'rejected': 2" in out, out
        repair = json.load(open(os.path.join(d1, "repair.json")))
        assert {r["id"] for r in repair} == {items[1]["id"], items[2]["id"]}
        assert "not verbatim" in [r["why"] for r in repair if r["id"] == items[1]["id"]][0]
        rec = json.load(open(REC))
        assert list(rec) == [items[0]["id"]] and rec[items[0]["id"]]["passes"]["blind"]["judge"] == "fake-judge"
        assert rec[items[0]["id"]]["passes"]["blind"]["pick_text"] == items[0]["options"]["B"]
        # 5. audit export now carries ONLY the reviewed item, with key + explanations + the judge's own result
        rc, out = run("export", "--pass", "audit", "--folder", FOLDER, "--files", FILES, "--sample", "3", "--seed", "4",
                      "--only-reviewed", "--from", d1, "--out", d2)
        assert rc == 0, out
        assert "29 skipped for having no blind record" in out and "(0 already have a current audit record" in out, out
        assert not os.path.exists(os.path.join(d2, "manifest.json")) and os.path.exists(R.manifest_path(d2))
        assert "explanations" not in open(os.path.join(d2, "unresolved.json")).read()
        ab = json.load(open(os.path.join(d2, "judge-001.json")))
        assert [i["id"] for i in ab["items"]] == [items[0]["id"]]
        it = ab["items"][0]
        assert set(it["explanations"]) == set("ABCD") and it["key"] in "ABCD" and it["your_blind_result"]["pick"] == "B"
        # 6. audit verdicts: a pass with a defect is rejected; a revise with a wrong quote is rejected; a good revise is kept
        aman = R.load_manifest(d2)["items"][items[0]["id"]]
        q = list(aman["evidence"].values())[1]
        bad1 = {"id": it["id"], "verdict": "pass", "defects": [{"field": "A", "reason": "r", "quote": first_quote(q)}], "fix": None}
        bad2 = {"id": it["id"], "verdict": "revise", "defects": [{"field": "A", "reason": "r", "quote": "an invented sentence about nothing at all"}], "fix": "f"}
        for bad in (bad1, bad2):
            json.dump([bad], open(os.path.join(d2, "verdict-001.json"), "w"))
            rc, out = run("import", "--pass", "audit", d2, "--judge", "fake-judge")
            assert rc == 1 and "'accepted': 0" in out, out
        ok = {"id": it["id"], "verdict": "revise", "defects": [{"field": "explanation B", "reason": "claim not on slide", "quote": first_quote(q)}],
              "fix": "Delete the unsupported claim."}
        json.dump([ok], open(os.path.join(d2, "verdict-001.json"), "w"))
        rc, out = run("import", "--pass", "audit", d2, "--judge", "fake-judge")
        assert rc == 0 and "'accepted': 1" in out, out
        # 7. report and check: scope is empty (nothing fails); --all fails and names the item
        rc, out = run("report", "--folder", FOLDER)
        assert rc == 0 and "AUDIT REVISE: 1" in out, out
        assert "blind CONFLICT flagged (slide vs slide, or slide vs established fact): 1" in out, out
        assert "blind MISSING detail flagged (stem or slides lack something decisive): 1" in out, out
        rc, out = run("check", "--folder", FOLDER)
        assert rc == 0 and "0 question(s) in scope" in out, out
        rc, out = run("check", "--folder", FOLDER, "--all")
        assert rc == 1 and "FAIL" in out, out
        # 7b. with the page in scope, check fails on the unreviewed questions and counts the denominator
        sc = os.path.join(tmp, "scope.json")
        json.dump({"globs": [FOLDER + "/" + FILES]}, open(sc, "w"))
        rc, out = run("--scope", sc, "check", "--folder", FOLDER)
        assert rc == 1 and "0/30 in-scope" in out, out
        # 8. an edit invalidates records: same stem, one option changed -> a different hash
        a = R.item_hash("S?", ["a", "b", "c", "d"])
        assert a == R.item_hash("S?", ["d", "c", "b", "a"]) and a != R.item_hash("S?", ["a", "b", "c", "e"])
        # 9. real deck with pictures: evidence_images listed in the batch, "[image]" quote accepted only for a slide that has them
        d5 = os.path.join(tmp, "micro")
        made.append(d5)
        rc, out = run("export", "--pass", "blind", "--folder", MICRO, "--files", MICRO_FILE, "--size", "40", "--out", d5)
        assert rc == 0, out
        mb = json.load(open(os.path.join(d5, "judge-001.json")))
        withimg = [i for i in mb["items"] if i.get("evidence_images")]
        without = [i for i in mb["items"] if not i.get("evidence_images")]
        if withimg:
            it = withimg[0]
            for k, paths in it["evidence_images"].items():
                assert k.split("|")[0] in it["evidence_slides"] and int(k.split("|")[1]) in it["evidence_slides"][k.split("|")[0]]
                assert paths and all(os.path.isabs(x) and os.path.exists(x) and not x.startswith(R.ROOT) for x in paths), paths
            key = sorted(it["evidence_images"])[0]
            v = [{"id": it["id"], "pick": "A", "cover_answer": "x", "why": "The picture shows it.", "quote": "[image] " + key,
                  "alt": None, "missing": None, "conflict": None, "clue": None}]
            if without:
                v.append({"id": without[0]["id"], "pick": "A", "cover_answer": "x", "why": "Picture.",
                          "quote": "[image] " + key, "alt": None, "missing": None, "conflict": None, "clue": None})
            json.dump(v, open(os.path.join(d5, "verdict-001.json"), "w"))
            rc, out = run("import", "--pass", "blind", d5, "--judge", "fake-judge")
            if without:
                assert "'accepted': 1" in out and "'rejected': 1" in out and "no listed evidence_images" in out, out
            else:
                assert "'accepted': 1" in out, out
        else:
            print("   (note: no picture slides among the %s items on this machine; real-deck image check skipped)" % MICRO_FILE)
        print("test_blind_review: all checks passed (exit 0)")
        return 0
    finally:
        for d in made + [os.path.join(tmp, "legacy")]:
            try:
                os.remove(R.manifest_path(d))
            except OSError:
                pass
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
