#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pull the PDM I Exam 3 Lecture 15 and 16 figures out of the decks (quiz + guide).

Lecture 15 = "15. EKG Ischemia and Infarction.pptx" (25 slides), Lecture 16 =
"16. EKG Pericarditis Electrolytes and Other Abnormalities.pptx" (51 slides), both
Scott Mathis (title slide). Every picture below was viewed at full size before it
was listed, and each crop was viewed after it was made.

QUIZ copies must not give the answer away. Most labels on these slides are
separate PowerPoint text boxes laid over the picture, so the extracted picture is
already clean. Where the label is BAKED INTO the picture it is cropped away here:
  - machine interpretation headers on the prehospital 12-leads (Lecture 15
    slides 16, 20, 21; Lecture 16 slide 19, 33) -- "Inferior ST elevation,
    CONSIDER ACUTE", "Left bundle branch block" and the like;
  - "pathologic Q wave - ECGpedia" under the Lecture 15 slide 12 complex;
  - "Diagnosis: hypokalemia K+ = 2.0" and the "U waves" label on Lecture 16
    slide 38 (only the V4-V6 column with its unlabeled arrows is kept);
  - the "Hypothermia" title and the feature box on Lecture 16 slide 17.
GUIDE copies keep their labels (the guide names the finding anyway).

QUIZ copies get NEUTRAL file names (case-a, rhythm-strip-a ...): the src is visible to anyone who
opens the image in a new tab, and "brugada-type-1.jpg" would answer its own question. What each
one shows is written beside it below. Guide copies are named for what they show.

This script only ever writes (never deletes) inside the shared Exam 3 image folders: other
lectures' builders write there too.

Pictures are addressed by (slide, n-th picture on the slide) and the pixel size is
asserted, so a re-exported deck that renumbers its pictures fails loudly instead of
silently pairing a caption with the wrong strip (image_only_slides memory).
"""
import io, os, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
INBOX = os.path.expanduser("~/Desktop/PA Quizzes/Semester 2/Principles of Diagnostic Medicine I Inbox/Exam 3")
DECK = {15: "15. EKG Ischemia and Infarction.pptx",
        16: "16. EKG Pericarditis Electrolytes and Other Abnormalities.pptx"}
OUT = os.path.join(ROOT, "Principles of Diagnostic Medicine I Exam 3")
QDIR = os.path.join(OUT, "pdm-exam-3-quiz-images")
GDIR = os.path.join(OUT, "pdm-exam-3-study-guide-images")
MAXW = 1400

# (lecture, slide, picture index on slide, expected size, crop box or None, out name)
QUIZ = [
 (15, 6, 2, (774, 410), None, "15-fast-rate-strip.jpg"),
 (15, 7, 1, (800, 163), None, "15-t-wave-strip-a.jpg"),  # normal asymmetrical T waves
 (15, 7, 2, (513, 301), None, "15-precordial-leads-a.jpg"),  # hyperacute T waves V1-V6
 (15, 9, 1, (1111, 345), (0, 22, 1111, 345), "15-12-lead-case-a.jpg"),  # global ST depression (subendocardial ischemia)
 (15, 11, 1, (900, 356), None, "15-12-lead-case-b.jpg"),  # V1-V4 ST elevation with hyperacute T waves
 (15, 12, 2, (300, 225), (0, 0, 300, 196), "15-single-complex-a.png"),  # deep, wide Q wave followed by an R wave
 (15, 12, 3, (200, 243), None, "15-single-complex-b.jpg"),  # QS complex, no R wave
 (15, 13, 1, (800, 446), None, "15-limb-leads-case-c.jpg"),  # Q waves and ST elevation in III and aVF
 (15, 16, 1, (1889, 678), (0, 182, 1889, 678), "15-12-lead-case-d.jpg"),  # inferior STEMI, reciprocal I and aVL
 (15, 17, 1, (1200, 446), None, "15-12-lead-case-e.jpg"),  # high lateral STEMI I and aVL, reciprocal III and aVF
 (15, 18, 1, (1200, 516), None, "15-12-lead-case-f.jpg"),  # anteroseptal STEMI V1-V3
 (15, 20, 1, (1600, 643), (0, 168, 1600, 600), "15-12-lead-case-g.jpg"),  # isolated ST depression V1-V3
 (15, 21, 1, (1600, 637), (0, 150, 1600, 600), "15-12-lead-case-h.jpg"),  # V4-V6 moved to V7-V9: elevation V8, V9
 (15, 23, 1, (1200, 493), None, "15-12-lead-case-i.jpg"),  # inferior STEMI, III > II, reciprocal I and aVL
 (15, 24, 1, (768, 401), None, "15-12-lead-case-j.jpg"),  # V4 moved to V4R (handwritten): elevation

 (16, 6, 1, (1024, 395), None, "16-12-lead-case-m.jpg"),  # acute pericarditis
 (16, 7, 2, (308, 164), None, "16-lead-v4-strip-a.jpg"),  # J wave notch, BER slide
 (16, 8, 1, (1024, 522), None, "16-12-lead-case-e.jpg"),  # benign early repolarization
 (16, 9, 1, (614, 320), None, "16-lead-v6-strip-a.png"),  # V6 ST/T ratio < 25% (BER)
 (16, 9, 2, (614, 320), None, "16-lead-v6-strip-b.png"),  # V6 ST/T ratio > 25% (pericarditis)
 (16, 11, 1, (1400, 722), None, "16-12-lead-case-a.jpg"),  # Brugada type 1 (coved V1-V2)
 (16, 11, 2, (300, 479), None, "16-v1-v3-sketch-a.jpg"),  # coved ST elevation V1-V3 sketch
 (16, 12, 1, (1600, 640), None, "16-12-lead-case-b.jpg"),  # Brugada type 2 (saddleback V2)
 (16, 14, 1, (1336, 761), None, "16-12-lead-case-q.jpg"),  # ventricular paced rhythm
 (16, 15, 1, (1024, 428), None, "16-12-lead-case-f.jpg"),  # LBBB with global ST and T changes
 (16, 16, 1, (250, 208), None, "16-single-complex-a.jpg"),  # Osborn wave (hypothermia slide)
 (16, 17, 1, (1920, 1080), (36, 173, 1884, 719), "16-12-lead-case-c.jpg"),  # hypothermia, Osborn waves
 (16, 19, 1, (1564, 612), (0, 156, 1564, 575), "16-12-lead-case-g.jpg"),  # LBBB meeting Sgarbossa
 (16, 28, 1, (1200, 824), None, "16-12-lead-case-o.jpg"),  # WPW type A (positive V1)
 (16, 29, 1, (1200, 719), (0, 60, 1200, 719), "16-12-lead-case-n.jpg"),  # WPW type B (negative V1)
 (16, 32, 1, (1200, 711), None, "16-12-lead-case-j.jpg"),  # orthodromic AVRT, narrow regular tachycardia
 (16, 33, 1, (1152, 471), (0, 46, 1152, 471), "16-rhythm-strip-b.jpg"),  # antidromic AVRT, wide regular tachycardia
 (16, 35, 1, (796, 441), None, "16-12-lead-case-l.jpg"),  # hyperkalemia, peaked T, flat P, bradycardia
 (16, 36, 1, (976, 415), None, "16-12-lead-case-r.jpg"),  # late hyperkalemia, wide QRS, no P waves
 (16, 38, 1, (2701, 1126), (1998, 40, 2690, 756), "16-leads-v4-v6-strip-a.jpg"),  # hypokalemia U waves (arrows only)
 (16, 41, 1, (897, 479), None, "16-12-lead-case-h.jpg"),  # hypocalcemia, long ST, normal T
 (16, 43, 1, (673, 271), None, "16-12-lead-case-p.jpg"),  # hypercalcemia, short ST
 (16, 45, 1, (850, 418), None, "16-12-lead-case-i.jpg"),  # hypothyroidism, low voltage
 (16, 47, 1, (1000, 206), None, "16-rhythm-strip-a.jpg"),  # digoxin toxicity: ventricular bigeminy
 (16, 49, 1, (1200, 637), None, "16-12-lead-case-k.jpg"),  # tricyclic antidepressant overdose
 (16, 50, 1, (797, 442), None, "16-12-lead-case-d.jpg"),  # pulmonary embolism, S1Q3T3, sinus tachycardia
]

GUIDE = [
 (15, 4, 1, (435, 500), None, "15-coronary-arteries.jpg"),
 (15, 5, 1, (680, 680), None, "15-j-point-st-segment.jpg"),
 (15, 6, 1, (363, 139), None, "15-tp-segment-strip.jpg"),
 (15, 6, 2, (774, 410), None, "15-fast-rate-strip.jpg"),
 (15, 7, 1, (800, 163), None, "15-normal-t-waves.jpg"),
 (15, 7, 2, (513, 301), None, "15-hyperacute-t-waves.jpg"),
 (15, 8, 1, (1346, 742), None, "15-subendocardial-ischemia.png"),
 (15, 9, 1, (1111, 345), (0, 22, 1111, 345), "15-global-st-depression.jpg"),
 (15, 10, 1, (1315, 760), None, "15-transmural-ischemia.png"),
 (15, 11, 1, (900, 356), None, "15-transmural-st-elevation.jpg"),
 (15, 12, 1, (532, 509), None, "15-pathologic-q-formation.jpg"),
 (15, 12, 2, (300, 225), None, "15-pathologic-q-complex.png"),
 (15, 12, 3, (200, 243), None, "15-pathologic-q-no-r.jpg"),
 (15, 13, 1, (800, 446), None, "15-inferior-q-waves.jpg"),
 (15, 14, 1, (836, 378), None, "15-reciprocal-changes.jpg"),
 (15, 15, 1, (953, 666), None, "15-lead-localization-table.png"),
 (15, 16, 1, (1889, 678), None, "15-inferior-stemi.jpg"),
 (15, 17, 1, (1200, 446), None, "15-high-lateral-stemi.jpg"),
 (15, 18, 1, (1200, 516), None, "15-anteroseptal-stemi.jpg"),
 (15, 19, 1, (1000, 800), None, "15-posterior-lead-placement.png"),
 (15, 20, 1, (1600, 643), None, "15-isolated-v1-v3-depression.jpg"),
 (15, 21, 1, (1600, 637), None, "15-posterior-stemi-v8-v9.jpg"),
 (15, 22, 1, (1000, 800), None, "15-v4r-placement.jpg"),
 (15, 23, 1, (1200, 493), None, "15-inferior-stemi-rv.jpg"),
 (15, 24, 1, (768, 401), None, "15-v4r-elevation.jpg"),

 (16, 5, 1, (300, 197), None, "16-pericarditis-v5.jpg"),
 (16, 5, 2, (300, 188), None, "16-pericarditis-avr.png"),
 (16, 6, 1, (1024, 395), None, "16-pericarditis-12-lead.jpg"),
 (16, 7, 1, (594, 435), None, "16-j-wave-diagram.png"),
 (16, 7, 2, (308, 164), None, "16-j-wave-v4.jpg"),
 (16, 8, 1, (1024, 522), None, "16-ber-12-lead.jpg"),
 (16, 9, 1, (614, 320), None, "16-v6-ber.png"),
 (16, 9, 2, (614, 320), None, "16-v6-pericarditis.png"),
 (16, 10, 1, (788, 635), None, "16-brugada-types.jpg"),
 (16, 11, 1, (1400, 722), None, "16-brugada-type-1.jpg"),
 (16, 12, 1, (1600, 640), None, "16-brugada-type-2.jpg"),
 (16, 13, 1, (550, 291), None, "16-lbbb-paced-normal.jpg"),
 (16, 14, 1, (1336, 761), None, "16-ventricular-paced.jpg"),
 (16, 15, 1, (1024, 428), None, "16-lbbb-st-t.jpg"),
 (16, 16, 1, (250, 208), None, "16-osborn-wave.jpg"),
 (16, 17, 1, (1920, 1080), None, "16-hypothermia-12-lead.jpg"),
 (16, 18, 1, (700, 492), None, "16-sgarbossa-criteria.jpg"),
 (16, 19, 1, (1564, 612), None, "16-sgarbossa-lbbb-example.jpg"),
 (16, 25, 1, (664, 493), None, "16-preexcitation.png"),
 (16, 26, 1, (492, 335), None, "16-wpw-bundle-of-kent.png"),
 (16, 27, 1, (850, 572), None, "16-delta-wave.png"),
 (16, 28, 1, (1200, 824), None, "16-wpw-type-a.jpg"),
 (16, 29, 1, (1200, 719), None, "16-wpw-type-b.jpg"),
 (16, 31, 1, (1485, 1091), None, "16-avrt-circuits.jpg"),
 (16, 32, 1, (1200, 711), None, "16-orthodromic-avrt.jpg"),
 (16, 33, 1, (1152, 471), None, "16-antidromic-avrt.jpg"),
 (16, 34, 1, (1020, 1120), None, "16-potassium-table.jpg"),
 (16, 35, 1, (796, 441), None, "16-hyperkalemia-peaked-t.jpg"),
 (16, 36, 1, (976, 415), None, "16-hyperkalemia-wide-qrs.jpg"),
 (16, 37, 1, (305, 165), None, "16-hypokalemia-diagram.png"),
 (16, 38, 1, (2701, 1126), None, "16-hypokalemia-u-waves.jpg"),
 (16, 40, 1, (600, 177), None, "16-hypocalcemia-diagram.png"),
 (16, 41, 1, (897, 479), None, "16-hypocalcemia-12-lead.jpg"),
 (16, 42, 1, (500, 147), None, "16-hypercalcemia-diagram.png"),
 (16, 43, 1, (673, 271), None, "16-hypercalcemia-12-lead.jpg"),
 (16, 45, 1, (850, 418), None, "16-hypothyroidism-12-lead.jpg"),
 (16, 47, 1, (1000, 206), None, "16-digoxin-bigeminy.jpg"),
 (16, 49, 1, (1200, 637), None, "16-tricyclic-overdose.jpg"),
 (16, 50, 1, (797, 442), None, "16-pulmonary-embolism.jpg"),
]


def pictures(lec):
    prs = Presentation(os.path.join(INBOX, DECK[lec]))
    out = {}

    def walk(shapes, sn):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, sn)
            elif hasattr(sh, "image"):
                try:
                    out.setdefault(sn, []).append(sh.image.blob)
                except Exception:
                    pass
    for i, s in enumerate(prs.slides, 1):
        walk(s.shapes, i)
    return out


# Lecture 15 slide 21: the handwritten relabel of the moved V4 electrode is illegible (it reads
# like "V4R"), and the slide itself lays a "V7" text box over it. Reproduce that overlay on the
# quiz copy so the picture says what the slide says. Box in post-crop pixels.
OVERLAY = {"15-12-lead-case-h.jpg": ((1168, 0, 1372, 44), "V7")}


def save(blob, size, crop, path):
    im = Image.open(io.BytesIO(blob))
    assert im.size == size, "%s: picture is %r, expected %r (deck renumbered?)" % (path, im.size, size)
    if crop:
        im = im.crop(crop)
    ov = OVERLAY.get(os.path.basename(path)) if os.path.dirname(path) == QDIR else None
    if ov:
        from PIL import ImageDraw, ImageFont
        box, text = ov
        im = im.convert("RGB"); d = ImageDraw.Draw(im)
        d.rectangle(box, fill=(253, 238, 240))
        try:
            f = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 30)
        except Exception:
            f = ImageFont.load_default()
        d.text((box[0] + 70, box[1] + 6), text, fill=(20, 60, 20), font=f)
    if im.width > MAXW:
        im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
    if path.endswith(".jpg"):
        im.convert("RGB").save(path, quality=86, optimize=True)
    else:
        im.save(path, optimize=True)
    return im.size


def main():
    os.makedirs(QDIR, exist_ok=True); os.makedirs(GDIR, exist_ok=True)
    pics = {15: pictures(15), 16: pictures(16)}
    sizes = {}
    for table, d in ((QUIZ, QDIR), (GUIDE, GDIR)):
        names = [t[-1] for t in table]
        assert len(names) == len(set(names)), "duplicate output name"
        for lec, sn, k, size, crop, name in table:
            blob = pics[lec][sn][k - 1]
            sizes[(d, name)] = save(blob, size, crop, os.path.join(d, name))
    print("quiz images %d -> %s" % (len(QUIZ), os.path.relpath(QDIR, ROOT)))
    print("guide images %d -> %s" % (len(GUIDE), os.path.relpath(GDIR, ROOT)))
    return sizes


if __name__ == "__main__":
    main()
