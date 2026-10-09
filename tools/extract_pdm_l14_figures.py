#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pull the Lecture 14 (Axis, Bundle Branch Blocks, and Chamber Enlargement) figures for
PDM I Exam 3: the 12-lead and strip images for the quizzes and the teaching figures for the
guide.

Deck: "14. EKG Axis, BBB.pptx" (Scott Mathis, title slide). Every picture in the deck was
extracted and viewed at full size before it was chosen here, including the three stored as
.bin media that python-pptx cannot type (slides 23, 51 and 52). Slide images are cleared for
use provided the slide is cited.

QUIZ images carry NO label that names the finding. Baked-in labels were masked or cropped:
  slide 12  V1/V6 left bundle branch block drawing: "rS" and "R" masked
  slide 16  V1/V6 right bundle branch block drawing: "rSR" and "qRs" masked
  slide 30  "Left Atrial Abnormality" title cropped off the A/B panel
  slide 31  "P mitrale" and "> 40ms" masked on the lead II strip
  slide 36  the left ventricular hypertrophy drawing: ECG panel only (the thick-walled heart
            and the "left ventricle hypertrophy" caption cropped off)
  slide 59  hospital header with two personal names cropped off (also in the guide copy)

TWO DECK FIGURES ARE WRONG AND ARE NOT USED AS TEACHING FIGURES:
  slide 22  the quadrant drawing's right-axis panel shows lead I up and aVF down. The lecturer
            said so in class ("Ignore the bottom left, that's incorrect ... it should be in the
            reverse", part 2, 42:28; and 44:31 "aVF should be up, and one should be down").
            Slide 25's own text gives it correctly. The guide teaches the quadrants as a table.
  slide 21  the causes wheel lists "Left ventricular hypertrophy (LVH)" under RIGHT axis
            deviation; slide 36 says LVH causes left axis deviation, and the true cause on that
            list is right ventricular hypertrophy. Used in the guide with that correction
            beside it; no question is built on its right-axis list.

Left out: slide 40 (stock photograph), slide 43/44 (inferior infarction: Lecture 15's
subject), slide 29's V1-V3 panel (the biphasic P is not clear enough to teach from).
"""
import io, os, re, zipfile
from PIL import Image
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DECK = os.path.expanduser(
    "~/Desktop/PA Quizzes/Semester 2/Principles of Diagnostic Medicine I Inbox/Exam 3/"
    "14. EKG Axis, BBB.pptx")
FOLDER = os.path.join(ROOT, "Principles of Diagnostic Medicine I Exam 3")
QUIZ = os.path.join(FOLDER, "pdm-exam-3-quiz-images")
GUIDE = os.path.join(FOLDER, "pdm-exam-3-study-guide-images")
MAXW = 1200

# (slide, picture shape id, output stem, destinations, crop box or None, white masks)
FIGS = [
 (6, 4, "l14-conduction-system", "G", None, []),
 (7, 5, "l14-lafb-12-lead", "QG", None, []),
 (9, 5, "l14-lpfb-limb-leads", "QG", None, []),
 (10, 4, "l14-left-bundle-fascicles", "G", None, []),
 (11, 4, "l14-lbbb-conduction-path", "G", None, []),
 (12, 4, "l14-lbbb-v1-v6-labeled", "G", None, []),
 (12, 4, "l14-lbbb-v1-v6", "Q", None, [(245, 595, 330, 660), (845, 465, 900, 525)]),
 (13, 4, "l14-lbbb-12-lead-a", "QG", None, []),
 (14, 4, "l14-lbbb-12-lead-b", "Q", None, []),
 (16, 4, "l14-rbbb-v1-v6-labeled", "G", None, []),
 (16, 4, "l14-rbbb-v1-v6", "Q", None, [(57, 120, 91, 139), (209, 120, 245, 139)]),
 (17, 4, "l14-rbbb-criteria-card", "G", None, []),
 (17, 5, "l14-rbbb-v1-morphologies", "G", None, []),
 (18, 4, "l14-rbbb-12-lead-b", "Q", None, []),
 (19, 4, "l14-rbbb-12-lead-a", "QG", None, []),
 (20, 6, "l14-axis-wheel", "G", None, []),
 (21, 4, "l14-axis-causes", "G", None, []),
 (23, 4, "l14-axis-normal", "QG", None, []),
 (24, 4, "l14-axis-left", "QG", None, []),
 (25, 5, "l14-axis-right", "QG", None, []),
 (26, 5, "l14-axis-extreme", "QG", None, []),
 (27, 4, "l14-chamber-dilation", "G", None, []),
 (28, 4, "l14-rae-criteria", "G", None, []),
 (29, 4, "l14-rae-limb-leads", "QG", None, []),
 (29, 13, "l14-rae-lead-ii", "Q", None, []),
 (30, 4, "l14-lae-panels", "Q", (0, 18, 292, 171), []),
 (30, 7, "l14-p-mitrale-schematic", "G", None, []),
 (31, 3, "l14-lae-lead-ii", "Q", None, [(138, 147, 242, 182), (316, 64, 400, 93)]),
 (31, 5, "l14-lae-v1", "QG", None, []),
 (31, 7, "l14-lae-ii-v1", "Q", None, []),
 (31, 11, "l14-lae-notched-p", "Q", None, []),
 (32, 4, "l14-rvh-heart", "G", None, []),
 (32, 5, "l14-lvh-heart", "G", None, []),
 (33, 4, "l14-strain-pattern", "QG", None, []),
 (34, 4, "l14-rvh-12-lead-a", "QG", None, []),
 (35, 4, "l14-rvh-12-lead-b", "Q", None, []),
 (36, 4, "l14-lvh-schematic", "Q", (0, 275, 300, 578), []),
 (36, 4, "l14-lvh-schematic-labeled", "G", None, []),
 (37, 4, "l14-lvh-12-lead-a", "QG", None, []),
 (38, 4, "l14-lvh-12-lead-b", "Q", None, []),
 (39, 4, "l14-lvh-avl", "QG", None, []),
 (41, 8, "l14-normal-12-lead", "Q", None, []),
 (45, 2, "l14-axis-left-b", "Q", None, []),
 (49, 2, "l14-axis-left-c", "Q", None, []),
 (51, 2, "l14-lbbb-12-lead-c", "Q", None, []),
 (55, 3, "l14-rae-right-axis", "Q", None, []),
 (59, 5, "l14-lvh-stemi-mimic", "QG", (0, 48, 1148, 548), []),
]


def blobs():
    """{(slide, shape id): bytes}, including pictures python-pptx cannot type (.bin media)."""
    out = {}
    prs = Presentation(DECK)
    z = zipfile.ZipFile(DECK)
    for i, s in enumerate(prs.slides, 1):
        rels = {r.rId: r.target_ref for r in s.part.rels.values()}
        for sh in s.shapes:
            el = sh._element
            emb = el.xpath('.//a:blip/@r:embed')
            if not emb:
                continue
            tgt = rels.get(emb[0])
            if not tgt or "media/" not in tgt:
                continue
            path = "ppt/media/" + tgt.split("media/")[-1]
            out[(i, sh.shape_id)] = z.read(path)
    return out


def main():
    os.makedirs(QUIZ, exist_ok=True)
    os.makedirs(GUIDE, exist_ok=True)
    B = blobs()
    for slide, sid, stem, dest, crop, masks in FIGS:
        assert (slide, sid) in B, "slide %d has no picture %d" % (slide, sid)
        im = Image.open(io.BytesIO(B[(slide, sid)]))
        fmt = im.format
        if im.mode in ("P", "RGBA", "LA"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        if crop:
            im = im.crop(crop)
        for box in masks:
            im.paste((255, 255, 255), box)
        if im.width < 700:                         # small deck strips read badly at natural size
            f = min(4, -(-800 // im.width))
            im = im.resize((im.width * f, im.height * f), Image.LANCZOS)
        if im.width > MAXW:
            im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
        ext = ".jpg" if fmt == "JPEG" and not masks else ".png"
        for d in ([QUIZ] if "Q" in dest else []) + ([GUIDE] if "G" in dest else []):
            p = os.path.join(d, stem + ext)
            if ext == ".jpg":
                im.save(p, quality=92)
            else:
                im.save(p, optimize=True)
                if os.path.getsize(p) > 200 * 1024:      # a photographed tracing: JPEG is a tenth the size
                    os.remove(p)
                    p = p[:-4] + ".jpg"
                    im.save(p, quality=90)
            print("%-6s slide %2d  %-44s %4dx%-4d %3d KB" % (os.path.basename(d)[11:16], slide, os.path.basename(p),
                  im.width, im.height, os.path.getsize(p) // 1024))


if __name__ == "__main__":
    main()
