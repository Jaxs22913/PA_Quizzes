#!/usr/bin/env python3
"""Write the Lecture 15 (Fungi) and Lecture 16 (Food Safety) QUIZ pictures into
"Microbiology Exam 3/micro-exam-3-quiz-images/".

Quiz pictures get NEUTRAL, numbered file names (l15-fig-01.jpg ...; memory quiz_image_filenames): a picture opened
on its own must never name its answer. Numbers follow slide order and are fixed here, so a rebuild writes the same
names. The study guide keeps its own descriptive copies (micro-exam-3-study-guide-images/).

QUIZ images carry NO label that names the answer, so each is a crop of the slide's only picture:
  Lecture 15 (PAJ5200.Fungi of Medical importance.pptx)
    slide 27  both colonies, the "(a)"/"(b)" letters and the credit cropped off (the slide's "< 35C" and "37C"
              text boxes are not part of the picture)
    slide 41  the forearm photograph, the copyright line and credit cropped off
    slide 47  panel (a) only, the scalp; panel (b), the letters and the credit cropped off
    slide 50  panel (b) only, the hair-shaft nodule, cropped inside its "Hair shaft" / "Black piedra" leader
              lines and labels
    slide 54  panel (a) only, the tongue; panel (b), the letters and the credit cropped off
    slide 59  the tissue section below the "Cyst, containing several boides" label and its two leader lines;
              the credit cropped off
  Lecture 16 (PAJ5200.FoodSafety-micro.pptx)
    slide 14  the RIGHT picture by geometry (the agar plate; the left one is a drawing of a rod)
    slide 24  the RIGHT picture by geometry (the spread diagram), with its "Listeria monocytogenes" heading
              painted out white (the words "contaminated food" stay)
    slide 34  the gel and its lane labels; the header naming the organism and the "PFGE-XbaI" method label
              cropped off
Each crop box was read off the original pixels and the result viewed.

    python3 tools/extract_micro_l15_l16_quiz_images.py
"""
import io, os
from pptx import Presentation
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
INBOX = os.path.expanduser("~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 3/")
L15 = INBOX + "PAJ5200.Fungi of Medical importance.pptx"
L16 = INBOX + "PAJ5200.FoodSafety-micro.pptx"
OUT = os.path.join(ROOT, "Microbiology Exam 3", "micro-exam-3-quiz-images")

# (deck, slide, which picture by position: "only" | "left" | "right", crop box on the original pixels or None,
#  white-out box on the original pixels or None, name)
JOBS = [
    (L15, 27, "only", (0, 25, 1200, 430), None, "l15-fig-01.jpg"),
    (L15, 41, "only", (20, 20, 790, 540), None, "l15-fig-02.jpg"),
    (L15, 47, "only", (80, 15, 440, 400), None, "l15-fig-03.jpg"),
    (L15, 50, "only", (10, 510, 380, 840), None, "l15-fig-04.jpg"),
    (L15, 54, "only", (40, 25, 525, 330), None, "l15-fig-05.jpg"),
    (L15, 59, "only", (0, 305, 747, 580), None, "l15-fig-06.jpg"),
    (L16, 14, "right", None, None, "l16-fig-01.jpg"),
    (L16, 24, "right", None, (14, 24, 530, 72), "l16-fig-02.jpg"),
    (L16, 34, "only", (0, 68, 500, 383), None, "l16-fig-03.jpg"),
]
MAXW = 600


def pictures(shapes):
    """Pictures, picture placeholders and grouped pictures, in shape-tree order."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from pictures(sh.shapes)
            continue
        try:
            sh.image
        except Exception:
            continue
        yield sh


def pick(slide, which):
    pics = list(pictures(slide.shapes))
    if which == "only":
        assert len(pics) == 1, "expected one picture, found %d" % len(pics)
        return pics[0]
    pics.sort(key=lambda sh: sh.left)
    assert len(pics) == 2, "expected two pictures side by side, found %d" % len(pics)
    return pics[0] if which == "left" else pics[1]


def main():
    decks = {}
    os.makedirs(OUT, exist_ok=True)
    for deck, n, which, box, white, name in JOBS:
        prs = decks.setdefault(deck, Presentation(deck))
        sh = pick(prs.slides[n - 1], which)
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
        if white:
            ImageDraw.Draw(im).rectangle(white, fill=(255, 255, 255))
        if box:
            im = im.crop(box)
        if im.width > MAXW:
            im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
        im.save(os.path.join(OUT, name), quality=85)
        print("wrote %-16s %dx%d  (%s slide %d, %s picture)"
              % (name, im.width, im.height, os.path.basename(deck)[8:20], n, which))


if __name__ == "__main__":
    main()
