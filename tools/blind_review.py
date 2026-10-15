#!/usr/bin/env python3
"""BLIND review + AUDIT of every Semester 2 question against the deck it cites.

Why this exists (2026-10-15): the writer checks its own work, so a wrong key or a second defensible answer
can ride through every deterministic checker. tools/question_writing/WRITER.md ("After writing") adds an
independent reader in two passes, modeled on audit_guide_links.py (batches out, JSON verdicts back, every
quote machine-checked):

  PASS 1 "blind"  the judge sees the stem, the options (A-D) and the deck evidence -- the cited slide plus
                  one either side -- and NOT the key, NOT the explanations, NOT `mis`. It first writes the
                  answer it would give from the stem alone, then picks, or says ambiguous / insufficient.
  PASS 2 "audit"  only for an item that already has a blind record. The judge gets the key, every option
                  explanation (and `mis`) and its own blind result, and returns pass / revise / hold with
                  defects, each carrying a verbatim quote.

Instructions the judges read: tools/question_writing/REVIEW_BLIND.md and REVIEW_AUDIT.md.
Judges run in a Claude Code session (usage, not API credits) -- NEVER through Pip's helper or an API batch
(memory no_bulk_model_runs). This script only writes batches and reads verdicts back.

UNIT. One question, identified by a hash of its exact stem + option texts (order-independent: sorting numeric
options, or a master exam permuting positions, does not change the item). An edit to the stem or any option
text makes a new item with no record. The AUDIT record also carries a second hash over the key and every
explanation and `mis`, so editing an explanation alone makes the audit stale (the blind record stays).
A master-exam question whose text matches a reviewed topic question is the SAME item: it reuses the record.

EVIDENCE. Decks stay OUT of the repo (the repo is public): slide text comes from
~/Desktop/PA Quizzes/Semester 2/<Class> Inbox/, extracted with python-pptx ("=== SLIDE N ===" markers) and
cached per deck content hash in ~/.cache/pa-quizzes/slides/. Export directories hold slide text, so --out
inside the repo is refused. Records (tools/blind_review.json) hold judge verdicts and quotes <= 200 chars.
A slide's pictures are not in the text; the evidence says so. Every evidence slide whose deck slide has
pictures gets them extracted once (python-pptx; EMF/WMF/TIFF converted to PNG with sips; pictures under
~8 KB or 64x64 skipped) into ~/.cache/pa-quizzes/slides-img/<deck content hash>/ -- never into the repo -- and
each batch item lists them: "evidence_images": {"<deck>|<slide>": ["/abs/path.png", ...]}. The judge opens
them with the Read tool; a justification that rests on a picture may quote "[image] <deck>|<slide>", which
import accepts only when that slide has listed images.

Usage:
  blind_review.py export --pass blind --folder "<folder>" [--files GLOB] [--sample N --seed S] [--size 25]
                         [--ids FILE] [--include-reviewed] --out DIR
  blind_review.py import --pass blind DIR --judge MODEL_ID
  blind_review.py export --pass audit --folder "<folder>" [--files GLOB] [--from BLIND_DIR] --out DIR
  blind_review.py import --pass audit DIR --judge MODEL_ID
  blind_review.py report [--folder "<folder>" ...]
  blind_review.py check  --folder "<folder>" [--all]

--ids FILE (both passes): export only the item ids in FILE (one per line, or a JSON list); an unknown id is
reported, not fatal. Combine with --include-reviewed to re-read an id that already has a record.

Files in an export directory (judge-visible, so NO KEYS): judge-NNN.json (the batch; the audit batch carries the
key by design), unresolved.json (id, file, stem, reason: questions with no usable cite/deck, listed so nothing is
skipped silently). The manifest (keys, explanations, slide text: what import needs) is written OUTSIDE it, to
~/.cache/pa-quizzes/review-manifests/<sha1 of the absolute --out path>.json. import and `export --pass audit
--from DIR` look it up there and fall back to DIR/manifest.json for directories exported before this change.
The judge writes verdict-NNN.json next to each judge-NNN.json. `import` writes repair.json (rejected items
and why) and records only the accepted ones.

`check` exits 1 when any question in scope (tools/question_writing/scope.json, globs of quiz files built under
the v1 rules; `--all` = every Semester 2 question in the folder) lacks a CURRENT blind record AND a CURRENT
audit PASS. Scope is EMPTY today, so nothing existing fails.

Question shape (page): {q, opts:[[text, explanation, (mis)], ...], c, cite}. `mis` (the misconception each
distractor stands for, WRITER.md rule 5) is not rendered into pages, so it is read from the pool/sets sources
(tools/*_sets.json, and the pool modules that mention `mis`) by exact stem, falling back to the option set.
Accepted source spellings: a third list element, an option dict with "mis", or a question-level "mis" that is
a list aligned with opts or a {option text | letter: reason} dict.
"""
import argparse
import datetime
import difflib
import fnmatch
import glob
import hashlib
import json
import os
import random
import re
import runpy
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_guide_links as B  # noqa: E402  (deck_of, slides_in: the same cite parsing as the guide-link tools)
import check_exam_standard as X  # noqa: E402  (load, is_frozen: the same question reader and frozen list)

RECORDS = os.path.join(HERE, "blind_review.json")
SCOPE = os.path.join(HERE, "question_writing", "scope.json")
CACHE = os.path.expanduser("~/.cache/pa-quizzes/slides")
IMG_CACHE = os.path.expanduser("~/.cache/pa-quizzes/slides-img")
MANIFESTS = os.path.expanduser("~/.cache/pa-quizzes/review-manifests")
MIN_IMG_BYTES = 8 * 1024                    # smaller pictures are decorative (bullets, logos, rules)
MIN_IMG_PX = 64                             # ... and so is anything under 64x64 pixels
IMAGE_QUOTE = "[image] "
INBOX = os.path.expanduser("~/Desktop/PA Quizzes/Semester 2")
PROMPT_VERSION = "v1"
LETTERS = "ABCD"
BATCH = 25
WHY_WORDS = 40
QUOTE_KEEP = 200
S2_PREFIXES = ("Clinical", "Microbiology", "Pharmacology I", "Physical Diagnosis 2", "Principles", "Interpretation")
BLIND_VERDICTS = LETTERS + "|ambiguous|insufficient"


# ----------------------------------------------------------------------------- text helpers
def norm(t):
    return re.sub(r"\s+", " ", t or "").strip()


def qnorm(t):
    """Quote comparison form: case, curly quotes, dashes, ellipses and whitespace do not matter."""
    t = (t or "").lower()
    for a, b in (("’", "'"), ("‘", "'"), ("“", '"'), ("”", '"'), ("–", "-"),
                 ("—", "-"), ("−", "-"), (" ", " "), ("…", " "), ("...", " ")):
        t = t.replace(a, b)
    return re.sub(r"\s+", " ", t).strip()


def sha(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest()[:16]


def item_hash(stem, opt_texts):
    return sha(json.dumps([norm(stem), sorted(norm(o) for o in opt_texts)], ensure_ascii=False))


def audit_hash(key_text, opts):
    """Key + every (option text, explanation, mis), order-independent. A change makes the audit stale."""
    return sha(json.dumps([norm(key_text), sorted([norm(t), norm(e), norm(m)] for t, e, m in opts)], ensure_ascii=False))


# ----------------------------------------------------------------------------- reading questions
def s2_folders():
    out = []
    for d in sorted(os.listdir(ROOT)):
        if not d.startswith(S2_PREFIXES) or d.startswith(X.TEXTBOOK) or X.is_frozen(d):
            continue
        if os.path.isdir(os.path.join(ROOT, d)):
            out.append(d)
    return out


class Q:
    """One question as shipped on a page, joined to its source record (mis) where one is found."""

    def __init__(self, folder, rel, idx, q):
        self.folder, self.file, self.idx = folder, rel, idx
        self.stem = q["q"]
        self.cite = q.get("cite", "") or ""
        self.key = q["c"]
        self.opts = []                      # [(text, explanation, mis)]
        for i, o in enumerate(q["opts"]):
            mis = o[2] if len(o) > 2 and isinstance(o[2], str) else ""
            self.opts.append((o[0], o[1], mis))
        if isinstance(q.get("mis"), (list, dict)):
            self._apply_mis(q["mis"])
        self.h = item_hash(self.stem, [o[0] for o in self.opts])
        self.sourced = False

    def _apply_mis(self, m):
        new = []
        for i, (t, e, cur) in enumerate(self.opts):
            v = cur
            if isinstance(m, list) and i < len(m) and isinstance(m[i], str):
                v = v or m[i]
            elif isinstance(m, dict):
                v = v or m.get(t) or m.get(LETTERS[i]) or ""
            new.append((t, e, v or ""))
        self.opts = new

    @property
    def key_text(self):
        return self.opts[self.key][0]

    @property
    def ah(self):
        return audit_hash(self.key_text, self.opts)

    def letter_of(self, text):
        for i, (t, _, _) in enumerate(self.opts):
            if norm(t) == norm(text):
                return LETTERS[i]
        return None


def quiz_files(folder, files_glob=None):
    out = []
    for f in sorted(glob.glob(os.path.join(ROOT, folder, "*.html"))):
        if files_glob and not fnmatch.fnmatch(os.path.basename(f), files_glob):
            continue
        out.append(f)
    return out


def load_folder(folder, files_glob=None, unparseable=None):
    qs = []
    for f in quiz_files(folder, files_glob):
        rel = os.path.relpath(f, ROOT)
        try:
            raw, _ = X.load(f)
        except X.Unparseable as e:
            if unparseable is not None:
                unparseable.append((rel, str(e)))
            continue
        for i, q in enumerate(raw or []):
            if len(q.get("opts", [])) != 4:
                continue                                    # the four-option rule is check_exam_standard's job
            qs.append(Q(folder, rel, i, q))
    return qs


# ----------------------------------------------------------------------------- source records (for `mis`)
_SRC = None


def _walk_questions(obj):
    if isinstance(obj, dict):
        if "q" in obj and "opts" in obj:
            yield obj
        else:
            for v in obj.values():
                yield from _walk_questions(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _walk_questions(v)


def source_index():
    """{normalized stem: record} and {sorted option texts: record} from tools/*_sets.json and the pool
    modules that carry `mis`. A pool module is only executed when its text contains a `mis` field."""
    global _SRC
    if _SRC is not None:
        return _SRC
    by_stem, by_opts = {}, {}

    def add(r):
        try:
            texts = [o[0] if isinstance(o, (list, tuple)) else o["t"] for o in r["opts"]]
        except Exception:
            return
        by_stem.setdefault(norm(r["q"]), r)
        by_opts.setdefault(tuple(sorted(norm(t) for t in texts)), r)

    for p in sorted(glob.glob(os.path.join(HERE, "*.json"))):
        if not p.endswith("sets.json") and "_sets" not in p:
            continue
        try:
            for r in _walk_questions(json.load(open(p, encoding="utf-8"))):
                add(r)
        except Exception:
            continue
    for p in sorted(glob.glob(os.path.join(HERE, "*pool*.py")) + glob.glob(os.path.join(HERE, "*_vig_*.py"))):
        try:
            txt = open(p, encoding="utf-8").read()
        except Exception:
            continue
        if not re.search(r"""\bmis\s*=|["']mis["']""", txt):
            continue
        try:
            ns = runpy.run_path(p)
        except BaseException:
            continue
        for r in _walk_questions(ns.get("QUESTIONS", [])):
            add(r)
    _SRC = (by_stem, by_opts)
    return _SRC


def _src_opt_mis(r, text, i):
    opts = r["opts"]
    for j, o in enumerate(opts):
        t = o[0] if isinstance(o, (list, tuple)) else o.get("t", "")
        if norm(t) == norm(text):
            if isinstance(o, (list, tuple)) and len(o) > 2 and isinstance(o[2], str):
                return o[2]
            if isinstance(o, dict) and o.get("mis"):
                return o["mis"]
            m = r.get("mis")
            if isinstance(m, list) and j < len(m) and isinstance(m[j], str):
                return m[j]
            if isinstance(m, dict):
                return m.get(text) or m.get(LETTERS[j]) or ""
    return ""


def attach_sources(qs):
    by_stem, by_opts = source_index()
    for q in qs:
        r = by_stem.get(norm(q.stem)) or by_opts.get(tuple(sorted(norm(o[0]) for o in q.opts)))
        if not r:
            continue
        q.sourced = True
        q.opts = [(t, e, m or _src_opt_mis(r, t, i)) for i, (t, e, m) in enumerate(q.opts)]
        if not q.cite and isinstance(r.get("cite"), str):
            q.cite = r["cite"]
    return qs


# ----------------------------------------------------------------------------- evidence (deck text)
def class_of(folder):
    return re.sub(r" (Exam \d+|Book Questions)$", "", folder)


_DECKS = {}


def deck_index(cls):
    if cls not in _DECKS:
        idx = {}
        base = os.path.join(INBOX, cls + " Inbox")
        for dp, dn, fn in os.walk(base):
            dn[:] = [d for d in dn if d != "recordings"]
            for f in fn:
                if f.lower().endswith((".pptx", ".ppt")) and not f.startswith(("~", ".")):
                    idx.setdefault(norm(f).lower(), os.path.join(dp, f))
        _DECKS[cls] = idx
    return _DECKS[cls]


def find_deck(cls, name):
    idx = deck_index(cls)
    k = norm(name).lower()
    if k in idx:
        return idx[k]
    stems = {os.path.splitext(a)[0]: a for a in idx}
    m = difflib.get_close_matches(os.path.splitext(k)[0], list(stems), n=1, cutoff=0.9)
    if m:
        return idx[stems[m[0]]]
    base = os.path.splitext(k)[0]
    pre = [a for st, a in stems.items() if len(base) >= 12 and (st.startswith(base) or base.startswith(st))]
    return idx[pre[0]] if len(pre) == 1 else None


def file_hash(path):
    h = hashlib.sha1()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def extract_deck(path):
    """{slide number: text with its '=== SLIDE N ===' marker}. Text frames, tables and groups; pictures
    and charts are flagged because their content is NOT in this text. Speaker notes are not included
    (slides-only grounding)."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE as T
    prs = Presentation(path)
    out = {}
    for n, slide in enumerate(prs.slides, 1):
        lines, flags = [], {"picture": 0, "chart": 0}

        def walk(shapes):
            for sh in shapes:
                st = getattr(sh, "shape_type", None)
                if st == T.GROUP:
                    walk(sh.shapes)
                    continue
                if st == T.PICTURE:
                    flags["picture"] += 1
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    flags["chart"] += 1
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        cells = [norm(c.text) for c in row.cells]
                        if any(cells):
                            lines.append(" | ".join(cells))
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        t = norm("".join(r.text for r in p.runs) or p.text)
                        if t:
                            lines.append(t)
        walk(slide.shapes)
        if flags["picture"] or flags["chart"]:
            lines.append("[this slide has %d picture(s) and %d chart(s); their content is NOT in this text]"
                         % (flags["picture"], flags["chart"]))
        out[n] = "=== SLIDE %d ===\n%s" % (n, "\n".join(lines))
    return out


def deck_slides(path):
    """Cached extraction, keyed by the deck's CONTENT hash, kept outside the repo (private)."""
    os.makedirs(CACHE, mode=0o700, exist_ok=True)
    cp = os.path.join(CACHE, file_hash(path) + ".json")
    if os.path.exists(cp):
        return {int(k): v for k, v in json.load(open(cp, encoding="utf-8")).items()}
    slides = extract_deck(path)
    json.dump({str(k): v for k, v in slides.items()}, open(cp, "w", encoding="utf-8"), ensure_ascii=False)
    return slides


_HASHES = {}


def deck_hash(path):
    if path not in _HASHES:
        _HASHES[path] = file_hash(path)
    return _HASHES[path]


_UNCONVERTED = []                                           # pictures sips could not convert (e.g. some WMF/EMF)


def _to_png(src, dst):
    """EMF/WMF/TIFF/GIF/BMP -> PNG with the macOS `sips`. True on success. (qlmanage was tried for the formats
    sips rejects and hangs on large WMF, so it is not used: those pictures are skipped and counted.)"""
    import subprocess
    try:
        r = subprocess.run(["sips", "-s", "format", "png", src, "--out", dst], capture_output=True, timeout=30)
        return r.returncode == 0 and os.path.exists(dst) and os.path.getsize(dst) > 0
    except Exception:
        return False


def extract_images(path, outdir):
    """{slide number: [absolute png/jpg paths]} for every slide picture that is not decorative. Pictures under
    ~8 KB or under 64x64 pixels are skipped; EMF/WMF/TIFF and other formats are converted to PNG."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE as T
    prs = Presentation(path)
    out = {}
    for n, slide in enumerate(prs.slides, 1):
        found, seen = [], set()

        def walk(shapes):
            for sh in shapes:
                if getattr(sh, "shape_type", None) == T.GROUP:
                    walk(sh.shapes)
                    continue
                try:
                    img = sh.image
                    blob, ext = img.blob, (img.ext or "").lower().lstrip(".")
                except Exception:
                    continue
                if len(blob) < MIN_IMG_BYTES:
                    continue
                try:
                    w, h = img.size
                    if w < MIN_IMG_PX and h < MIN_IMG_PX:
                        continue
                except Exception:
                    pass                                    # EMF/WMF: size unknown, keep (the byte floor passed)
                d = hashlib.sha1(blob).hexdigest()
                if d in seen:
                    continue
                seen.add(d)
                found.append((blob, ext))
        walk(slide.shapes)
        paths = []
        for k, (blob, ext) in enumerate(found, 1):
            base = os.path.join(outdir, "s%03d-%d" % (n, k))
            if ext in ("png", "jpg", "jpeg"):
                with open(base + "." + ext, "wb") as fh:
                    fh.write(blob)
                paths.append(base + "." + ext)
            else:
                raw = base + ".orig." + (ext or "bin")
                with open(raw, "wb") as fh:
                    fh.write(blob)
                ok = _to_png(raw, base + ".png")
                os.remove(raw)
                if ok:
                    paths.append(base + ".png")
                else:
                    _UNCONVERTED.append("%s slide %d (.%s)" % (os.path.basename(path)[:40], n, ext))
        if paths:
            out[n] = paths
    return out


def deck_images(path):
    """Cached per deck CONTENT hash under ~/.cache/pa-quizzes/slides-img/ (outside the repo, never copied into it)."""
    outdir = os.path.join(IMG_CACHE, deck_hash(path))
    idx = os.path.join(outdir, "index.json")
    if os.path.exists(idx):
        return {int(k): [os.path.join(outdir, f) for f in v] for k, v in json.load(open(idx, encoding="utf-8")).items()}
    os.makedirs(outdir, mode=0o700, exist_ok=True)
    got = extract_images(path, outdir)
    tmp = idx + ".tmp"
    json.dump({str(k): [os.path.basename(p) for p in v] for k, v in got.items()}, open(tmp, "w", encoding="utf-8"))
    os.replace(tmp, idx)                                    # the index appears only when extraction finished
    return got


def has_pictures(text):
    return "picture(s)" in text and not re.search(r"has 0 picture\(s\)", text)


_IMGS = {}                                                  # "<deck>|<slide>" -> [paths] for slides evidence_for resolved


def parse_cite(cite):
    """'Deck.pptx, Slide 12; Other, Slides 3-4' -> [(deck name, [slide numbers])]. The deck name may
    omit its extension; a segment with only 'Slide 5' inherits the previous deck."""
    out, last = [], None
    for seg in re.split(r"[;\n]", cite or ""):
        seg = seg.strip()
        if not seg:
            continue
        m = re.match(r"(.*?\.pptx?)", seg, re.I) or re.match(r"(.*?)[\s,;:(\-]*\bSlides?\b", seg, re.I)
        deck = m.group(1).strip(" ,;:-") if m and m.group(1).strip(" ,;:-") else last
        slides = sorted(B.slides_in(seg))
        if deck and slides:
            out.append((deck, slides))
            last = deck
    return out


def evidence_for(q, cache):
    """-> ([(deck file name, window slide numbers, {n: text})], None) or (None, reason). Every deck the cite
    names must resolve, or the item is not reviewable (a judge must not rule on partial evidence)."""
    parts = parse_cite(q.cite)
    if not parts:
        return None, "no deck + slide number in the cite %r" % q.cite[:80]
    cls = class_of(q.folder)
    merged = {}
    for deck, slides in parts:
        merged.setdefault(deck, set()).update(slides)
    out = []
    for deck, slides in merged.items():
        path = find_deck(cls, deck)
        if not path:
            return None, "deck %r not found in %s Inbox" % (deck[:60], cls)
        if path.lower().endswith(".ppt"):
            return None, "old .ppt deck cannot be read (python-pptx): %s" % deck[:60]
        if path not in cache:
            cache[path] = deck_slides(path)
        allslides = cache[path]
        if not any(s in allslides for s in slides):
            return None, "cited slide(s) %s not in %s (%d slides)" % (sorted(slides), deck[:50], len(allslides))
        want = set()
        for s_ in slides:
            want.update((s_ - 1, s_, s_ + 1))
        window = sorted(s_ for s_ in want if s_ in allslides)
        if any(has_pictures(allslides[s_]) for s_ in window):
            try:
                imgs = deck_images(path)
            except Exception as e:                           # a broken picture must not stop the text review
                print("   (could not extract pictures of %s: %s)" % (os.path.basename(path)[:50], str(e)[:80]))
                imgs = {}
            for s_ in window:
                if imgs.get(s_):
                    _IMGS["%s|%d" % (os.path.basename(path), s_)] = imgs[s_]
        out.append((os.path.basename(path), window, {s_: allslides[s_] for s_ in window}))
    return out, None


def window_lines(texts):
    """-> (normalized joined text, set of normalized lines) for quote checking."""
    joined = qnorm(" ".join(texts))
    lines = {qnorm(l) for t in texts for l in t.split("\n")}
    return joined, lines


def image_quote_ok(quote, m):
    """'[image] <deck>|<slide>' stands for what is clearly visible in that slide's picture. Accepted only when
    that slide has listed images in the item's manifest entry."""
    qn = qnorm(quote)
    if not qn.startswith("[image]"):
        return False
    ref = qn[len("[image]"):].strip()
    return any(qnorm(k) == ref and v for k, v in (m.get("evidence_images") or {}).items())


def quote_ok(quote, joined, lines):
    """Verbatim (after normalization) and not a stub: 5+ words, or a whole line of the evidence."""
    qn = qnorm(quote)
    if not qn:
        return False
    if qn not in joined:
        return False
    return len(qn.split()) >= 5 or qn in lines


# ----------------------------------------------------------------------------- records
def load_records(path=None):
    p = path or RECORDS
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else {}


def save_records(rec, path=None):
    p = path or RECORDS
    tmp = p + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:             # one line per item keeps diffs reviewable
        fh.write("{\n")
        keys = sorted(rec)
        for i, k in enumerate(keys):
            fh.write(" %s: %s%s\n" % (json.dumps(k), json.dumps(rec[k], ensure_ascii=False, sort_keys=True),
                                      "," if i < len(keys) - 1 else ""))
        fh.write("}\n")
    os.replace(tmp, p)                                      # never a half-written records file


def records_lock():
    """Several judges import at once (2026-10-15: six in parallel): hold an exclusive lock from load to save so no
    import overwrites another's records. The lock file lives outside the repo; it is released when the process ends."""
    import fcntl
    d = os.path.expanduser("~/.cache/pa-quizzes")
    os.makedirs(d, exist_ok=True)
    fh = open(os.path.join(d, "blind_review.lock"), "w")
    fcntl.flock(fh, fcntl.LOCK_EX)
    return fh


def blind_of(rec, h):
    return (rec.get(h) or {}).get("passes", {}).get("blind")


def audit_current(rec, q):
    a = (rec.get(q.h) or {}).get("passes", {}).get("audit")
    return a if a and a.get("ah") == q.ah and blind_of(rec, q.h) else None


def load_scope(path=None):
    p = path or SCOPE
    if not os.path.exists(p):
        return []
    return json.load(open(p, encoding="utf-8")).get("globs", [])


def in_scope(rel, globs):
    return any(fnmatch.fnmatch(rel, g) for g in globs)


def dedupe(qs):
    """One unit per hash: a master-exam copy of a topic question is the same item."""
    seen, out, dup = {}, [], 0
    for q in qs:
        if q.h in seen:
            dup += 1
            continue
        seen[q.h] = q
        out.append(q)
    return out, dup


# ----------------------------------------------------------------------------- export
def refuse_inside_repo(out):
    ap = os.path.realpath(out)
    if ap == ROOT or ap.startswith(ROOT + os.sep):
        sys.exit("refusing --out inside the repo: batches hold deck slide text and the repo is public. "
                 "Use a directory outside it (e.g. ~/.cache/pa-quizzes/review/<name>).")


def ev_images(ev):
    """{"<deck>|<slide>": [absolute image paths]} for the evidence slides that have extracted pictures."""
    out = {}
    for deck, window, _ in ev:
        for s_ in window:
            k = "%s|%d" % (deck, s_)
            if _IMGS.get(k):
                out[k] = list(_IMGS[k])
    return out


def item_payload(q, ev, pass_, rec):
    d = {"id": q.h, "stem": q.stem, "options": {LETTERS[i]: o[0] for i, o in enumerate(q.opts)},
         "evidence_slides": {deck: window for deck, window, _ in ev}}
    imgs = ev_images(ev)
    if imgs:
        d["evidence_images"] = imgs
    if pass_ == "audit":
        b = blind_of(rec, q.h)
        d["key"] = LETTERS[q.key]
        d["explanations"] = {LETTERS[i]: o[1] for i, o in enumerate(q.opts)}
        mis = {LETTERS[i]: o[2] for i, o in enumerate(q.opts) if o[2] and i != q.key}
        if mis:
            d["mis"] = mis
        d["your_blind_result"] = {
            "pick": b["pick"] if b["pick"] in ("ambiguous", "insufficient") else q.letter_of(b["pick_text"]),
            "cover_answer": b["cover_answer"], "alt": q.letter_of(b["alt_text"]) if b.get("alt_text") else None,
            "missing": b.get("missing"), "conflict": b.get("conflict"), "clue": b.get("clue")}
    return d


def read_ids(path):
    """Item ids from FILE: a JSON list (strings, or objects with "id") or one id per line (# comments ok)."""
    txt = open(os.path.expanduser(path), encoding="utf-8").read()
    try:
        data = json.loads(txt)
    except ValueError:
        data = None
    if isinstance(data, dict):
        data = data.get("ids") or list(data)
    if isinstance(data, list):
        return [str(x.get("id") if isinstance(x, dict) else x).strip() for x in data if x]
    return [l.split("#")[0].strip() for l in txt.splitlines() if l.split("#")[0].strip()]    # one id per line


def manifest_path(d):
    """The manifest holds the writer's keys and the slide text, so it lives OUTSIDE the judge-visible export
    directory: ~/.cache/pa-quizzes/review-manifests/<sha1 of the absolute export path>.json."""
    return os.path.join(MANIFESTS, hashlib.sha1(os.path.abspath(os.path.expanduser(d)).encode("utf-8")).hexdigest() + ".json")


def load_manifest(d):
    """Cache lookup first; DIR/manifest.json only for directories exported before the move."""
    for p in (manifest_path(d), os.path.join(d, "manifest.json")):
        if os.path.exists(p):
            return json.load(open(p, encoding="utf-8"))
    sys.exit("no manifest for %s (looked in %s and %s)" % (d, manifest_path(d), os.path.join(d, "manifest.json")))


def cmd_export(a):
    refuse_inside_repo(a.out)
    rec = load_records(a.records)
    unparse = []
    qs = attach_sources(load_folder(a.folder, a.files, unparse))
    for rel, err in unparse:
        print("UNPARSEABLE (not exported): %s: %s" % (rel, err))
    total = len(qs)
    qs, dup = dedupe(qs)
    if a.ids:
        want = read_ids(a.ids)
        have = {q.h for q in qs}
        unknown = [i for i in want if i not in have]
        qs = [q for q in qs if q.h in set(want)]
        print("--ids: %d id(s) listed, %d found in %s%s" % (len(set(want)), len(qs), a.folder, (" (--files %s)" % a.files) if a.files else ""))
        if unknown:
            print("   UNKNOWN id(s), not exported: %s" % ", ".join(unknown[:20]) + (" ..." if len(unknown) > 20 else ""))
    skipped_nb = 0
    if a.pass_ == "blind":
        pending = [q for q in qs if a.include_reviewed or not blind_of(rec, q.h)]
    else:
        nob = [q for q in qs if not blind_of(rec, q.h)]
        skipped_nb = len(nob) if a.only_reviewed else 0
        if nob:
            print("REFUSED: %d of %d item(s) have no blind record; the audit pass needs the blind pass first "
                  "(run export/import --pass blind). First few:" % (len(nob), len(qs)))
            for q in nob[:5]:
                print("   %s  %s" % (q.file, q.stem[:70]))
            if not a.only_reviewed:
                return 1
        pending = [q for q in qs if blind_of(rec, q.h) and (a.include_reviewed or not audit_current(rec, q))]
    already = len(qs) - len(pending) - skipped_nb
    if a.sample is not None and len(pending) > a.sample:
        pending = random.Random(a.seed).sample(sorted(pending, key=lambda q: q.h), a.sample)

    cache, ready, unresolved = {}, [], []
    for q in pending:
        ev, why = evidence_for(q, cache)
        if ev is None:
            unresolved.append({"id": q.h, "file": q.file, "stem": q.stem[:90], "why": why})
        else:
            ready.append((q, ev))
    ready.sort(key=lambda t: (t[1][0][0], t[1][0][1][0], t[0].h))

    os.makedirs(a.out, exist_ok=True)
    os.chmod(a.out, 0o700)
    batches = []
    if a.pass_ == "audit" and a.from_dir:
        old = {}
        for f in sorted(glob.glob(os.path.join(a.from_dir, "judge-*.json"))):
            old[os.path.basename(f)] = json.load(open(f, encoding="utf-8"))
        place = {}
        for name, b in old.items():
            for it in b["items"]:
                place[it["id"]] = name
        groups = {}
        for q, ev in ready:
            groups.setdefault(place.get(q.h, "~rest"), []).append((q, ev))
        for name in sorted(groups):
            batches.append((name, groups[name]))
    else:
        for i in range(0, len(ready), a.size):
            batches.append(("judge-%03d.json" % (i // a.size + 1), ready[i:i + a.size]))

    manifest = {"pass": a.pass_, "prompt_version": PROMPT_VERSION, "folder": a.folder, "items": {}}
    for n, (name, group) in enumerate(batches, 1):
        evid, items = {}, []
        for q, ev in group:
            for deck, window, texts in ev:
                for s in window:
                    evid.setdefault(deck, {})[s] = texts[s]
            items.append(item_payload(q, ev, a.pass_, rec))
            manifest["items"][q.h] = {
                "folder": q.folder, "file": q.file, "stem": q.stem,
                "evidence_slides": {deck: window for deck, window, _ in ev},
                "letters": {LETTERS[i]: o[0] for i, o in enumerate(q.opts)},
                "key": LETTERS[q.key], "key_text": q.key_text, "ah": q.ah,
                "item_text": [q.stem] + [o[0] for o in q.opts] + [o[1] for o in q.opts],
                "evidence": {"%s|%d" % (deck, s): texts[s] for deck, window, texts in ev for s in window},
                "evidence_images": ev_images(ev),
                "batch": name}
        batch = {"pass": a.pass_, "prompt_version": PROMPT_VERSION,
                 "instructions": "tools/question_writing/REVIEW_%s.md" % a.pass_.upper(),
                 "evidence": {d: "\n".join(v[s] for s in sorted(v)) for d, v in evid.items()},
                 "items": items}
        out_name = name if name.startswith("judge-") else "judge-%03d.json" % n
        json.dump(batch, open(os.path.join(a.out, out_name), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        for it in items:
            manifest["items"][it["id"]]["batch"] = out_name
    os.makedirs(MANIFESTS, mode=0o700, exist_ok=True)
    mp = manifest_path(a.out)
    json.dump(manifest, open(mp, "w", encoding="utf-8"), ensure_ascii=False)
    os.chmod(mp, 0o600)
    legacy = os.path.join(a.out, "manifest.json")
    if os.path.exists(legacy):
        os.remove(legacy)                                   # an older export left keys in the judge-visible folder
    # unresolved.json is judge-visible: id, file, stem and the reason only -- never a key or an explanation
    json.dump([{k: u[k] for k in ("id", "file", "stem", "why")} for u in unresolved],
              open(os.path.join(a.out, "unresolved.json"), "w", encoding="utf-8"), indent=1)
    print("scope: %d question(s) read, %d duplicate copy(ies) folded (master reuse), %d unique item(s)"
          % (total, dup, len(qs)))
    skip_note = ("; %d skipped for having no blind record (--only-reviewed)" % skipped_nb) if skipped_nb else ""
    print("exported %d item(s) in %d batch(es) -> %s   (%d already have a current %s record%s; %d not reviewable: see unresolved.json)"
          % (len(ready), len(batches), a.out, already, a.pass_, skip_note, len(unresolved)))
    if _UNCONVERTED:
        print("   (%d picture(s) could not be converted by sips and are NOT listed, e.g. %s)" % (len(_UNCONVERTED), _UNCONVERTED[0]))
    nimg = sum(len(v) for q, ev in ready for v in ev_images(ev).values())
    if nimg:
        print("pictures: %d image file(s) listed as evidence_images (cached under %s; open them with Read)" % (nimg, IMG_CACHE))
    for u in unresolved[:10]:
        print("   unresolved: %s | %s | %s" % (u["file"], u["stem"][:50], u["why"]))
    return 0


# ----------------------------------------------------------------------------- import
def today():
    return datetime.date.today().isoformat()


def words(t):
    return len((t or "").split())


def check_blind(e, m):
    """-> (record or None, rejection reason or None)."""
    pick = e.get("pick")
    if pick not in tuple(LETTERS) + ("ambiguous", "insufficient"):
        return None, "pick must be A-D, ambiguous or insufficient (got %r)" % (pick,)
    if not norm(e.get("cover_answer") or ""):
        return None, "cover_answer is empty (answer from the stem alone BEFORE reading the options)"
    why = norm(e.get("why") or "")
    if not why:
        return None, "why is empty"
    if words(why) > WHY_WORDS:
        return None, "why is %d words (limit %d)" % (words(why), WHY_WORDS)
    alt = e.get("alt")
    if alt is not None and (alt not in tuple(LETTERS) or alt == pick):
        return None, "alt must be another letter A-D or null (got %r)" % (alt,)
    joined, lines = window_lines(list(m["evidence"].values()))
    quote = e.get("quote")
    if pick in tuple(LETTERS) and not norm(quote or ""):
        return None, "a letter pick needs a verbatim quote from the evidence"
    if norm(quote or "") and not quote_ok(quote, joined, lines) and not image_quote_ok(quote, m):
        if qnorm(quote).startswith("[image]"):
            return None, "image quote %r: that slide has no listed evidence_images (use '[image] <deck>|<slide>' for a listed slide)" % norm(quote)[:80]
        return None, "quote is not verbatim in the evidence (or under 5 words and not a whole line)"
    r = {"pick": pick, "pick_text": m["letters"].get(pick) if pick in tuple(LETTERS) else None,
         "cover_answer": norm(e["cover_answer"]), "why": why, "quote": norm(quote or "")[:QUOTE_KEEP] or None,
         "alt": alt, "alt_text": m["letters"].get(alt) if alt else None}
    for f in ("missing", "conflict", "clue"):
        v = e.get(f)
        r[f] = norm(v) if isinstance(v, str) and norm(v) else None
    return r, None


def check_audit(e, m):
    v = e.get("verdict")
    if v not in ("pass", "revise", "hold"):
        return None, "verdict must be pass, revise or hold (got %r)" % (v,)
    defects = e.get("defects") or []
    if not isinstance(defects, list):
        return None, "defects must be a list"
    if v == "pass" and defects:
        return None, "a pass cannot list defects (use revise or hold)"
    if v != "pass" and not defects:
        return None, "%s needs at least one defect" % v
    if v == "revise" and not norm(e.get("fix") or ""):
        return None, "revise needs the smallest correction in `fix`"
    joined, lines = window_lines(list(m["evidence"].values()))
    ijoined, ilines = window_lines(m["item_text"])
    out = []
    for d in defects:
        if not isinstance(d, dict) or not norm(d.get("field") or "") or not norm(d.get("reason") or ""):
            return None, "every defect needs field and reason"
        if not norm(d.get("quote") or ""):
            return None, "every defect needs a verbatim quote (slide text, or the item's own text for an unsupported claim)"
        if quote_ok(d["quote"], joined, lines):
            src = "slide"
        elif image_quote_ok(d["quote"], m):
            src = "image"
        elif quote_ok(d["quote"], ijoined, ilines):
            src = "item"
        else:
            return None, "defect quote %r is not verbatim in the slides or the item" % norm(d["quote"])[:50]
        out.append({"field": norm(d["field"]), "reason": norm(d["reason"]), "quote": norm(d["quote"])[:QUOTE_KEEP], "src": src})
    return {"verdict": v, "defects": out, "fix": norm(e.get("fix") or "") or None}, None


def cmd_import(a):
    manifest = load_manifest(a.dir)
    if manifest["pass"] != a.pass_:
        sys.exit("manifest in %s is a %s export, not %s" % (a.dir, manifest["pass"], a.pass_))
    if manifest["prompt_version"] != PROMPT_VERSION:
        sys.exit("export was prompt %s, this tool is %s" % (manifest["prompt_version"], PROMPT_VERSION))
    lock = records_lock()                                   # noqa: F841 (held until the process exits)
    rec = load_records(a.records)
    res = {"accepted": 0, "rejected": 0, "missing": 0}
    repair, seen = [], set()
    for f in sorted(glob.glob(os.path.join(a.dir, "verdict-*.json"))):
        for e in json.load(open(f, encoding="utf-8")):
            k = e.get("id")
            m = manifest["items"].get(k)
            if not m:
                continue
            seen.add(k)
            if a.pass_ == "audit" and not blind_of(rec, k):
                res["rejected"] += 1
                repair.append({"id": k, "file": m["file"], "why": "REFUSED: audit result for an item with no blind record"})
                continue
            r, why = (check_blind if a.pass_ == "blind" else check_audit)(e, m)
            if r is None:
                res["rejected"] += 1
                repair.append({"id": k, "file": m["file"], "stem": m["stem"][:80], "why": why})
                continue
            r.update({"judge": a.judge, "date": today(), "prompt": PROMPT_VERSION})
            if a.pass_ == "audit":
                r["ah"] = m["ah"]
                r["key_text"] = m["key_text"]
            slot = rec.setdefault(k, {"meta": {}, "passes": {}})
            slot["meta"] = {"folder": m["folder"], "file": m["file"], "evidence_slides": m["evidence_slides"]}
            slot["passes"][a.pass_] = r
            if a.pass_ == "blind":
                slot["passes"].pop("audit", None)               # a fresh blind read invalidates an old audit
            res["accepted"] += 1
    res["missing"] = len(set(manifest["items"]) - seen)
    save_records(rec, a.records)
    json.dump(repair, open(os.path.join(a.dir, "repair.json"), "w", encoding="utf-8"), indent=1)
    print("imported (%s): %s | rejected items listed in %s" % (a.pass_, res, os.path.join(a.dir, "repair.json")))
    for r in repair[:10]:
        print("   REJECTED %s | %s" % (r["id"], r["why"]))
    return 1 if res["rejected"] else 0


# ----------------------------------------------------------------------------- report / check
def gather(folders, scope_only=False, files_glob=None):
    globs = load_scope() if scope_only else None
    allq, unparse = [], []
    for f in folders:
        for q in load_folder(f, files_glob, unparse):
            if scope_only and not in_scope(q.file, globs):
                continue
            allq.append(q)
    return allq, unparse


def snip(q, n=70):
    return "%s | %s" % (q.file, q.stem[:n])


def cmd_report(a):
    rec = load_records(a.records)
    folders = a.folder or s2_folders()
    qs, unparse = gather(folders)
    uq, dup = dedupe(qs)
    n = len(uq)
    have_b = [q for q in uq if blind_of(rec, q.h)]
    wrong, amb, ins, alts, revise, hold, passed, stale, disagree_pass = [], [], [], [], [], [], [], [], []
    conflicts, missings = [], []
    for q in have_b:
        b = blind_of(rec, q.h)
        if b["pick"] in LETTERS and norm(b["pick_text"]) != norm(q.key_text):
            wrong.append(q)
        elif b["pick"] == "ambiguous":
            amb.append(q)
        elif b["pick"] == "insufficient":
            ins.append(q)
        if b.get("alt"):
            alts.append(q)
        if b.get("conflict"):
            conflicts.append(q)
        if b.get("missing"):
            missings.append(q)
    for q in have_b:
        a_ = (rec[q.h]["passes"].get("audit"))
        if not a_:
            continue
        if a_["ah"] != q.ah:
            stale.append(q)
        elif a_["verdict"] == "revise":
            revise.append(q)
        elif a_["verdict"] == "hold":
            hold.append(q)
        else:
            passed.append(q)
            b = blind_of(rec, q.h)
            if b["pick"] not in LETTERS or norm(b["pick_text"]) != norm(q.key_text):
                disagree_pass.append(q)
    current = {q.h for q in uq}
    orphans = [k for k in rec if k not in current]

    def show(title, items, extra=None):
        print("\n%s: %d" % (title, len(items)))
        for q in items[:a.limit]:
            print("   %s%s" % (snip(q), ("  -> " + extra(q)) if extra else ""))
        if len(items) > a.limit:
            print("   ... %d more (--limit)" % (len(items) - a.limit))

    print("Blind review report  (prompt %s)  folders: %s" % (PROMPT_VERSION, ", ".join(folders) if a.folder else "all Semester 2"))
    print("questions read: %d, unique items: %d (%d master/duplicate copies fold into them)" % (len(qs), n, dup))
    print("blind reviewed: %d / %d   audited (current): %d / %d   audit PASS: %d / %d"
          % (len(have_b), n, len(passed) + len(revise) + len(hold), n, len(passed), n))
    show("WRONG-KEY SUSPECTS (blind pick != key)", wrong,
         lambda q: "key %s, judge %s: %s" % (LETTERS[q.key], q.letter_of(blind_of(rec, q.h)["pick_text"]), blind_of(rec, q.h)["why"][:80]))
    show("AMBIGUOUS", amb, lambda q: blind_of(rec, q.h)["why"][:90])
    show("INSUFFICIENT", ins, lambda q: (blind_of(rec, q.h).get("missing") or blind_of(rec, q.h)["why"])[:90])
    show("second defensible answer named (alt)", alts, lambda q: "alt %s" % q.letter_of(blind_of(rec, q.h)["alt_text"]))
    show("blind CONFLICT flagged (slide vs slide, or slide vs established fact)", conflicts, lambda q: blind_of(rec, q.h)["conflict"][:90])
    show("blind MISSING detail flagged (stem or slides lack something decisive)", missings, lambda q: blind_of(rec, q.h)["missing"][:90])
    show("AUDIT REVISE", revise, lambda q: "; ".join(d["reason"][:50] for d in rec[q.h]["passes"]["audit"]["defects"][:2]))
    show("AUDIT HOLD", hold, lambda q: "; ".join(d["reason"][:50] for d in rec[q.h]["passes"]["audit"]["defects"][:2]))
    show("audit PASS although the blind judge disagreed with the key (read these; agreement/pass is not proof)", disagree_pass)
    show("STALE audit (explanations or key edited since)", stale)
    print("\nrecords for items no longer on any page: %d" % len(orphans))
    for rel, err in unparse:
        print("UNPARSEABLE %s: %s" % (rel, err))
    print("report: %d wrong-key suspect(s), %d ambiguous, %d insufficient, %d revise, %d hold, "
          "%d conflict flag(s), %d missing flag(s) of %d unique items (%d blind-reviewed)  (exit 0)"
          % (len(wrong), len(amb), len(ins), len(revise), len(hold), len(conflicts), len(missings), n, len(have_b)))
    return 0


def cmd_check(a):
    rec = load_records(a.records)
    globs = load_scope()
    unparse = []
    qs = load_folder(a.folder, None, unparse)
    scoped = qs if a.all else [q for q in qs if in_scope(q.file, globs)]
    bad = []
    reused = 0
    for q in scoped:
        b = blind_of(rec, q.h)
        au = audit_current(rec, q)
        why = None
        if not b:
            why = "no blind record"
        elif not au:
            old = rec[q.h]["passes"].get("audit")
            why = "audit is stale (key or an explanation changed)" if old else "no audit record"
        elif au["verdict"] != "pass":
            why = "audit verdict %s" % au["verdict"]
        if why:
            bad.append((q, why))
        elif rec[q.h]["meta"].get("file") != q.file:
            reused += 1
    for rel, err in unparse:
        print("UNPARSEABLE %s: %s -- NOT CHECKED" % (rel, err))
    if not scoped:
        print("blind review: %s: 0 question(s) in scope (scope.json globs: %d) -- nothing to check (exit 0)" % (a.folder, len(globs)))
        return 1 if unparse else 0
    for q, why in bad[:20]:
        print("   FAIL %s | %s" % (snip(q, 60), why))
    if len(bad) > 20:
        print("   ... %d more" % (len(bad) - 20))
    ok = len(scoped) - len(bad)
    print("blind review: %s: %d/%d in-scope question(s) have a current blind + audit PASS (%d via reused record) -> %s (exit %d)"
          % (a.folder, ok, len(scoped), reused, "FAIL" if bad else "PASS", 1 if (bad or unparse) else 0))
    return 1 if (bad or unparse) else 0


# ----------------------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--records", default=None, help="records file (default tools/blind_review.json; tests use a temp file)")
    ap.add_argument("--scope", default=None, help="scope file (default tools/question_writing/scope.json)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--pass", dest="pass_", choices=("blind", "audit"), required=True)
    e.add_argument("--folder", required=True)
    e.add_argument("--files", default=None, help="glob on file names inside the folder, e.g. '*-quiz*.html'")
    e.add_argument("--sample", type=int, default=None)
    e.add_argument("--seed", type=int, default=1)
    e.add_argument("--size", type=int, default=BATCH)
    e.add_argument("--include-reviewed", action="store_true")
    e.add_argument("--ids", default=None, help="FILE with item ids (one per line, or a JSON list): export only these; unknown ids are reported")
    e.add_argument("--from", dest="from_dir", default=None, help="audit: reuse the batching/evidence of this blind export dir")
    e.add_argument("--only-reviewed", action="store_true", help="audit: export the items that have a blind record, skip the rest")
    e.add_argument("--out", required=True)
    i = sub.add_parser("import")
    i.add_argument("--pass", dest="pass_", choices=("blind", "audit"), required=True)
    i.add_argument("dir")
    i.add_argument("--judge", required=True, help="judge model id recorded with every verdict, e.g. claude-sonnet-5-5")
    r = sub.add_parser("report")
    r.add_argument("--folder", action="append")
    r.add_argument("--limit", type=int, default=25)
    c = sub.add_parser("check")
    c.add_argument("--folder", required=True)
    c.add_argument("--all", action="store_true", help="every question in the folder, not just scope.json")
    a = ap.parse_args(argv)
    global SCOPE
    SCOPE = a.scope or SCOPE
    if a.cmd == "export" and not os.path.isdir(os.path.join(ROOT, a.folder)):
        sys.exit("no such folder in repo: %r" % a.folder)
    return {"export": cmd_export, "import": cmd_import, "report": cmd_report, "check": cmd_check}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main())
