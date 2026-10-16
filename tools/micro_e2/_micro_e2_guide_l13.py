# -*- coding: utf-8 -*-
"""Section 13 of the Microbiology Exam 2 guide -- Miscellaneous Bacterial Agents of Disease (Dr. Fair).

Same shape as _micro_e2_guide_l11.py (FIG, DECK, fig(), TOC, SECTION, TEST, IMAGES, extract()):
build_micro_e2_guide.py imports it, calls extract() for the figures and splices TOC, SECTION and TEST.

Attribution: Dr. Fair. Her PAJ5200.* decks never carry a lecturer name; the calendar row for 16 October
2026 (3-5 pm) names Dr. Fair. The recording will confirm it.

Instructional objectives are VERBATIM from the syllabus (Micro.pdf, page 4): the two numbered objectives
under "Misc. Bacterial Agents of Human Diseases", which points to Appendix 2's organism list (page 8). The
deck's own objective slide (slide 2) matches them. Objective 1 is answered organism by organism (13.1-13.15),
objective 2 is drawn together in 13.16.

Content is from the VISIBLE slides only, text and pictures (speaker notes do not count: "prostitute's pupil",
"CJT", the Lyme state ranking and "un-Gramstainable" are notes-only and left out). Pictures read at full size,
including the ones that carry content no text box has: slide 11 (the syphilis stage table), slide 25 (the
relapsing-fever graph), slide 28 (the Lyme two-year cycle), slide 33 (the cholera water diagram), slide 41
(the urease gel), slide 47 (the rickettsia table), slide 49 (the spotted-fever cycle), slide 57 (the
chlamydia life cycle), slides 65-66 and 68-70 (the dental diagrams).

AUDIO: not yet folded in (the recording is pending); phase 2 adds the recording box.
"""
import os

FIG = "micro-exam-2-study-guide-images"
DECK = "PAJ5200.Misc inf agents-2.pptx"
INBOX_DECK = os.path.expanduser(
    "~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 2/PAJ5200.Misc inf agents-2.pptx")

# (slide, picture index on that slide in shape-tree order counting group members and placeholder pictures,
# published file name). Multi-picture slides resolved against each picture's content: on slide 14 picture 1 is
# the gumma, picture 2 the pupil; on slide 33 picture 2 is panel (b), the cholera-toxin water loss.
IMAGES = [
    (4, 1, "l13-s004-spirochete-structure.jpg"),
    (13, 1, "l13-s013-secondary-syphilis-rash.jpg"),
    (14, 1, "l13-s014-gumma.jpg"),
    (15, 1, "l13-s015-congenital-syphilis.jpg"),
    (25, 1, "l13-s025-relapsing-fever-pattern.jpg"),
    (28, 1, "l13-s028-lyme-cycle.jpg"),
    (30, 1, "l13-s030-erythema-migrans.jpg"),
    (33, 2, "l13-s033-cholera-water-loss.jpg"),
    (41, 1, "l13-s041-helicobacter-urease-test.jpg"),
    (49, 1, "l13-s049-spotted-fever-cycle.jpg"),
    (50, 1, "l13-s050-spotted-fever-rash.jpg"),
    (57, 1, "l13-s057-chlamydia-life-cycle.jpg"),
    (66, 1, "l13-s066-dental-disease-flow.jpg"),
]


def _pictures(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pictures(sh.shapes)
            continue
        try:
            sh.image
        except Exception:
            continue
        yield sh


def extract(imgdir):
    """Write IMAGES into imgdir from the deck (pictures, picture placeholders and grouped pictures)."""
    from pptx import Presentation
    prs = Presentation(INBOX_DECK)
    os.makedirs(imgdir, exist_ok=True)
    for slide, k, name in IMAGES:
        pics = list(_pictures(prs.slides[slide - 1].shapes))
        assert len(pics) >= k, "%s slide %d has %d picture(s)" % (DECK, slide, len(pics))
        img = pics[k - 1].image
        assert name.endswith("." + img.ext.replace("jpeg", "jpg")), (name, img.ext)
        open(os.path.join(imgdir, name), "wb").write(img.blob)
    return len(IMAGES)


def fig(name, w, h, alt, cap, slide, mh=True):
    note = "Copyright &copy; The McGraw-Hill Companies. " if mh else ""
    return ('<figure class="fig"><img width="%d" height="%d" loading="lazy" src="%s/%s" alt="%s">'
            '<figcaption>%s <span class="src">%sSource: %s, Slide %d.</span></figcaption></figure>'
            % (w, h, FIG, name, alt, cap, note, DECK, slide))


TOC = """
  <a class="top-link" href="#misc-bacteria">13 &middot; Miscellaneous Bacterial Agents of Disease</a>
  <a class="sub-link" href="#mba-spirochetes">13.1 Objective 1 &mdash; The spirochetes</a>
  <a class="sub-link" href="#mba-syphilis">13.2 Objective 1 &mdash; <em>Treponema pallidum</em> and syphilis</a>
  <a class="sub-link" href="#mba-treponematoses">13.3 Objective 1 &mdash; Nonsyphilitic treponematoses</a>
  <a class="sub-link" href="#mba-leptospira">13.4 Objective 1 &mdash; <em>Leptospira</em> and leptospirosis</a>
  <a class="sub-link" href="#mba-relapsing">13.5 Objective 1 &mdash; <em>Borrelia</em> and relapsing fever</a>
  <a class="sub-link" href="#mba-lyme">13.6 Objective 1 &mdash; <em>Borrelia burgdorferi</em> and Lyme disease</a>
  <a class="sub-link" href="#mba-cholera">13.7 Objective 1 &mdash; The curviform bacteria and <em>Vibrio cholerae</em></a>
  <a class="sub-link" href="#mba-seafood">13.8 Objective 1 &mdash; Vibrios carried by seafood</a>
  <a class="sub-link" href="#mba-campylobacter">13.9 Objective 1 &mdash; <em>Campylobacter</em></a>
  <a class="sub-link" href="#mba-helicobacter">13.10 Objective 1 &mdash; <em>Helicobacter pylori</em></a>
  <a class="sub-link" href="#mba-rickettsia">13.11 Objective 1 &mdash; Rickettsias and <em>Ehrlichia</em></a>
  <a class="sub-link" href="#mba-coxiella">13.12 Objective 1 &mdash; <em>Coxiella</em> and <em>Bartonella</em></a>
  <a class="sub-link" href="#mba-chlamydia">13.13 Objective 1 &mdash; <em>Chlamydia</em> and <em>Chlamydophila</em></a>
  <a class="sub-link" href="#mba-mycoplasma">13.14 Objective 1 &mdash; Mycoplasmas</a>
  <a class="sub-link" href="#mba-dental">13.15 Objective 1 &mdash; Bacteria in dental disease</a>
  <a class="sub-link" href="#mba-transmission">13.16 Objective 2 &mdash; Modes of transmission</a>
  <a class="sub-link" href="#mba-notes">13.17 Where the slides and current practice differ</a>
"""

SECTION_HEAD = """
<section class="deck" id="misc-bacteria">
  <h2 class="deck-title">13 &middot; Miscellaneous Bacterial Agents of Disease</h2>
  <div class="io-box">
    <h3>Instructional Objectives</h3>
    <ol>
      <li>Compare and contrast the morphology, medium, laboratory tests, biochemical reactions, diseases,
      clinical manifestations, methods of treatment, and prevention of miscellaneous bacterial agents of
      human diseases.</li>
      <li>Describe the modes of transmission by which humans acquire miscellaneous bacterial agents of
      human diseases.</li>
    </ol>
  </div>

  <div class="callout"><strong>Three odd shapes and three odd lifestyles.</strong> The
  <strong>spirochetes</strong> (<em>Treponema</em>, <em>Leptospira</em>, <em>Borrelia</em>) are Gram-negative
  spirals that swim with flagella <em>inside</em> the cell. The <strong>curviform</strong> Gram-negatives
  (<em>Vibrio</em>, <em>Campylobacter</em>, <em>Helicobacter</em>) are commas and curved rods of the gut. The
  <strong>unique morphologies</strong> are the obligate intracellular rickettsias and chlamydias and the
  wall-less mycoplasmas. The syllabus points to its Appendix 2 list for this lecture: <em>Treponema
  pallidum</em> and <em>Borrelia</em>; <em>Rickettsia</em>, <em>Ehrlichia</em>, <em>Anaplasma</em>,
  <em>Chlamydia</em> and <em>Chlamydophila</em>; <em>Helicobacter</em>; and the mycoplasmas. The deck adds
  <em>Leptospira</em>, the vibrios, <em>Campylobacter</em>, <em>Coxiella</em>, <em>Bartonella</em> and the
  bacteria of dental disease, and never names <em>Anaplasma</em> (13.17). One drug recurs:
  <strong>tetracycline</strong>, a very broad-spectrum antibiotic that targets the prokaryotic ribosome
  (<strong>30S subunit</strong>) (slide 53).</div>
"""

BODY = """
  <h3 class="sub" id="mba-spirochetes">13.1 &middot; Objective 1 &mdash; The spirochetes</h3>
  <ul>
    <li><strong>Gram-negative and spiral shaped</strong>.</li>
    <li><strong>Motile by endoflagella</strong>, which are found <strong>inside the cell</strong> (in the
    periplasmic space, under the outer membrane).</li>
    <li>Free-living saprobes or commensals of animals. Three genera: <em>Treponema</em>,
    <em>Leptospira</em>, <em>Borrelia</em>.</li>
    <li>While most bacteria have circular chromosomes, these narrow cells may have <strong>linear
    double-stranded deoxyribonucleic acid and plasmids</strong>, especially <em>Borrelia</em>.</li>
  </ul>
  """ + fig("l13-s004-spirochete-structure.jpg", 800, 507,
      "Top, a drawing of a spiral cell with endoflagella running along the cell body inside the outer membrane in the periplasmic space; below, three micrographs of spirochetes, two glowing orange spirals and one gray spiral.",
      "<b>Spirochete structure.</b> The endoflagella lie in the periplasmic space, between the cell body and the "
      "outer membrane, which is why spirochetes can swim without external flagella.", 4) + """

  <h3 class="sub" id="mba-syphilis">13.2 &middot; Objective 1 &mdash; <em>Treponema pallidum</em> and syphilis</h3>
  <table>
    <tr><th>Heading</th><th><em>Treponema pallidum</em>, the spirochete of syphilis</th></tr>
    <tr><td>Morphology</td><td>Thin, regular, coiled cells (genus <em>Treponema</em>)</td></tr>
    <tr><td>Habitat and host</td><td>The genus lives in the oral cavity, intestinal tract and perianal and
    perineal regions of humans and animals. For <em>T. pallidum</em>, <strong>humans are the natural
    hosts</strong></td></tr>
    <tr><td>Growth and culture</td><td>Strict parasites with complex growth requirements; <strong>extremely
    fastidious</strong> and sensitive to environmental conditions, so they <strong>cannot survive long outside
    the human host</strong>; <strong>difficult to culture: they require live host cells</strong> for
    cultivation in the laboratory</td></tr>
    <tr><td>Transmission</td><td><strong>Sexual and transplacental</strong></td></tr>
    <tr><td>Pathogenesis</td><td>Binds to epithelial tissue (mucous membrane or abraded skin), multiplies,
    then <strong>penetrates capillaries</strong>, moves into the circulation and multiplies further</td></tr>
    <tr><td>Communicability</td><td>Spirochetes appear in lesions and in blood during the <strong>first two
    stages</strong>; syphilis is <strong>communicable in the primary and secondary stages</strong></td></tr>
    <tr><td>Diagnosis</td><td>The stages can <strong>mimic other diseases</strong>, so clinicians weigh
    symptoms, history, microscopy and serology. False-positive rapid plasma reagin results can occur with
    human immunodeficiency virus infection, Lyme disease, some pneumonias, malaria and lupus</td></tr>
    <tr><td>Treatment</td><td><strong>Penicillin G or doxycycline</strong>; erythromycin and tetracycline are
    less effective but can be used if penicillin is contraindicated, for example by allergy</td></tr>
  </table>
  <p><strong>The clinical stages of untreated syphilis</strong> (slides 9&ndash;10):</p>
  <ul>
    <li><strong>Primary</strong>: a <strong>hard chancre</strong> at the site of inoculation; it often heals
    spontaneously, so the illness may go undiagnosed.</li>
    <li><strong>Secondary</strong>: an upper-respiratory-like illness with fever, headache and sore throat,
    and a <strong>red or brown rash on the skin, palms and soles</strong>; the rash often disappears
    spontaneously.</li>
    <li><strong>Tertiary</strong>: reached by about 30% of infections; can last 20 years or longer; numerous
    complications in susceptible tissues and organs: <strong>neurological and cardiovascular</strong>
    symptoms, and <strong>gummas</strong>. The <strong>Argyll Robertson pupil</strong> (an irregular pupil
    with eventual loss of the light reflex) is a tertiary sign.</li>
    <li><strong>Congenital</strong>: <strong>transplacental</strong> infection of the fetus causes post-natal
    nasal discharge (<strong>&ldquo;snuffles&rdquo;</strong>), skin eruptions, <strong>bone
    deformation</strong> (<strong>&ldquo;Hutchinson&rsquo;s teeth&rdquo;</strong>) and nervous system
    abnormalities.</li>
  </ul>
  <p><strong>The stage table</strong> (slide 11, a picture):</p>
  <table>
    <tr><th>Stage</th><th>Average duration</th><th>Clinical setting</th><th>Diagnosis</th></tr>
    <tr><td>Incubation</td><td>3 weeks</td><td>No lesion; treponemes adhere to and penetrate the epithelium,
    multiply and disseminate</td><td>Asymptomatic phase</td></tr>
    <tr><td>Primary</td><td>2&ndash;6 weeks</td><td>The chancre appears at the inoculation site; intense
    treponemal activity in the body; the chancre later disappears</td><td>Dark-field microscopy; serological
    tests</td></tr>
    <tr><td>Primary latency</td><td>2&ndash;8 weeks</td><td>Healed chancre with little scarring; treponemes in
    the blood; few if any symptoms</td><td>Serological tests positive</td></tr>
    <tr><td>Secondary</td><td>2&ndash;6 weeks after the chancre leaves</td><td>Skin and mucous membrane
    lesions; hair loss; <strong>patient highly infectious</strong>; fever, lymphadenopathy; symptoms can
    persist for months</td><td>Dark-field testing of lesions; serological tests positive</td></tr>
    <tr><td>Latency</td><td>6 months to 8 or more years</td><td>Treponemes quiescent unless relapse occurs;
    lesions can reappear</td><td>Seropositive blood test</td></tr>
    <tr><td>Tertiary</td><td>Variable, up to 20 years</td><td>Neural and cardiovascular symptoms; gummas
    develop in organs; seropositivity</td><td>Treponeme may be shown by deoxyribonucleic acid analysis of
    tissue</td></tr>
  </table>
  <p>The table&rsquo;s treatment column is <strong>penicillin G</strong> (benzathine, or aqueous benzyl or
  procaine) at every stage, in double doses from the secondary stage on.</p>
  """ + fig("l13-s013-secondary-syphilis-rash.jpg", 1200, 527,
      "Left, a person's back and arms covered in a dense red spotted rash; right, two palms with scattered reddish-brown spots.",
      "<b>Secondary syphilis.</b> A widespread red or brown rash that, unusually, includes the "
      "<b>palms</b> (and soles).", 13) + """
  """ + fig("l13-s014-gumma.jpg", 800, 573,
      "A raised, red, granular, ulcerated nodule on the skin of a face below the eye.",
      "<b>Tertiary syphilis: a gumma.</b> A destructive granulomatous lesion; gummas develop in organs in "
      "the tertiary stage.", 14) + """
  """ + fig("l13-s015-congenital-syphilis.jpg", 795, 900,
      "Top, an infant with crusted discharge around the nose and mouth; bottom, a child's teeth with notched, peg-shaped upper incisors.",
      "<b>Congenital syphilis.</b> (a) <b>&ldquo;Snuffles&rdquo;</b>, the nasal discharge; (b) "
      "<b>Hutchinson&rsquo;s teeth</b>, notched incisors from the bone and tooth deformation.", 15) + """
  <p><strong>Serological tests</strong> (slide 17): the <strong>rapid plasma reagin</strong> test detects
  syphilis antibodies; the <strong>VDRL</strong> (Venereal Disease Research Laboratory) test is the older
  version of it; the <strong>FTA-ABS</strong> is the fluorescent treponemal antibody-absorption test. A blood
  test may still be required for a marriage license in one state; others impose waiting periods.</p>
  <p><strong>The &ldquo;Tuskegee Study&rdquo; of Untreated Syphilis (1932&ndash;1972)</strong> is the deck&rsquo;s
  history note (the book <em>Bad Blood</em> and the film <em>Miss Evers&rsquo; Boys</em> tell it).</p>

  <h3 class="sub" id="mba-treponematoses">13.3 &middot; Objective 1 &mdash; Nonsyphilitic treponematoses</h3>
  <p>The non-venereal treponematoses <strong>resemble syphilis but are rarely transmitted sexually or
  congenitally</strong>: cutaneous and bone diseases <strong>endemic to specific regions</strong>, usually arid
  climates and/or countries of the Global South.</p>
  <table>
    <tr><th>Disease</th><th>Agent</th><th>What it does</th><th>Where</th></tr>
    <tr><td><strong>Bejel</strong> (sahel)</td><td><em>T. pallidum</em> subspecies <em>endemicum</em></td>
    <td>A deforming <strong>childhood</strong> infection of the mouth, nasal cavity, body and hands</td>
    <td>West Africa and the eastern Mediterranean</td></tr>
    <tr><td><strong>Yaws</strong></td><td><em>T. pallidum</em> subspecies <em>pertenue</em></td><td>Invades a
    <strong>skin cut</strong>, causing a primary ulcer that seeds a <strong>second crop of lesions</strong></td>
    <td>South America, Africa, Asia and Oceania</td></tr>
    <tr><td><strong>Pinta</strong></td><td><em>T. carateum</em></td><td>A <strong>superficial</strong> skin
    lesion that <strong>depigments and scars</strong> the skin</td><td></td></tr>
  </table>

  <h3 class="sub" id="mba-leptospira">13.4 &middot; Objective 1 &mdash; <em>Leptospira</em> and leptospirosis</h3>
  <ul>
    <li><strong>Morphology</strong>: tight, regular individual coils with a <strong>bend or hook at one or
    both ends</strong>.</li>
    <li><em>L. biflexa</em> is a <strong>harmless, free-living saprobe</strong>; <em>L. interrogans</em> causes
    <strong>leptospirosis</strong>, a <strong>zoonosis transmitted mainly via mammal urine</strong>.</li>
    <li>The bacteria are shed in urine and infection occurs mainly by <strong>contact with contaminated
    urine</strong>, so <strong>moisture is required</strong>. Those likely to be infected include
    <strong>surfers, water-skiers, veterinarians, farmers and slaughterhouse workers</strong>.</li>
    <li><strong>Targets</strong>: kidneys, liver, brain and eyes. <strong>Signs</strong>: sudden high fever,
    chills, headache, muscle aches, <strong>conjunctivitis</strong> and vomiting; long-term infection may
    affect the kidneys and liver.</li>
  </ul>

  <h3 class="sub" id="mba-relapsing">13.5 &middot; Objective 1 &mdash; <em>Borrelia</em> and relapsing fever</h3>
  <p><strong><em>Borrelia</em></strong> are the <strong>arthropod-borne spirochetes</strong>: large, with
  <strong>3&ndash;10 irregularly spaced coils</strong>, transmitted by arthropod vectors such as
  <strong>ticks</strong>. <em>B. hermsii</em> causes <strong>relapsing fever</strong>; <em>B.
  burgdorferi</em> causes <strong>Lyme disease</strong> (13.6).</p>
  <table>
    <tr><th>Heading</th><th><em>Borrelia hermsii</em>: relapsing fever</th></tr>
    <tr><td>Reservoirs</td><td>Mammals: <strong>squirrels, chipmunks, wild rodents</strong></td></tr>
    <tr><td>Vector</td><td><strong>Tick-borne</strong></td></tr>
    <tr><td>Disease</td><td>After a 2&ndash;15-day incubation: high fever, shaking chills, headache and fatigue;
    nausea, vomiting, muscle aches, abdominal pain; extensive damage to liver, spleen, heart, kidneys and
    cranial nerves</td></tr>
    <tr><td>Why it relapses</td><td>The organism <strong>changes in the patient&rsquo;s body</strong> while
    the immune system tries to control it, giving <strong>recurrent febrile and afebrile cycles</strong></td></tr>
    <tr><td>Treatment</td><td><strong>Tetracycline</strong>, the most common option</td></tr>
  </table>
  """ + fig("l13-s025-relapsing-fever-pattern.jpg", 800, 514,
      "Graph of body temperature over 24 days: three fever peaks, each labeled an antigenic challenge, separated by troughs at normal temperature labeled first, second and third antibody response; notes explain that each antibody response reduces fever and a new antigenic form causes the next relapse.",
      "<b>The relapsing pattern.</b> Fever with the first infection; the first antibody response brings it "
      "down; a <b>mutant <em>Borrelia</em> with a new antigen</b> causes a relapse; a second antibody response, "
      "and again. Each relapse is a new antigenic form escaping the last antibody.", 25) + """

  <h3 class="sub" id="mba-lyme">13.6 &middot; Objective 1 &mdash; <em>Borrelia burgdorferi</em> and Lyme disease</h3>
  <table>
    <tr><th>Heading</th><th><em>Borrelia burgdorferi</em>: Lyme disease</th></tr>
    <tr><td>Reservoir and vector</td><td><strong>Carried by the white-footed mouse</strong>,
    <strong>transmitted by <em>Ixodes</em> ticks</strong>, in a complex <strong>two-year cycle</strong> involving
    mice and deer</td></tr>
    <tr><td>Course</td><td>A nonfatal, slowly progressive syndrome that <strong>mimics neuromuscular and
    rheumatoid conditions</strong> (it can be mistaken for an autoimmune disorder)</td></tr>
    <tr><td>Early signs</td><td>A <strong>bull&rsquo;s-eye rash</strong> at the tick-bite site
    (<strong>erythema migrans</strong>) in 50&ndash;70%; flu-like fever, headache, stiff neck and
    dizziness</td></tr>
    <tr><td>Untreated</td><td>Can progress to <strong>cardiac and neurological</strong> symptoms and
    <strong>polyarthritis</strong></td></tr>
    <tr><td>Treatment</td><td><strong>Tetracycline or amoxicillin</strong></td></tr>
    <tr><td>Prevention</td><td>Insect repellent containing <strong>DEET</strong> (diethyltoluamide) is the best
    prevention; also, <strong>check for ticks</strong>. A <strong>vaccine for dogs</strong> is available; the
    human vaccine (LYMErix, approved 1998) was <strong>withdrawn from the United States market in
    2002</strong></td></tr>
  </table>
  """ + fig("l13-s028-lyme-cycle.jpg", 1200, 892,
      "A two-year cycle diagram: newly hatched larval ticks become infected feeding on a mouse carrying Borrelia burgdorferi in the first year; in the second year they molt into nymphs, an aggressive feeding stage, that feed on deer and humans, the human being an accidental host; on deer the nymphs mature into adult ticks that mate, and the female lays eggs in plant litter. Inset, ticks beside a pin head.",
      "<b>The two-year cycle.</b> Year 1: larval ticks pick up <em>Borrelia</em> from <b>mice</b>. Year 2: "
      "<b>nymphs</b>, the aggressive feeding stage, bite deer and humans (an <b>accidental host</b>); adults "
      "mate on <b>deer</b> and the eggs start the cycle again.", 28) + """
  """ + fig("l13-s030-erythema-migrans.jpg", 800, 581,
      "A large, round, evenly red patch of skin on the back of an arm near the armpit.",
      "<b>Erythema migrans</b>, the bull&rsquo;s-eye rash at the site of the tick bite.", 30) + """
  <p>The deck&rsquo;s map (slide 29) shows predicted Lyme transmission concentrated in the Northeast and the
  upper Midwest of the United States.</p>

  <h3 class="sub" id="mba-cholera">13.7 &middot; Objective 1 &mdash; The curviform bacteria and <em>Vibrio cholerae</em></h3>
  <table>
    <tr><th>Genus</th><th>Shape</th><th>Flagella</th></tr>
    <tr><td><em>Vibrio</em></td><td><strong>Comma-shaped</strong> rods</td><td>A <strong>single polar</strong>
    flagellum</td></tr>
    <tr><td><em>Campylobacter</em></td><td>Short spirals or curved rods</td><td><strong>One</strong>
    flagellum</td></tr>
    <tr><td><em>Helicobacter</em></td><td>Tight spirals</td><td><strong>Several polar</strong> flagella</td></tr>
  </table>
  <table>
    <tr><th>Heading</th><th><em>Vibrio cholerae</em></th></tr>
    <tr><td>Morphology and antigens</td><td>Comma-shaped, with unique <strong>O and H antigens</strong>; the
    <strong>El Tor</strong> biotype <strong>survives longer and is more infectious</strong> than other
    strains</td></tr>
    <tr><td>Transmission and dose</td><td>Ingested in <strong>contaminated water</strong>. The infectious dose
    is large, because the vibrios <strong>must first pass through the stomach</strong> to reach the
    intestine</td></tr>
    <tr><td>Pathogenesis</td><td>Infects the <strong>mucous barrier of the small intestine</strong> and is
    <strong>noninvasive</strong>. <strong>Cholera toxin</strong> causes loss of electrolytes and water by
    <strong>secretory diarrhea</strong>, the <strong>&ldquo;rice water stool&rdquo;</strong>; the resulting
    dehydration leads to muscle, circulatory and neurological symptoms</td></tr>
    <tr><td>Treatment and prevention</td><td><strong>Oral rehydration</strong> and tetracycline; a
    <strong>vaccine</strong> is available</td></tr>
  </table>
  """ + fig("l13-s033-cholera-water-loss.jpg", 595, 213,
      "Diagram of an intestinal cell between the intestinal lumen and a blood vessel; arrows carry water from the blood vessel through the cell into the lumen, and chloride, bicarbonate, sodium and potassium flow out into the lumen.",
      "<b>What cholera toxin does.</b> Normally water and electrolytes are absorbed from the lumen into the "
      "blood (slide 33, panel a). With the toxin the flow reverses: chloride, bicarbonate, sodium, potassium "
      "and water pour out into the lumen.", 33) + """

  <h3 class="sub" id="mba-seafood">13.8 &middot; Objective 1 &mdash; Vibrios carried by seafood</h3>
  <p>Salt-tolerant inhabitants of <strong>coastal waters</strong> that associate with marine
  invertebrates.</p>
  <ul>
    <li><strong><em>Vibrio parahaemolyticus</em></strong>: gastroenteritis from <strong>raw
    seafood</strong>, with symptoms similar to cholera.</li>
    <li><strong><em>Vibrio vulnificus</em></strong>: gastroenteritis from <strong>raw oysters</strong>, with
    <strong>serious complications in people with diabetes or liver disease</strong>; one of the
    &ldquo;flesh-eating&rdquo; bacteria, and a problem with <strong>sewage overflow after hurricanes</strong>
    and other natural disasters.</li>
    <li><strong>Treatment</strong>: fluid and electrolyte replacement is the best option; antimicrobial drugs
    are used occasionally.</li>
    <li><strong>Which seafood</strong>: peak predators (squid, mackerel, tuna), smaller fish (sardines,
    herring) and filter or bottom feeders (crab, shrimp, oysters, clams) may all contribute to
    <strong>bioaccumulation</strong> of the pathogens.</li>
  </ul>

  <h3 class="sub" id="mba-campylobacter">13.9 &middot; Objective 1 &mdash; <em>Campylobacter</em></h3>
  <p>Gram-negative, slender, curved or spiral rods, often <strong>S-shaped or in &ldquo;gull-winged&rdquo;
  pairs</strong>, with a <strong>single polar flagellum</strong>; common residents of the intestinal tract,
  genitourinary tract and oral cavity of birds and mammals. The two that matter are <em>C. jejuni</em> and
  <em>C. fetus</em>.</p>
  <table>
    <tr><th></th><th><em>Campylobacter jejuni</em> enteritis</th><th><em>Campylobacter fetus</em></th></tr>
    <tr><td>Who</td><td>An important cause of <strong>bacterial gastroenteritis</strong></td><td>An
    <strong>opportunist</strong> in <strong>debilitated persons</strong> or <strong>women late in
    pregnancy</strong></td></tr>
    <tr><td>Transmission</td><td><strong>Beverages and food</strong>: water, milk, meat and chicken are
    common</td><td>To humans, <strong>uncertain</strong>; in veterinary medicine a sexually transmitted
    disease of domestic herd animals (cattle, sheep, goats)</td></tr>
    <tr><td>Disease</td><td>Reaches the mucosa of the <strong>last segment of the small intestine near the
    colon</strong>, adheres to and burrows through the mucus, and multiplies; a <strong>heat-labile
    enterotoxin</strong> stimulates a <strong>secretory diarrhea like that of cholera</strong>; headache,
    fever, abdominal pain, <strong>bloody or watery diarrhea</strong></td><td>Meningitis, pneumonia, arthritis;
    <strong>septicemia in the newborn</strong></td></tr>
    <tr><td>Treatment</td><td><strong>Rehydration and electrolyte balance</strong> (rarely antibiotics)</td>
    <td></td></tr>
  </table>

  <h3 class="sub" id="mba-helicobacter">13.10 &middot; Objective 1 &mdash; <em>Helicobacter pylori</em></h3>
  <ul>
    <li><strong>Gram-negative, motile, curved cells</strong> with several polar flagella, discovered in
    <strong>stomach biopsies from ulcer patients</strong> by the Australian doctors <strong>Warren and
    Marshall</strong> (Nobel Prize in Medicine, 2005).</li>
    <li>Causes most <strong>stomach and duodenal ulcers</strong> and is an apparent <strong>cofactor in
    gastric cancer</strong>.</li>
    <li>Produces <strong>urease</strong>, which converts urea into ammonium and bicarbonate, so the bacterium
    survives the stomach by <strong>raising its local pH</strong>.</li>
    <li>The deck&rsquo;s world map shows reported infection prevalence varying widely between countries.</li>
  </ul>
  """ + fig("l13-s041-helicobacter-urease-test.jpg", 583, 900,
      "Top, a colored scanning micrograph of curved orange Helicobacter cells on purple stomach mucosa; bottom, a small plastic test card with a gel well and a color key reading pH 5.2 negative, pH 5.8 negative, pH 6.5 positive slightly, pH 7.2 positive moderately, pH 7.5 positive markedly, with the instruction to insert the biopsy deep into the gel.",
      "<b>Finding it.</b> (a) <em>Helicobacter</em> on the stomach mucosa. (b) A rapid <b>urease</b> test: a "
      "biopsy is pushed into the gel, and the urease&rsquo;s ammonia raises the pH, turning the gel from "
      "negative (pH 5.2&ndash;5.8) to positive (pH 6.5 and up).", 41) + """

  <h3 class="sub" id="mba-rickettsia">13.11 &middot; Objective 1 &mdash; Rickettsias and <em>Ehrlichia</em></h3>
  <p>The bacteria of <strong>unique morphology and biology</strong> (slide 42): the rickettsias and the
  chlamydias are <strong>Gram-negative coccobacilli</strong>; the mycoplasmas are unique among all bacteria in
  <strong>lacking a cell wall</strong>, and therefore peptidoglycan.</p>
  <table>
    <tr><th>Heading</th><th><em>Rickettsia</em> (family Rickettsiaceae)</th></tr>
    <tr><td>Morphology</td><td><strong>Among the smallest of the bacteria</strong>; Gram-negative cell wall;
    <strong>nonmotile pleomorphic rods or coccobacilli</strong></td></tr>
    <tr><td>Growth</td><td><strong>Obligate intracellular parasites</strong>, so they require <strong>live
    cells for cultivation</strong>; the family has about 23 species, mainly in the genus
    <em>Rickettsia</em>, causing the &ldquo;rickettsioses&rdquo;</td></tr>
    <tr><td>Life cycle</td><td>Spend part of it in <strong>arthropod vectors</strong>: <strong>ticks, fleas and
    lice</strong></td></tr>
    <tr><td>Pathogenesis</td><td>Enter <strong>endothelial cells</strong> and cause necrosis of the vascular
    lining: <strong>vasculitis, vascular leakage and thrombosis</strong></td></tr>
  </table>
  <table>
    <tr><th>Disease</th><th>Agent</th><th>Vector and reservoir</th><th>Features</th></tr>
    <tr><td><strong>Epidemic typhus</strong></td><td><em>R. prowazekii</em></td><td><strong>Lice</strong> (the
    body louse); reservoir humans; louse feces rubbed into the bite, or inhaled</td><td>High fever, chills,
    headache, rash; <strong>Brill-Zinsser disease</strong> is its chronic, recurrent form</td></tr>
    <tr><td><strong>Endemic (murine) typhus</strong></td><td><em>R. typhi</em></td><td><strong>Fleas</strong>;
    harbored by <strong>mice and rats</strong></td><td>Sporadic, in areas of heavy flea infestation;
    <strong>milder</strong></td></tr>
    <tr><td><strong>Rocky Mountain spotted fever</strong></td><td><em>R. rickettsii</em></td><td>A zoonosis
    carried by <strong>dog and wood ticks</strong>; small mammals</td><td>Most cases in the
    <strong>Southeast and the eastern seaboard</strong> of the United States; a distinct <strong>spotted
    rash</strong>; may damage the heart and central nervous system</td></tr>
    <tr><td>Rickettsialpox</td><td><em>R. akari</em></td><td>Mite; mice</td><td>Worldwide</td></tr>
    <tr><td>Scrub typhus</td><td><em>Orientia tsutsugamushi</em></td><td>Immature mite; rodents</td><td>Asia,
    Australia, Pacific Islands</td></tr>
    <tr><td><strong>Human monocytic ehrlichiosis</strong></td><td><em>Ehrlichia chaffeensis</em></td>
    <td><strong>Tick</strong> bite</td><td>The genus <em>Ehrlichia</em> holds tick-borne rickettsias that
    cause human monocytic and granulocytic ehrlichiosis (13.17)</td></tr>
  </table>
  <p><strong>The Weil-Felix test</strong> rests on a <strong>cross-reaction between <em>Proteus</em> bacteria
  and <em>R. prowazekii</em></strong>. During World War II two Polish doctors used it to make the occupying
  Germans, who feared typhus outbreaks, avoid towns in southeastern Poland.</p>
  """ + fig("l13-s049-spotted-fever-cycle.jpg", 800, 507,
      "Diagram: a tick feeding on a dog, a female tick with eggs, an egg, and a tick biting a human; arrows from the human lead to pictures of central nervous system damage and coma, a rash, and vascular damage to the heart.",
      "<b>Rocky Mountain spotted fever.</b> The tick passes <em>Rickettsia rickettsii</em> between dogs and "
      "ticks (and through the tick&rsquo;s eggs); a bite infects a human, and the endothelial damage shows as "
      "rash, vascular damage, and central nervous system damage up to coma.", 49) + """
  """ + fig("l13-s050-spotted-fever-rash.jpg", 800, 558,
      "A child's wrist and hand covered with many small, red, blotchy spots.",
      "<b>The spotted rash.</b> Early on it may be <b>misdiagnosed as measles</b>, a viral disease with very "
      "different treatment.", 50, mh=True) + """

  <h3 class="sub" id="mba-coxiella">13.12 &middot; Objective 1 &mdash; <em>Coxiella</em> and <em>Bartonella</em></h3>
  <table>
    <tr><th></th><th><em>Coxiella burnetii</em></th><th><em>Bartonella</em> species</th></tr>
    <tr><td>Organism</td><td>An <strong>intracellular parasite</strong> that produces an <strong>unusually
    resistant endospore</strong> (13.17)</td><td>Small, <strong>Gram-negative, fastidious</strong>;
    <strong>cultured on blood agar</strong></td></tr>
    <tr><td>Disease</td><td><strong>Q fever</strong> (Q for &ldquo;query&rdquo;, unknown; historic cases most
    likely arose from contaminated milking barns): usually <strong>inhaled</strong>, causing pneumonitis,
    fever and inflammation of the liver</td><td><strong>Trench fever</strong>, spread by
    <strong>lice</strong>; <strong>cat-scratch disease</strong>, a lymphatic infection after a clawing injury by
    a cat; <strong>bacillary angiomatosis</strong> in patients with acquired immunodeficiency syndrome</td></tr>
    <tr><td>Reservoir and transmission</td><td>Harbored by a wide range of vertebrates and arthropods;
    infectious material includes <strong>urine, feces, milk and airborne particles</strong></td><td>Lice (trench
    fever); cats (cat-scratch disease)</td></tr>
    <tr><td>Treatment</td><td><strong>Tetracycline</strong></td><td><strong>Tetracycline, erythromycin,
    rifampin</strong></td></tr>
    <tr><td>Prevention</td><td>A <strong>vaccine</strong> is available for those at highest risk:
    <strong>veterinarians, farmers, laboratory technicians</strong></td><td></td></tr>
  </table>

  <h3 class="sub" id="mba-chlamydia">13.13 &middot; Objective 1 &mdash; <em>Chlamydia</em> and <em>Chlamydophila</em></h3>
  <p>The <strong>Chlamydiaceae</strong> are <strong>obligate intracellular parasites</strong>, small cells with a
  Gram-negative cell wall, that <strong>alternate between two stages</strong>:</p>
  <table>
    <tr><th>Elementary body</th><th>Reticulate body</th></tr>
    <tr><td>Small, <strong>metabolically inactive</strong>, <strong>extracellular</strong> and
    <strong>infectious</strong>; the form released by the infected host cell</td><td><strong>Noninfectious</strong>
    and <strong>actively dividing</strong>; grows inside host-cell vacuoles, but <strong>cannot make its own
    adenosine triphosphate</strong></td></tr>
  </table>
  """ + fig("l13-s057-chlamydia-life-cycle.jpg", 586, 900,
      "Top, a micrograph of chlamydias among the microvilli of a host cell; below, a cycle diagram: elementary bodies enter a host cell into phagosomes, become reticulate bodies that divide by binary fission, reorganize into elementary bodies, and are released to infect a new host cell.",
      "<b>The life cycle.</b> Elementary bodies (EB) attach and are taken into a phagosome, turn into "
      "reticulate bodies (RB) that divide by binary fission, become elementary bodies again, and are "
      "released to infect new cells.", 57) + """
  <table>
    <tr><th>Organism</th><th>Disease</th></tr>
    <tr><td><em>Chlamydia trachomatis</em> (human reservoir; two strains, trachoma and lymphogranuloma
    venereum)</td><td>The trachoma strain attacks the mucous membranes of the <strong>eyes, genitourinary tract
    and lungs</strong>. <strong>Ocular trachoma</strong>: severe, deforms the eyelid and cornea, may cause
    <strong>blindness</strong>; it affects people in tropical climates with poor sanitation, flying insects and
    poor hygiene. <strong>Inclusion conjunctivitis</strong>: acquired as the baby passes through the birth
    canal; prevented by prophylaxis. <strong>Sexually transmitted infection</strong>: the first or second most
    prevalent in the United States in recent years; urethritis, cervicitis, salpingitis (pelvic inflammatory
    disease), infertility, scarring. <strong>Lymphogranuloma venereum</strong>: a disfiguring disease of the
    external genitalia and pelvic lymphatics. Diagnosis can show an <strong>inclusion body</strong> next to
    the nucleus of an infected cell (slide 60)</td></tr>
    <tr><td><em>Chlamydophila</em>, a new genus of former <em>Chlamydia</em> members</td><td><em>C.
    pneumoniae</em>: an <strong>atypical pneumonia</strong>, serious in <strong>asthma</strong> patients. <em>C.
    psittaci</em>: <strong>ornithosis</strong>, a zoonosis transmitted to humans from <strong>birds</strong>
    (&ldquo;ornith&rdquo; = bird), highly communicable among birds; pneumonia or a flu-like infection with fever
    and lung congestion</td></tr>
  </table>

  <h3 class="sub" id="mba-mycoplasma">13.14 &middot; Objective 1 &mdash; Mycoplasmas</h3>
  <ul>
    <li>The <strong>mollicutes</strong> and other cell-wall-deficient bacteria: they <strong>naturally lack
    cell walls</strong> and are <strong>highly pleomorphic</strong> (cell shapes vary widely).</li>
    <li>They <strong>require special lipids from host membranes</strong>.</li>
    <li><strong>Treated with tetracycline or erythromycin</strong>. Not penicillin: <strong>no peptidoglycan is
    present</strong> for it to act on.</li>
    <li><strong><em>Mycoplasma pneumoniae</em></strong>: <strong>primary atypical pneumonia</strong>; the
    pathogen spreads slowly over the interior respiratory surfaces, causing fever, chest pain and sore
    throat.</li>
    <li><em>M. genitalium</em> and <em>Ureaplasma urealyticum</em>: <strong>weak sexually transmitted</strong>
    pathogens.</li>
  </ul>

  <h3 class="sub" id="mba-dental">13.15 &middot; Objective 1 &mdash; Bacteria in dental disease</h3>
  <ul>
    <li>The oral cavity is a complex, dynamic ecosystem of <strong>over 400 microbial species</strong>.</li>
    <li><strong>Dental caries</strong> (cavities) is a slow, progressive infection of irregular areas of the
    enamel surface, and the <strong>most commonly diagnosed of all human diseases</strong>.</li>
    <li>It begins with colonization by <strong>slime-forming <em>Streptococcus</em></strong> species, with
    cross-adherence of <em>Actinomyces</em> in the biofilm; <strong><em>Streptococcus mutans</em></strong> is
    commonly found. This forms <strong>plaque</strong>, a thick, adherent layer harboring masses of bacteria
    whose <strong>acid dissolves enamel</strong>; mature biofilms are polymicrobial.</li>
    <li>Plaque left on and between teeth brings <strong>secondary invaders</strong> such as
    <em>Lactobacillus</em>, <em>Bacteroides</em> and <em>Fusobacterium</em>.</li>
    <li><strong>Periodontal disease</strong>, a soft-tissue disease: plaque <strong>calcified into
    calculus</strong> above and below the gingiva irritates it, causing <strong>gingivitis</strong>; the
    pockets between tooth and gingiva are invaded by <strong>spirochetes and Gram-negative bacilli</strong>;
    the tooth socket may become involved (<strong>periodontitis</strong>) and the tooth may be lost.</li>
  </ul>
  """ + fig("l13-s066-dental-disease-flow.jpg", 514, 600,
      "Flow chart: tooth surface, then pellicle formation, then plaque formation, which branches into acidification and cavitation leading to dental caries, damage to enamel and tooth destruction, and into calculus formation leading to gingivitis and periodontal disease and bone resorption; both branches end in tooth loss.",
      "<b>Two roads from plaque to tooth loss.</b> Acid gives <b>caries</b> and tooth destruction; "
      "calculus gives <b>gingivitis, periodontal disease</b> and bone resorption.", 66) + """

  <h3 class="sub" id="mba-transmission">13.16 &middot; Objective 2 &mdash; Modes of transmission</h3>
  <table>
    <tr><th>Organism</th><th>How humans acquire it</th></tr>
    <tr><td><em>Treponema pallidum</em> (syphilis)</td><td><strong>Sexual</strong> contact and
    <strong>transplacental</strong> (congenital syphilis); humans are the only natural host</td></tr>
    <tr><td>Bejel, yaws, pinta</td><td><strong>Rarely</strong> sexual or congenital; endemic in particular
    regions; yaws invades a skin cut</td></tr>
    <tr><td><em>Leptospira interrogans</em></td><td><strong>Contact with mammal urine</strong>, in moist
    settings (surfers, water-skiers, veterinarians, farmers, slaughterhouse workers)</td></tr>
    <tr><td><em>Borrelia hermsii</em></td><td><strong>Tick</strong> bite; rodent reservoirs (squirrels,
    chipmunks)</td></tr>
    <tr><td><em>Borrelia burgdorferi</em></td><td><strong><em>Ixodes</em> tick</strong> bite; white-footed mouse
    reservoir, deer in the cycle, humans accidental hosts</td></tr>
    <tr><td><em>Vibrio cholerae</em></td><td><strong>Contaminated water</strong>, ingested</td></tr>
    <tr><td><em>Vibrio parahaemolyticus</em>, <em>V. vulnificus</em></td><td><strong>Raw seafood</strong>;
    raw oysters</td></tr>
    <tr><td><em>Campylobacter jejuni</em></td><td><strong>Food and beverages</strong>: water, milk, meat,
    chicken</td></tr>
    <tr><td><em>Campylobacter fetus</em></td><td>Uncertain in humans; sexually transmitted among herd
    animals</td></tr>
    <tr><td><em>Rickettsia prowazekii</em></td><td><strong>Lice</strong></td></tr>
    <tr><td><em>Rickettsia typhi</em></td><td><strong>Fleas</strong> (mice and rats)</td></tr>
    <tr><td><em>Rickettsia rickettsii</em></td><td><strong>Dog and wood ticks</strong></td></tr>
    <tr><td><em>Ehrlichia</em></td><td><strong>Ticks</strong></td></tr>
    <tr><td><em>Coxiella burnetii</em></td><td>Usually <strong>inhaled</strong> (urine, feces, milk, airborne
    particles from animals)</td></tr>
    <tr><td><em>Bartonella</em></td><td><strong>Lice</strong> (trench fever); a <strong>cat&rsquo;s
    scratch</strong> (cat-scratch disease)</td></tr>
    <tr><td><em>Chlamydia trachomatis</em></td><td><strong>Sexual</strong> contact; <strong>passage through the
    birth canal</strong> (inclusion conjunctivitis)</td></tr>
    <tr><td><em>Chlamydophila psittaci</em></td><td><strong>Birds</strong></td></tr>
    <tr><td><em>Mycoplasma genitalium</em>, <em>Ureaplasma urealyticum</em></td><td><strong>Sexual</strong>
    contact</td></tr>
  </table>

  <h3 class="sub" id="mba-notes">13.17 &middot; Where the slides and current practice differ</h3>
  <p>The quizzes key what the slides say. Where a slide differs from current usage, or contradicts another
  slide, it is noted here; <strong>points the deck contradicts itself on are not asked</strong>.</p>
  <ul>
    <li><strong><em>Helicobacter</em> as a &ldquo;spirochete&rdquo; (slide 31).</strong> Slide 39 calls it
    &ldquo;curved cells&rdquo;, and a spirochete swims with endoflagella (slide 3) while <em>Helicobacter</em>
    has several polar flagella outside the cell; it is a curved, spiral Gram-negative rod. Not asked.</li>
    <li><strong>Spirochetes &ldquo;human pathogens&rdquo; and &ldquo;not primary pathogens&rdquo;
    (slide 3)</strong> sit side by side. Not asked.</li>
    <li><strong>&ldquo;Rickettsioses identified only in the last 20&ndash;30 years&rdquo; (slide 43)</strong>
    conflicts with the deck&rsquo;s own wartime typhus history (slide 45). Not asked.</li>
    <li><strong>Human granulocytic ehrlichiosis (slides 46&ndash;47)</strong> is now called <strong>human
    granulocytic anaplasmosis</strong>, caused by <em>Anaplasma phagocytophilum</em>, also tick-borne. The
    syllabus lists <em>Anaplasma</em>; the deck never names it. Only monocytic ehrlichiosis is keyed.</li>
    <li><strong>The human Lyme vaccine.</strong> Slide 26 says it was discontinued &ldquo;due to adverse
    effects&rdquo;; slide 27 says researchers found no connection between the vaccine and the complaints. That
    it was withdrawn, and that a dog vaccine exists, are asked; the reason is not.</li>
    <li><strong><em>Coxiella</em>&rsquo;s &ldquo;endospore&rdquo; (slide 52).</strong> Current texts describe a
    resistant, spore-like small-cell form rather than a true endospore like those of <em>Bacillus</em> and
    <em>Clostridium</em>.</li>
    <li><strong>Spellings</strong>: the test is the rapid plasma <em>reagin</em> (slide 16 writes
    &ldquo;reagent&rdquo;); <em>Ureaplasma</em> (slide 62, &ldquo;Ureplasma&rdquo;); bacillary
    angiomatosis (slide 53, &ldquo;angiomatosus&rdquo;); <em>Coxiella burnetii</em> (slides 51&ndash;52,
    &ldquo;burnetti&rdquo;).</li>
    <li><strong>Percentages and counts</strong> (30% reaching tertiary syphilis, 50&ndash;70% with the Lyme
    rash, ulcer percentages, blood-type ratios, case counts, infectious-dose numbers) are context; no question
    keys one.</li>
  </ul>
"""

SECTION_FOOT = """
  <button type="button" class="test-yourself-btn" onclick="window.openTestYourself('Test yourself &mdash; Miscellaneous Bacterial Agents', TEST_YOURSELF.miscBacteria)">Test yourself! &rarr;</button>
  <p class="src">Source: <em>Miscellaneous bacterial agents of infection</em> (PAJ5200.Misc inf agents-2.pptx),
  Slides 1&ndash;70, and the PAJ 5200 syllabus instructional objectives and Appendix 2.</p>
</section>
"""

TEST = """
    miscBacteria: [
      {q:"How do spirochetes move?",
       choices:["Endoflagella inside the cell","A single external polar flagellum","Gliding on slime","They are nonmotile"],correct:0,
       explain:"Spirochetes swim with endoflagella that lie inside the cell, in the periplasmic space. Vibrio has one external polar flagellum."},
      {q:"Which stage of syphilis brings a rash on the palms and soles?",
       choices:["Primary","Tertiary","Secondary","Latent"],correct:2,
       explain:"Secondary syphilis brings fever, sore throat and a red or brown rash including the palms and soles. Primary is the hard chancre; tertiary brings gummas and neurological and cardiovascular disease."},
      {q:"Which tick transmits Borrelia burgdorferi?",
       choices:["Dog tick","Ixodes tick","Wood tick","Soft tick of rodents"],correct:1,
       explain:"Lyme disease is carried by the white-footed mouse and transmitted by Ixodes ticks. Dog and wood ticks carry Rocky Mountain spotted fever."},
      {q:"Why is penicillin useless against Mycoplasma?",
       choices:["It is destroyed by urease","Mycoplasma lives inside cells","Mycoplasma makes penicillinase","Mycoplasma has no peptidoglycan"],correct:3,
       explain:"Mycoplasmas naturally lack cell walls, so there is no peptidoglycan for penicillin to act on; tetracycline or erythromycin is used."},
      {q:"Which chlamydial form is infectious?",
       choices:["The elementary body","The reticulate body","The inclusion body","The endospore"],correct:0,
       explain:"The elementary body is the small, metabolically inactive, extracellular infectious form; the reticulate body divides inside the host cell and is noninfectious."}
    ],
"""


def section():
    return SECTION_HEAD + BODY + SECTION_FOOT


SECTION = None   # set by build(); build_micro_e2_guide.py reads SECTION after import


def build():
    global SECTION
    SECTION = section()
    return SECTION


build()
