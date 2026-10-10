# -*- coding: utf-8 -*-
"""Section 11 of the Microbiology Exam 2 guide -- Novel Antimicrobial Therapy (Dr. Fair).

Same shape as _micro_e2_guide_l10.py (FIG, DECK, fig(), TOC, SECTION, TEST, IMAGES), plus extract():
four of this deck's pictures sit in content PLACEHOLDERS, which build_micro_e2_guide.py's picture
extractor does not see, so this module writes its own figures and the builder only checks them.

Attribution: Dr. Fair. The title slide names no lecturer (her PAJ5200.* decks never do); slide 30
credits "Photos: D. Fair", the calendar row for 9 October 2026 names Dr. Fair, and the recording is hers.

Instructional objectives are VERBATIM from the syllabus (Micro.pdf): the four numbered objectives under
"Novel Antimicrobial Therapy", answered in order (11.1-11.3 objective 1, 11.4 objective 2, 11.5 objective
3, 11.6 objective 4). The deck teaches phage therapy before fecal transplant; the guide follows the
objectives.

AUDIO: the 9 October 2026 recording (45:42, one clip) read in BOTH transcripts (local faster-whisper
medium.en and Notability's) and diffed; quotes appear in both unless marked, cited by the local
transcript's timestamps.
"""
import os

FIG = "micro-exam-2-study-guide-images"
DECK = "PAJ5200.Novel Antimicrobial Therapies-2.pptx"
INBOX_DECK = os.path.expanduser(
    "~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 2/PAJ5200.Novel Antimicrobial Therapies-2.pptx")

# (slide, picture index on that slide counting placeholder pictures, published file name)
IMAGES = [
    (4, 1, "l11-s04-stewardship-cycle.jpg"),
    (5, 1, "l11-s05-six-antibiotic-facts.png"),
    (9, 1, "l11-s09-disk-diffusion.jpg"),
    (18, 1, "l11-s18-phage-structure.png"),
]


def extract(imgdir):
    """Write IMAGES into imgdir from the deck (pictures and picture placeholders alike)."""
    from pptx import Presentation
    prs = Presentation(INBOX_DECK)
    os.makedirs(imgdir, exist_ok=True)
    for slide, k, name in IMAGES:
        pics = []
        for sh in prs.slides[slide - 1].shapes:
            try:
                sh.image
            except Exception:
                continue
            pics.append(sh)
        assert len(pics) >= k, "%s slide %d has %d picture(s)" % (DECK, slide, len(pics))
        img = pics[k - 1].image
        assert name.endswith("." + img.ext.replace("jpeg", "jpg")), (name, img.ext)
        open(os.path.join(imgdir, name), "wb").write(img.blob)
    return len(IMAGES)


def fig(name, w, h, alt, cap, slide):
    return ('<figure class="fig"><img width="%d" height="%d" loading="lazy" src="%s/%s" alt="%s">'
            '<figcaption>%s <span class="src">Source: %s, Slide %d.</span></figcaption></figure>'
            % (w, h, FIG, name, alt, cap, DECK, slide))


TOC = """
  <a class="top-link" href="#novel-antimicrobial">11 &middot; Novel Antimicrobial Therapy</a>
  <a class="sub-link" href="#nat-audio">&#9733; What the recording adds</a>
  <a class="sub-link" href="#nat-why">11.1 Objective 1 &mdash; Why novel therapies: resistance and stewardship</a>
  <a class="sub-link" href="#nat-choosing">11.2 Objective 1 &mdash; How antibiotics are chosen: testing, antibiograms, empiric therapy</a>
  <a class="sub-link" href="#nat-prophylaxis">11.3 Objective 1 &mdash; Prophylaxis and the costs of antimicrobials</a>
  <a class="sub-link" href="#nat-fmt">11.4 Objective 2 &mdash; Fecal microbiota transplant</a>
  <a class="sub-link" href="#nat-phage">11.5 Objective 3 &mdash; Phage therapy</a>
  <a class="sub-link" href="#nat-other">11.6 Objective 4 &mdash; Other alternatives to antibiotics</a>
  <a class="sub-link" href="#nat-truth">11.7 Where the slide, the recording and current practice differ</a>
"""

AUDIO = """
<!--MICROL11AUDIO-->
  <div class="prof-flag" id="nat-audio"><span class="prof-flag-label">&#9733; From the lecture recording &mdash; 9 October 2026</span>
  <p>About 46 minutes with Dr. Fair, one recording that begins at slide 4 (the stewardship cycle), so the title,
  objectives and &ldquo;what is driving the need&rdquo; slides (1&ndash;3) have no audio. <b>Both transcripts read
  and compared</b> (Notability&rsquo;s own and an independent local transcription); every quote below appears in
  both unless marked. She signposts little, but two remarks cut the scope.</p>
  <table>
    <tr><th>What was said</th><th>What it means for you</th></tr>
    <tr><td>On penicillin&rsquo;s history: <em>&ldquo;<mark class="prof-highlight">Don&rsquo;t worry about
    dates</mark>.&rdquo;</em> [24:04]</td>
    <td>No question asks for a year: not slide 8&rsquo;s discovery, use and resistance dates, not the 1920s and 1930s
    of phage therapy, not the centuries of fecal therapy. Know the patterns instead.</td></tr>
    <tr><td>On how vancomycin works: <em>&ldquo;<mark class="prof-highlight">don&rsquo;t worry about this</mark>, I
    geeked out at lunchtime.&rdquo;</em> [08:18]</td>
    <td>Vancomycin&rsquo;s mechanism is not examined here. What she did say: it still works against Gram-positive
    infections but not Gram-negative ones, which it cannot get through the lipid-rich envelope of.</td></tr>
    <tr><td><em>&ldquo;They&rsquo;re very, very specific as to the bacterium that they&rsquo;re going to attack, and
    then <mark class="prof-highlight">once there are no more of that particular strain of bacteria, they die
    off</mark>.&rdquo;</em> [22:16]</td>
    <td>The two properties of phages to carry into any phage question: high host specificity, and dying off
    once the host is gone (11.5).</td></tr>
    <tr><td><em>&ldquo;And <mark class="prof-highlight">remember that the diagnosis might be time consuming</mark>
    too, especially if the bacterium causing the infection is a slow grower, or it&rsquo;s very, very
    fastidious.&rdquo;</em> [25:02]</td>
    <td>Phage therapy needs the causative organism identified first; for slow growers such as
    <em>Mycobacterium</em> or <em>Legionella</em> that may need immunological or DNA techniques (11.5).</td></tr>
    <tr><td>On fecal microbiota transplant: <em>&ldquo;it does seem to have a great success rate.&rdquo;</em>
    [30:28] The study she teaches from (the one highlighted on slide 27) gave it to patients with <b>recurring
    <em>Clostridioides difficile</em></b>, 14 by nasogastric tube and 14 by colonoscopy. [33:42&ndash;34:14]</td>
    <td>Recurrent <em>Clostridioides difficile</em> is the indication to know (11.4). The success figures she gave
    are context, not exam material.</td></tr>
  </table>
  <p><b>Things said that are not on a slide</b> &mdash; consistent with the deck, useful as hooks:</p>
  <ul>
    <li><b>Reading the disk-diffusion plates</b> (slide 9): penicillin does nothing to the <em>Escherichia coli</em>
    on the left, a Gram-negative, but works on that <em>Staphylococcus aureus</em> strain on the right. Leave a
    plate in the incubator a few days longer and colonies appear inside the clear zones: bacteria
    <b>evolving resistance on the plate</b> at a sub-lethal concentration of the drug. [06:15&ndash;07:14]</li>
    <li><b>Penicillin was found by accident</b>: a fungus landed on a plate of staphylococci and its product killed
    them, and even its discoverer warned in his Nobel lecture that antibiotics are not a magic bullet.
    [03:56&ndash;04:48]</li>
    <li><b>Animal bites</b>: <em>Capnocytophaga</em>, &ldquo;the dog bite bacterium&rdquo;, is her example of why a
    deep bite gets prophylaxis. [12:54]</li>
    <li><b>Phages</b> were discovered around 1915&ndash;1917 by two scientists working independently, who saw
    something clearing zones in their bacterial colonies; &ldquo;phage&rdquo; means to eat. [20:34&ndash;21:05]</li>
    <li><b>A fourth route for fecal transplant</b>: frozen oral capsules, besides nasogastric tube and
    colonoscopy. [32:19&ndash;32:43] (Oral capsule products exist; slide 26 lists tube, colonoscopy and enema.)</li>
    <li><b>Medical tourism</b> is one way patients reach phage therapy that is not approved where they live; in the
    United States the usual routes are investigational drug status and clinical studies. [26:07&ndash;27:20]</li>
    <li><b>Parasitic worms</b> (slide 28, about six minutes): in populations with heavy intestinal worm burdens,
    studies found fewer allergies, less diabetes and fewer other chronic inflammatory conditions, and one recent
    study also found more gut inflammation-linked cancers; whether the worms or their products are responsible
    is unknown. It is the hygiene hypothesis of Lecture 7. [34:45&ndash;40:06] Not quizzed: the slide&rsquo;s
    text states only the plant-diversity finding.</li>
    <li><b>Plant essential oils</b> have been studied as antibacterial, antifungal, antiprotozoal, antiviral and
    anti-inflammatory agents. [44:39]</li>
    <li>More antibiotics go into <b>animal feed</b> than to anything else. [45:31]</li>
  </ul>
  <p class="muted"><b>Slips to ignore.</b> Introducing slide 15 she said &ldquo;the non-steroidal
  anti-inflammatories&rdquo; [14:46]; the slide says <b>systemic corticosteroids</b>, which are steroids. She
  introduced empiric therapy as &ldquo;prescribing antibiotics before anything happens&rdquo; with the old dental
  example [09:00&ndash;09:30]; that is prophylaxis. Empiric therapy, as slide 11 defines it, <b>treats a suspected
  infection before the culture results are back</b>. Of the phage gene-transfer worry she asked &ldquo;what if the
  bacteriophage DNA is then incorporated into human cells?&rdquo; [25:32]; phages infect only bacteria, and the
  slide&rsquo;s &ldquo;host cell&rdquo; is the bacterium (11.7). Her ESKAPE letters began &ldquo;E for
  <em>E. coli</em>&rdquo; [19:18]; the usual first E is <em>Enterococcus faecium</em>. And penicillin&rsquo;s
  discovery was in London (St Mary&rsquo;s Hospital), not Edinburgh [03:27].</p>
  </div>
<!--/MICROL11AUDIO-->
"""

SECTION_HEAD = """
<section class="deck" id="novel-antimicrobial">
  <h2 class="deck-title">11 &middot; Novel Antimicrobial Therapy</h2>
  <div class="io-box">
    <h3>Instructional Objectives</h3>
    <ol>
      <li>Compare and contrast current antibiotic use related to novel antimicrobial therapy.</li>
      <li>Describe fecal microbiota transplant (FMT).</li>
      <li>Describe phage therapy.</li>
      <li>Describe other alternative therapies to antibiotic use.</li>
    </ol>
  </div>

  <div class="callout"><strong>One question runs through the whole lecture:</strong> <em>are we nearing a
  post-antibiotic era?</em> The answer on the title slide is &ldquo;yes, probably, but we do have
  options.&rdquo; Objective 1 is the problem and how antibiotics are used today: resistance and what drives
  it, stewardship, susceptibility testing, antibiograms and empiric therapy, the special circumstances for
  prophylaxis, and the direct and indirect costs of every course. Objectives 2&ndash;4 are the options:
  <strong>fecal microbiota transplant</strong> (restore the gut&rsquo;s own flora), <strong>phage
  therapy</strong> (use the viruses that kill bacteria) and the <strong>other alternatives</strong>
  (corticosteroids, donor antibodies, prebiotics and probiotics, a varied plant diet, plant essential
  oils).</div>
"""

BODY = """
  <h3 class="sub" id="nat-why">11.1 &middot; Objective 1 &mdash; Why novel therapies: resistance and stewardship</h3>
  <p><strong>What drives the need for novel therapies</strong> is antimicrobial and antibiotic
  <strong>resistance</strong>, fed by:</p>
  <ul>
    <li><strong>Over-use of antibiotics in agriculture</strong>, in plants and in animals;</li>
    <li><strong>clinically inappropriate prescriptions</strong>, which create the need for
    <strong>antimicrobial stewardship</strong>: education of clinicians and of patients;</li>
    <li><strong>antibiotic prophylaxis</strong>, used to reduce infection risk &ldquo;when possible&rdquo;
    (11.3 lists the circumstances where it is justified);</li>
    <li><strong>patient non-compliance</strong>, such as stopping a course early or leaving against
    medical advice.</li>
  </ul>
  <p><strong>Antibiotic resistance</strong> is a type of drug resistance that renders the antibiotic
  ineffective in killing or controlling bacterial growth (slide 6). It belongs to the bacteria, not to
  the patient: an allergy is the patient reacting to the drug; resistance is the drug no longer working
  on the organism.</p>
  """ + fig("l11-s04-stewardship-cycle.jpg", 673, 468,
      "A four-part cycle around the words cycle of bacterial disease management: diagnosis of bacterial disease; consideration of a non-antibiotic alternative to prevent, control or treat; judicious use of antimicrobial drugs when they are needed; re-evaluation of the need for antimicrobial use.",
      "<b>Antimicrobial stewardship as a cycle.</b> Diagnose a <em>bacterial</em> disease first; then ask "
      "whether a <b>non-antibiotic alternative</b> could prevent, control or treat it; use antimicrobials "
      "<b>judiciously, when they are needed</b>; and <b>re-evaluate the need</b> for them, which brings you "
      "back to the start.", 4) + """
  """ + fig("l11-s05-six-antibiotic-facts.png", 1024, 512,
      "Six numbered icons: antibiotics are life-saving drugs; antibiotics only treat bacterial infections; some ear infections do not require an antibiotic; most sore throats do not require an antibiotic; green colored mucus is not a sign that an antibiotic is needed; there are potential risks when taking any prescription drug, shown as allergy, diarrhea and rash.",
      "<b>Six facts about antibiotic use</b> to give patients: antibiotics are <b>life-saving</b>; they "
      "treat <b>only bacterial</b> infections; <b>some ear infections</b> and <b>most sore throats</b> do not "
      "need one; <b>green mucus is not</b> a sign that one is needed; and <b>every prescription drug carries "
      "risks</b> (allergy, rash, diarrhea).", 5) + """
  <p><strong>Resistance follows use.</strong> There are over 100 antibiotics on the market, and resistance is
  still becoming more prevalent. Slide 8 tracks three of them from discovery to widespread clinical use to
  the first resistance detected: penicillin, vancomycin and azithromycin each met resistant bacteria within
  a few years to about two decades of coming into wide use. (The exact years differ between sources; the
  pattern is the point.)</p>
  <p><strong>The &ldquo;superbugs&rdquo;</strong> (slide 16): methicillin-resistant and methicillin-susceptible
  <em>Staphylococcus aureus</em> (MRSA, MSSA), vancomycin-resistant and vancomycin-intermediate
  <em>Staphylococcus aureus</em> (VRSA, VISA), toxigenic <em>Escherichia coli</em>, <em>Clostridioides
  difficile</em>, <em>Neisseria gonorrhoeae</em>, <em>Mycobacterium tuberculosis</em>, carbapenem-resistant
  <em>Klebsiella pneumoniae</em>, vancomycin-resistant <em>Enterococcus</em> and <em>Pseudomonas</em>. The
  slide&rsquo;s notes add the <strong>ESKAPE pathogens</strong>, which top the lists of organisms causing
  nosocomial (hospital-acquired) infections.</p>

  <h3 class="sub" id="nat-choosing">11.2 &middot; Objective 1 &mdash; How antibiotics are chosen: testing, antibiograms, empiric therapy</h3>
  <p><strong>Disk diffusion (Kirby-Bauer).</strong> White paper disks treated with antibiotics are placed on
  an agar plate inoculated with the organism. The antibiotic diffuses out, and the circular <strong>zone of
  inhibition</strong> around each disk shows how effective, or how ineffective, that antimicrobial is
  against the organism: a wide clear ring means growth was stopped; growth up to the disk means the drug
  had little effect.</p>
  """ + fig("l11-s09-disk-diffusion.jpg", 403, 302,
      "Two agar plates on a dark surface above a ruler. The left plate shows bacterial growth close around its white antibiotic disks; the right plate shows wide clear rings around several of its disks.",
      "<b>Zones of inhibition.</b> <em>Escherichia coli</em> (left) and <em>Staphylococcus aureus</em> "
      "(right) on disk-diffusion plates: the clear rings on the right are where the antibiotics stopped "
      "growth; on the left the bacteria grow nearly up to the disks.", 9) + """
  <p><strong>Antibiograms</strong> are profiles of antimicrobial susceptibility, by percentage, used to guide
  <strong>empiric therapy</strong>. On the slide&rsquo;s example, methicillin-resistant <em>Staphylococcus
  aureus</em> is far less often susceptible than methicillin-susceptible strains to ciprofloxacin and
  erythromycin, while both are fully susceptible to vancomycin. The limitations:</p>
  <ul>
    <li>the <strong>minimum inhibitory concentration is not included</strong>;</li>
    <li>the <strong>patient&rsquo;s risk factors are not considered</strong>;</li>
    <li>there are <strong>no data on synergy</strong> when antimicrobials are used in combination;</li>
    <li>the data <strong>may apply only to a single hospital</strong> within a health system.</li>
  </ul>
  <p><strong>Empiric therapy</strong> means antibiotics prescribed <strong>before</strong> cultures are
  confirmed or other diagnostic results are available. It is common in human and veterinary medicine,
  especially for <strong>critically ill or emergent</strong> patients, is based on <strong>prior clinical
  experience and/or antibiograms</strong>, and is then <strong>adjusted</strong> once results arrive (for
  example, from intravenous to oral). &ldquo;Why empirical therapy? Because it works!&rdquo; (Klastersky,
  2009).</p>

  <h3 class="sub" id="nat-prophylaxis">11.3 &middot; Objective 1 &mdash; Prophylaxis and the costs of antimicrobials</h3>
  <p><strong>Special circumstances</strong> where antimicrobials are given to prevent rather than treat
  (slides 12&ndash;13):</p>
  <table>
    <tr><th>Circumstance</th><th>Why, and the examples</th></tr>
    <tr><td>Foreign body-associated infections</td><td>Prosthetic implants and devices are prone to microbial
    <strong>biofilm</strong> formation.</td></tr>
    <tr><td>Presurgical prophylaxis</td><td>Reduces the incidence of <strong>post-operative surgical site
    infections</strong>.</td></tr>
    <tr><td>Immunocompromised patients</td><td>Human immunodeficiency virus infection and acquired
    immunodeficiency syndrome, cancer chemotherapy, immunosuppressive therapy for organ transplant.</td></tr>
    <tr><td>Contacts of an infected patient</td><td>Prevents transmission to susceptible contacts, for example
    <strong>tuberculosis, meningitis or pertussis</strong>.</td></tr>
    <tr><td>Dental and other invasive procedures</td><td>Many dentists now limit it to a <strong>few high-risk
    scenarios</strong>, such as <strong>prosthetic heart valves or a history of endocarditis</strong>.</td></tr>
    <tr><td>High-risk traumatic injuries</td><td>For example <strong>animal bites</strong> or penetrating
    brain injury.</td></tr>
  </table>
  <p><strong>Adverse effects of antimicrobial drugs</strong> come in two kinds (slide 14):</p>
  <table>
    <tr><th>Direct (on the patient, from the drug)</th><th>Indirect (through other microbes)</th></tr>
    <tr><td>Allergy, possibly life-threatening; acute or chronic <strong>toxicity</strong> (such as kidney or
    liver toxicity); <strong>drug-drug interactions</strong>; <strong>therapeutic failure</strong>.</td>
    <td>Effects on the <strong>commensal microflora</strong> of humans or animals; <strong><em>Clostridioides
    difficile</em> infection</strong>; a greater chance of infection with <strong>drug-resistant
    pathogens</strong>; effects on <strong>environmental microflora</strong>, such as antibiotics in the food
    or water supply.</td></tr>
  </table>
  <p>The indirect column is the bridge to objectives 2&ndash;4: <em>Clostridioides difficile</em> infection
  is what happens when an antibiotic clears the commensal flora, and fecal microbiota transplant is the
  therapy that puts that flora back.</p>

  <h3 class="sub" id="nat-fmt">11.4 &middot; Objective 2 &mdash; Fecal microbiota transplant</h3>
  <p><strong>Fecal microbiota transplant</strong> gives a patient screened donor stool to realign their
  gastrointestinal microflora, &ldquo;the power of poop&rdquo;.</p>
  <ul>
    <li><strong>It is old.</strong> It has been used for centuries in human and animal medicine; old Chinese
    texts give recipes for &ldquo;yellow soup&rdquo; (Ge Hong in the 4th century and Li Shizhen in the 16th
    are the sources usually cited).</li>
    <li><strong>Its researched use is <em>Clostridioides difficile</em> infection</strong>, about 500,000 cases a
    year and about 30,000 deaths in the United States alone. Preventing <strong>recurrent</strong>
    <em>Clostridioides difficile</em> infection is its established use (see 11.7 on the slide&rsquo;s other
    conditions).</li>
    <li><strong>How it is given:</strong> by nasogastric tube, colonoscopy or enema, using screened donor
    stool.</li>
    <li><strong>Stool banks</strong> make it more accessible. <strong>Donor screening is similar to blood bank
    screening</strong>, to reduce the risk from blood-borne and other pathogens: human immunodeficiency virus,
    hepatitis A, B and C, syphilis, <em>Clostridioides difficile</em>, and the parasites <em>Giardia</em> and
    <em>Cryptosporidium</em>.</li>
    <li>Veterinarians call it <strong>&ldquo;transfaunation&rdquo;</strong>, and it works for animals too.</li>
  </ul>

  <h3 class="sub" id="nat-phage">11.5 &middot; Objective 3 &mdash; Phage therapy</h3>
  <p><strong>Bacteriophages</strong> (&ldquo;phages&rdquo;) are very small viruses that specifically attack
  <strong>bacteria and no other types of cells</strong>. Like all viruses they carry DNA or RNA, never both,
  and they replicate by a <strong>lytic</strong> or a <strong>lysogenic</strong> life cycle. For therapy you
  need <strong>lytic phages with DNA</strong>: phages that destroy the bacterium rather than settling into its
  genome. &ldquo;The enemy of my enemy is my friend&rdquo;: the bacteria&rsquo;s own enemies, put to work for
  the patient.</p>
  """ + fig("l11-s18-phage-structure.png", 910, 570,
      "Left, an electron micrograph of tailed bacteriophages with a 500 angstrom scale bar; right, a labeled drawing of a phage: a hexagonal head containing coiled DNA, a tail tube inside a tail sheath, a baseplate at the end of the tail, and two long tail fibers.",
      "<b>A tailed bacteriophage.</b> The <b>head</b> holds the DNA; the <b>tail tube</b>, inside its "
      "<b>sheath</b>, delivers it through the bacterial wall; the <b>baseplate</b> and <b>long tail fibers</b> "
      "attach the phage to its host.", 18) + """
  <p><strong>Why it is an option now.</strong> Phages exhibit <strong>high host specificity</strong>, so they
  could attack highly antibiotic-resistant clinical strains. The idea is not new: it was explored in eastern
  Europe in the 1920s and 1930s and <strong>abandoned in the West with the emergence of antibiotic
  therapy</strong> (penicillin became clinically available in the 1940s). One obstacle today is translating
  the original publications, which are in non-English journals.</p>
  <table>
    <tr><th>Potential benefits (slide 22)</th><th>The bad news (slide 21)</th></tr>
    <tr><td>
      <ul>
        <li>Could be used against <strong>multidrug-resistant</strong> pathogens (extensively and totally
        drug-resistant strains are a concern too).</li>
        <li><strong>Narrow spectrum</strong>, which would preserve the patient&rsquo;s existing
        microbiome.</li>
        <li>Potentially <strong>low risk of side effects</strong>.</li>
        <li><strong>Wide distribution</strong> in the body after systemic administration.</li>
        <li>Potential to <strong>reduce the inflammatory response</strong>.</li>
        <li>Could eventually be <strong>more cost effective</strong> than antibiotics.</li>
        <li>Greater efficacy: phages target specific bacterial cells and then <strong>die off when their host
        is no longer present</strong>.</li>
      </ul></td>
    <td>
      <ul>
        <li><strong>Not a magic bullet</strong>: the optimal dose, route, frequency and duration still need
        to be precisely defined.</li>
        <li>Clinicians <strong>must know the causative agent</strong>, because host specificity is so high;
        identifying it can be time-consuming and cost-prohibitive in resource-limited settings.</li>
        <li><strong>Gene transfer</strong>: what if the phage transfers DNA to its host cell?</li>
        <li>Not covered by public health insurance in some countries (a few European countries are
        encouraging research).</li>
        <li><strong>No phage product is yet approved as a medicine</strong> (the slide: bacterial viruses
        &ldquo;are not currently recognized as medicinal products&rdquo;; in the United States the Food and Drug
        Administration handles them as investigational biologics), so regulatory clearance and unknown safety
        issues are concerns; <strong>investigational drug status</strong> and <strong>emergency use
        authorization</strong> are the useful routes around it.</li>
      </ul></td></tr>
  </table>
  <p>Outside the clinic, phage solutions have also been used to <strong>reduce bacterial contamination of
  food</strong>.</p>

  <h3 class="sub" id="nat-other">11.6 &middot; Objective 4 &mdash; Other alternatives to antibiotics</h3>
  <table>
    <tr><th>Option</th><th>What it does</th></tr>
    <tr><td><strong>Systemic corticosteroids</strong></td><td>Decrease the <strong>host&rsquo;s inflammatory
    response</strong>; they do not kill the pathogen.</td></tr>
    <tr><td><strong>Intravenous immunoglobulin G</strong> and <strong>convalescent plasma</strong></td><td>Give
    ready-made antibodies from donors; convalescent plasma comes from people who have recovered from the
    infection.</td></tr>
    <tr><td><strong>Prebiotics</strong></td><td>Fiber-rich foods that provide <strong>nutrients to the gut
    microflora</strong>.</td></tr>
    <tr><td><strong>Probiotics</strong></td><td>The beneficial organisms themselves: <em>Lactobacillus</em> (a
    bacterium) and <em>Saccharomyces</em> (a fungus, a yeast), often found in <strong>fermented
    foods</strong>.</td></tr>
    <tr><td><strong>A varied plant diet</strong></td><td>From the crowd-sourced American Gut project: <strong>the
    more diverse your plant intake, the more diverse your gut microbiota</strong> (slide 28).</td></tr>
    <tr><td><strong>Plant essential oils</strong></td><td>Show antibacterial activity <strong>in vitro</strong>;
    better experimental design is needed to test what that means for patients.</td></tr>
  </table>
  <p><strong>Why herbs and spices?</strong> The hypotheses (slide 31): herbs and spices season food across
  cultures; some kill bacteria and fungi, especially those that spoil food; they may provide macro- and
  micronutrients; they enhance taste and smell; and spice use should be heaviest where a <strong>hot
  climate</strong> goes with <strong>rapid food spoilage</strong>. The ones named (slide 32) include garlic,
  onion, allspice, oregano, thyme, cinnamon, tarragon, cumin, coriander, basil, ginger, cloves, lemongrass, bay
  leaf, capsicums, rosemary, marjoram, mustard, caraway, mint, parsley, cardamom and eucalyptus.
  &ldquo;Let your food be your medicine, and your medicine be your food&rdquo; (Hippocrates).</p>

  <h3 class="sub" id="nat-truth">11.7 &middot; Where the slide, the recording and current practice differ</h3>
  <p>The quizzes key what is medically true. Where a slide says something else, it is listed here rather than
  silently changed. <strong>None of these is the basis of a quiz question</strong> except where noted.</p>
  <ul>
    <li><strong>&ldquo;Transduction&rdquo; (slide 17 notes).</strong> The notes call a phage injecting its own
    nucleic acid into a bacterium &ldquo;transduction&rdquo;. Injecting its genome is simply how a phage
    infects; <em>transduction</em> is a phage carrying <em>bacterial</em> genes from one bacterium to another.
    That second process is the real &ldquo;what if the phage transfers DNA to its host cell?&rdquo; worry of
    slide 21, since it can move genes such as toxin or resistance genes. <em>Keyed</em> as gene transfer.</li>
    <li><strong>What fecal microbiota transplant is proven for (slide 24).</strong> The slide says it has been
    &ldquo;a successful option&rdquo; for irritable bowel syndrome, ulcerative colitis and chronic constipation.
    Its established use, and the only one with approved products, is preventing <strong>recurrent
    <em>Clostridioides difficile</em> infection</strong>; for the other conditions the evidence is limited and
    mixed. <em>Keyed</em> as <em>Clostridioides difficile</em> only.</li>
    <li><strong>&ldquo;Serologic testing&rdquo; of donor stool (slide 26).</strong> Serology is a blood test.
    Donors are screened with <strong>blood tests</strong> (human immunodeficiency virus, hepatitis, syphilis)
    and <strong>stool tests</strong> (<em>Clostridioides difficile</em>, <em>Giardia</em>,
    <em>Cryptosporidium</em>). <em>Keyed</em> as &ldquo;screened for&rdquo;, without the test type.</li>
    <li><strong>&ldquo;No new major antibiotic in 30 years&rdquo; (slide 6).</strong> Out of date: new agents,
    including new classes, have been approved since. The infographic&rsquo;s percentages are not asked
    either.</li>
    <li><strong>The resistance timeline (slide 8).</strong> Dates for discovery, widespread use and first
    resistance differ between sources (a bacterial enzyme that destroys penicillin was described in 1940,
    before the drug was in wide use). <em>Keyed</em> only as the pattern: resistance followed widespread use within years.</li>
  </ul>
"""

SECTION_FOOT = """
  <button type="button" class="test-yourself-btn" onclick="window.openTestYourself('Test yourself &mdash; Novel Antimicrobial Therapy', TEST_YOURSELF.novelAntimicrobial)">Test yourself! &rarr;</button>
  <p class="src">Source: <em>Novel Antimicrobial Therapies</em> (PAJ5200.Novel Antimicrobial Therapies-2.pptx),
  Slides 1&ndash;34, the 9 October 2026 recording, and the PAJ 5200 syllabus instructional objectives.</p>
</section>
"""

TEST = """
    novelAntimicrobial: [
      {q:"When is antibiotic therapy called empiric?",
       choices:["When it follows culture results","When it starts before results are available","When it is given to contacts","When it is given before surgery"],correct:1,
       explain:"Empiric therapy is started before cultures or other results are available, from clinical experience and antibiograms, then adjusted. Treating contacts or surgical patients is prophylaxis."},
      {q:"Clostridioides difficile infection after antibiotics is which kind of adverse effect?",
       choices:["Direct, an allergy","Direct, a toxicity","Indirect, through the commensal flora","Not an adverse effect"],correct:2,
       explain:"It is indirect: the antibiotic disturbs the commensal microflora and Clostridioides difficile takes over. Allergy, toxicity, drug interactions and therapeutic failure are the direct effects."},
      {q:"Which phages are needed for phage therapy?",
       choices:["Lytic phages","Lysogenic phages","Phages that infect human cells","Phages with no genome"],correct:0,
       explain:"Lytic phages replicate and destroy the bacterium. Lysogenic phages settle into the host's genome instead, and phages never infect human cells."},
      {q:"What must a clinician know before using phage therapy?",
       choices:["The patient's blood type","The causative agent","Only the infection site","The patient's diet"],correct:1,
       explain:"Phage host specificity is very high, so the causative organism must be identified first, which can be slow and costly in resource-limited settings."},
      {q:"How does a prebiotic differ from a probiotic?",
       choices:["Prebiotics feed gut flora; probiotics are live microbes","They are the same thing","Prebiotics are antibiotics","Probiotics are fiber"],correct:0,
       explain:"Prebiotics are fiber-rich foods that nourish the gut microflora; probiotics are beneficial organisms such as Lactobacillus (a bacterium) and Saccharomyces (a fungus)."}
    ],
"""


def section():
    return SECTION_HEAD + AUDIO + BODY + SECTION_FOOT


SECTION = None   # set by build(); build_micro_e2_guide.py reads SECTION after calling build()


def build():
    global SECTION
    SECTION = section()
    return SECTION


build()
