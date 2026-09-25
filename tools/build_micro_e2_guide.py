#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Microbiology Exam 2 study guide (Lectures 7 and 8 so far).

Skeleton lifted from the Microbiology Exam 1 guide -- head, styles, back bar,
pull-to-refresh, footer and scripts -- so the class reads as one thing; only
the title, header, table of contents, body and TEST_YOURSELF bank change. The
forest-green palette is kept on purpose: the Exam 2 quizzes already share
Exam 1's palette.

Sections are numbered by LECTURE (7, 8, ...) exactly as the Exam 1 guide is,
and each lecture's instructional objectives are quoted verbatim from the
syllabus and answered in order. Later Exam 2 lectures (9-13) get their own
_micro_e2_guide_lN.py and a line in SECTIONS below.

Figures are re-extracted from the decks each run (python-pptx, by slide and
image index) into micro-exam-2-study-guide-images/, and every one cites its
slide. Each was viewed at full size before being captioned.

    python3 tools/build_micro_e2_guide.py
"""
import io, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import _micro_e2_guide_l7 as L7
import _micro_e2_guide_l8 as L8

SECTIONS = [L7, L8]
DONOR = os.path.join(ROOT, "Microbiology Exam 1", "micro-exam-1-study-guide.html")
OUTDIR = os.path.join(ROOT, "Microbiology Exam 2")
OUT = os.path.join(OUTDIR, "micro-exam-2-study-guide.html")
IMGDIR = os.path.join(OUTDIR, "micro-exam-2-study-guide-images")
INBOX = os.path.expanduser("~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 2")
DECK7 = os.path.join(INBOX, "Lecture 7  Dr. Webster  Disorders in Immunity.pptx")
DECK8 = os.path.join(INBOX, "PAJ5200.Diagnosing Infections-2.pptx")

# (deck, slide, picture index on that slide) -> published file name
IMAGES = [
    (DECK7, 4, 1, "l7-s04-immunopathology-overview.png"),
    (DECK7, 13, 1, "l7-s13-type-i-sensitization.png"),
    (DECK7, 22, 1, "l7-s22-anaphylaxis-epinephrine.jpg"),
    (DECK7, 46, 1, "l7-s46-acute-rejection.png"),
    (DECK7, 64, 1, "l7-s64-primary-immunodeficiency.png"),
    (DECK8, 14, 1, "l8-s14-specimen-scheme.jpg"),
    (DECK8, 46, 1, "l8-s46-complement-fixation.jpg"),
    (DECK8, 48, 1, "l8-s48-fluorescent-antibody.jpg"),
    (DECK8, 51, 1, "l8-s51-elisa.jpg"),
]

# No abbreviations in guide prose (no_abbreviations_content_policy): the two
# nucleic acids are written out, as the Exam 1 guide does. Quotes keep the
# speaker's words but gain the expansion in brackets.
PROSE_FIXES = [
    (r"(?<![\w(\[])DNA(?![\w)\]]| \[)", "deoxyribonucleic acid"),
    (r"(?<![\w(\[])RNA(?![\w)\]]| \[)", "ribonucleic acid"),
    (r"dies of HIV\.", "dies of HIV [human immunodeficiency virus]."),
    (r"as well as HIV and", "as well as HIV [human immunodeficiency virus] and"),
    (r"making the IgG antibodies", "making the IgG [immunoglobulin G] antibodies"),
    (r"allergy as &ldquo;IgA mediated&rdquo;", "allergy as &ldquo;IgA [immunoglobulin A] mediated&rdquo;"),
]


def extract_images():
    from pptx import Presentation
    os.makedirs(IMGDIR, exist_ok=True)
    decks = {}
    for deck, slide, k, name in IMAGES:
        prs = decks.setdefault(deck, Presentation(deck))
        pics = [sh for sh in prs.slides[slide - 1].shapes if sh.shape_type == 13]
        assert len(pics) >= k, "%s slide %d has %d picture(s)" % (os.path.basename(deck), slide, len(pics))
        blob = pics[k - 1].image.blob
        ext = pics[k - 1].image.ext
        assert name.endswith("." + ext.replace("jpeg", "jpg")), (name, ext)
        io.open(os.path.join(IMGDIR, name), "wb").write(blob)
    return len(IMAGES)


def main():
    n_img = extract_images()
    donor = io.open(DONOR, encoding="utf-8").read()
    head = donor[:donor.index('<div class="layout wrap"')]
    foot = donor[donor.index('<footer class="guide-foot">'):donor.index("</footer>") + len("</footer>")]
    tail = donor[donor.index("</main>"):]
    ts = tail.index("var TEST_YOURSELF = {")
    te = tail.index("\n  };", ts)
    bank = "".join(s.TEST for s in SECTIONS).rstrip().rstrip(",")
    tail = tail[:ts] + "var TEST_YOURSELF = {\n" + bank + "\n" + tail[te:]

    head = re.sub(r"<title>.*?</title>",
                  "<title>Microbiology &middot; Exam 2 &mdash; Study Guide</title>", head, count=1, flags=re.S)
    head = re.sub(r"<header class=\"top\">.*?</header>",
        '<header class="top">\n  <h1>Microbiology &middot; Exam 2 &mdash; Study Guide</h1>\n'
        '  <p>PAJ 5200 Microbiology &middot; Class of 2028</p>\n'
        '  <p>Covers Lectures 7 and 8 so far &mdash; Disorders in Immunity and Diagnosing Infections &middot; '
        'Exam 2 (Lectures 7&ndash;13) is on Friday 23 October 2026; further sections are added as each '
        'lecture is posted &middot; Instructional Objectives (IOs) taken verbatim from the syllabus</p>\n</header>',
        head, count=1, flags=re.S)
    # the Word-export link is re-added by tools/build_guide_docx.py for THIS page
    head = re.sub(r'\n?\s*<link rel="alternate" type="application/vnd\.openxmlformats[^>]*>', "", head)

    toc = '<nav class="toc">\n  <h2>Contents</h2>' + "".join(s.TOC for s in SECTIONS) + "</nav>"
    body = "<main>\n" + "\n".join(s.SECTION for s in SECTIONS) + "\n" + foot + "\n"
    for pat, rep in PROSE_FIXES:
        body = re.sub(pat, rep, body)
    html = head + '<div class="layout wrap" data-readable>\n' + toc + "\n\n" + body + tail

    # guards: structure the rest of the site relies on
    assert html.count('<section class="deck"') == len(SECTIONS)
    assert "guide-back-bar" in html and "window.reportMistake()" in html and "../index.html" in html
    assert "data-audio-dir" not in html
    for m in re.finditer(r'href="#([\w-]+)"', toc):
        assert 'id="%s"' % m.group(1) in html, "TOC anchor without target: " + m.group(1)
    for m in re.finditer(r'src="(micro-exam-2-study-guide-images/[^"]+)"', html):
        assert os.path.exists(os.path.join(OUTDIR, m.group(1))), m.group(1)
    for s in SECTIONS:
        key = re.search(r"TEST_YOURSELF\.(\w+)", s.SECTION).group(1)
        assert key + ":" in html, "no TEST_YOURSELF bank for " + key

    io.open(OUT, "w", encoding="utf-8").write(html)
    print("wrote %s (%d KB, %d figures)" % (os.path.relpath(OUT, ROOT), len(html) // 1024, n_img))


if __name__ == "__main__":
    main()
