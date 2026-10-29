#!/usr/bin/env python3
"""Write the Lecture 14 (Gram-Negative Bacilli) QUIZ pictures into
"Microbiology Exam 3/micro-exam-3-quiz-images/".

Quiz pictures get NEUTRAL, numbered file names (l14-fig-01.jpg ...; memory quiz_image_filenames): a picture opened
on its own must never name its answer. The numbers follow slide order and are fixed here, so a rebuild writes the
same names. The study guide keeps its own, descriptive copies (micro-exam-3-study-guide-images/).

QUIZ images carry NO label that names the answer, so each is a crop of the slide's picture:
  slide 16  the TOP plate only (greenish-blue agar around the growth); the lower plate and the two credit lines
            cropped off
  slide 44  panel (b) only, the invasive pattern (bacteria in the lining with blood); panel (a), the "(a)/(b)"
            letters and the "Intestinal epithelium" label cropped off (the "Bacteria" and "Blood" labels stay: they
            name what the picture shows, not the organism)
  slide 64  the whole photograph (Rainbow agar; no label)
  slide 67  the micrograph below the "Cell body" / "Capsule" labels (their leader lines remain, without words)
  slide 71  the LEFT picture by geometry (the swarming plate; the right one is the flagella stain)
  slide 82  the smear and its inset below the "White blood cell" / "Y. pestis" labels
Each picture is chosen by its position on the slide (python-pptx left), never by an extraction index, and each
crop box was read off the original pixels and the result viewed.

    python3 tools/extract_micro_l14_quiz_images.py
"""
import io, os
from pptx import Presentation
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DECK = os.path.expanduser("~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 3/"
                          "Lecture 14  Dr. Webster Gram Negative Bacilli of Medical Importance.pptx")
OUT = os.path.join(ROOT, "Microbiology Exam 3", "micro-exam-3-quiz-images")

# (slide, which picture by position: "only" | "left" | "right", crop box on the original pixels or None, name)
JOBS = [
    (16, "only", (20, 0, 425, 395), "l14-fig-01.jpg"),
    (44, "only", (425, 75, 860, 860), "l14-fig-02.jpg"),
    (64, "only", None, "l14-fig-03.jpg"),
    (67, "only", (0, 62, 641, 409), "l14-fig-04.jpg"),
    (71, "left", None, "l14-fig-05.jpg"),
    (82, "only", (0, 42, 791, 484), "l14-fig-06.jpg"),
]
MAXW = 600


def pictures(shapes):
    """Pictures, picture placeholders and grouped pictures, in shape-tree order."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from pictures(sh.shapes)
            continue
        try:
            sh.image
        except Exception:
            continue
        yield sh


def pick(slide, which):
    pics = list(pictures(slide.shapes))
    if which == "only":
        assert len(pics) == 1, "expected one picture, found %d" % len(pics)
        return pics[0]
    pics.sort(key=lambda sh: sh.left)
    assert len(pics) == 2, "expected two pictures side by side"
    return pics[0] if which == "left" else pics[1]


def main():
    prs = Presentation(DECK)
    os.makedirs(OUT, exist_ok=True)
    for n, which, box, name in JOBS:
        sh = pick(prs.slides[n - 1], which)
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
        if box:
            im = im.crop(box)
        if im.width > MAXW:
            im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
        im.save(os.path.join(OUT, name), quality=85)
        print("wrote %-16s %dx%d  (slide %d, %s picture)" % (name, im.width, im.height, n, which))


if __name__ == "__main__":
    main()
