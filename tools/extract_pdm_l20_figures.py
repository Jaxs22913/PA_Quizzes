#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pull the Lecture 20 (Pulmonary Function Testing and Oxygenation) figures for PDM I Exam 4: the tracings,
loops and reports the quizzes ask the student to read, and the teaching figures for the guide.

Deck: "Pulmonary Function Testing - Prof. Fowler.pptx" (Professor Fowler, MMS, PA-C, title slide; 26 October
2026). All 57 slides were read and every picture in the deck was viewed before it was chosen here (the deck has
no rendering of its own here, so each embedded picture was pulled at full size and read). Slide images are
cleared for use provided the slide is cited.

QUIZ images carry NO label that names the answer. What was done to each:
  slide 3   the two films (hyperinflated; small-volume) cut apart; the slide's captions are separate text, not
            in the pictures.
  slide 20  the lung-volume tracing (slides 20, 22, 25, 27 repeat it) cropped to the tracing, the labelled
            boxes on the right cut away, one compartment shaded in yellow per question.
  slide 28  panel B only, its "AIRWAY OBSTRUCTION" title masked (the printed ratio stays: reading it is the task).
  slide 29  the volume-time curve with its three labels masked and one arrow, drawn by us, to the point at one
            second.
  slide 30  the top panel (the loop) with its "Peak expiratory flow", FEV1 and VC labels masked; the slide's own
            arrowhead at the peak stays (its shaft went with the label).
  slide 34  the obstructive volume-time curve, unlabelled as it stands.
  slide 36  the right-hand (obstructive) loop alone; slide 37 the severe-obstruction loop alone (its title is
            separate slide text; the picture's own ghost lettering is illegible).
  slide 40  the normal-against-restrictive volume-time curves with NORMAL and RESTRICTIVE masked; the
            normal-against-restrictive loops with RESTRICTIVE masked.
  slide 44  the three loops cut apart above their Normal / Obstructive / Restrictive captions.
  slide 45  the mild-obstruction, severe-obstruction and restriction panels cut apart above their captions,
            every text box masked (the dashed normal outline stays).
  slides 47, 51, 56  the three reports with the patient's (fictional) name masked; age and sex stay. The quiz
            copies of slides 47 and 51 are cropped to the report card (slide 51's "Case 2" heading and both films
            cut away); the guide keeps the whole picture.

Left out: slide 4 (stock photograph), slide 14 (three spirometer product photographs; the guide names the
device instead), slide 15's two home peak flow meters (one kept), slide 17 (incentive spirometer photograph:
the slide carries only its title), slides 33/38 (stock background photographs), slide 57 (closing picture).
"""
import io, os, sys, zipfile
from PIL import Image, ImageDraw
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from extract_pdm_l17_figures import apply as _apply, WHITE, MAXW  # noqa: E402  (same drawing helpers as Lecture 17)

ROOT = os.path.dirname(HERE)
DECK = os.path.expanduser(
    "~/Desktop/PA Quizzes/Semester 2/Principles of Diagnostic Medicine I Inbox/Exam 4/"
    "Pulmonary Function Testing - Prof. Fowler.pptx")
FOLDER = os.path.join(ROOT, "Principles of Diagnostic Medicine I Exam 4")
QUIZ = os.path.join(FOLDER, "pdm-exam-4-quiz-images")
GUIDE = os.path.join(FOLDER, "pdm-exam-4-study-guide-images")

# The slide-20 tracing, in its own pixels: horizontal lines at 80 (y 10), 37 (y 388), 30 (y 508), 15 (y 693)
# and 0 (y 860) millilitres per kilogram; the tracing itself ends at x 850, where the labelled boxes begin.
VOL_CROP = (0, 0, 852, 865)
BANDS = {"rv": (100, 694, 848, 858), "erv": (100, 509, 848, 692), "irv": (100, 12, 848, 387),
         "frc": (100, 509, 848, 858)}

# s45 text boxes (full-picture pixels), masked before each panel is cut out
P45_MILD = [(325, 10, 446, 60), (302, 174, 426, 223), (336, 219, 414, 254), (322, 256, 362, 284)]
P45_SEVERE = [(536, 10, 645, 60), (517, 174, 632, 223), (543, 243, 621, 279)]
P45_RESTRICT = [(764, 10, 885, 60), (741, 174, 861, 224), (774, 248, 870, 282)]


def m(boxes):
    return [(b, None) for b in boxes]


# (slide, picture shape id, output stem, destinations, ops)
# ops: everything extract_pdm_l17_figures.apply takes (crop LAST, masks, arrows, stack), plus band=(x0,y0,x1,y1).
FIGS = [
 # ---- the two breathless patients
 (3, 128, "l20-film-hyperinflated", "Q", {}),
 (3, 129, "l20-film-small-volume", "Q", {}),
 (3, 128, "l20-two-patients-a", "G", {}),
 (3, 129, "l20-two-patients-b", "G", {}),
 # ---- equipment
 (12, 200, "l20-equipment", "G", {}),
 (13, 206, "l20-plethysmograph", "G", {}),
 (15, 223, "l20-peak-flow-meter", "G", {}),
 (16, 232, "l20-pulse-oximeters", "G", {}),
 # ---- volumes, capacities and the spirogram
 (20, 258, "l20-lung-volumes", "G", {}),
 (20, 258, "l20-volumes-band-rv", "Q", {"band": BANDS["rv"], "crop": VOL_CROP}),
 (20, 258, "l20-volumes-band-erv", "Q", {"band": BANDS["erv"], "crop": VOL_CROP}),
 (20, 258, "l20-volumes-band-irv", "Q", {"band": BANDS["irv"], "crop": VOL_CROP}),
 (20, 258, "l20-volumes-band-frc", "Q", {"band": BANDS["frc"], "crop": VOL_CROP}),
 (28, 306, "l20-volume-time-normal-obstructed", "G", {}),
 (28, 306, "l20-vt-panel-b", "Q", {"masks": m([(50, 240, 340, 268)]), "crop": (0, 236, 425, 451)}),
 (29, 314, "l20-volume-time-fev1", "G", {}),
 (29, 314, "l20-vt-arrow-one-second", "Q", {"masks": [(b, WHITE) for b in ((80, 56, 149, 86), (150, 92, 202, 130), (428, 47, 480, 68))],
                                            "arrows": [((250, 205), (152, 122))]}),
 (30, 320, "l20-flow-volume-normal", "G", {}),
 (30, 320, "l20-loop-arrow-top", "Q", {"masks": m([(97, 6, 260, 32), (163, 94, 202, 119), (122, 279, 152, 296)]),
                                       "crop": (0, 0, 275, 300)}),
 # ---- obstruction
 (34, 353, "l20-vt-slow-rise", "QG", {}),
 (36, 364, "l20-loop-normal-slide36", "G", {}),
 (36, 365, "l20-loop-concave", "QG", {}),
 (37, 372, "l20-loop-flattened", "QG", {}),
 # ---- restriction
 (40, 470, "l20-restrictive-volume-time", "G", {}),
 (40, 470, "l20-vt-two-curves", "Q", {"masks": [(b, (239, 239, 239)) for b in ((345, 58, 445, 100), (330, 222, 452, 262))]}),
 (40, 472, "l20-restrictive-loop", "G", {}),
 (40, 472, "l20-loop-two-overlaid", "Q", {"masks": m([(600, 140, 925, 245)])}),
 # ---- interpretation
 (41, 479, "l20-interpretation-table", "G", {}),
 (42, 485, "l20-spirometry-algorithm", "G", {}),
 (44, 505, "l20-loop-shapes", "G", {}),
 (44, 505, "l20-loop-panel-left", "Q", {"crop": (0, 0, 372, 412)}),
 (44, 505, "l20-loop-panel-middle", "Q", {"crop": (372, 0, 768, 412)}),
 (44, 505, "l20-loop-panel-right", "Q", {"crop": (768, 0, 1170, 412)}),
 (45, 512, "l20-loop-variations", "G", {}),
 (45, 512, "l20-variation-second", "Q", {"masks": m(P45_MILD), "crop": (237, 0, 452, 408)}),
 (45, 512, "l20-variation-third", "Q", {"masks": m(P45_SEVERE), "crop": (455, 0, 669, 408)}),
 (45, 512, "l20-variation-fourth", "Q", {"masks": m(P45_RESTRICT), "crop": (672, 0, 886, 408)}),
 # ---- the three reports
 (47, 526, "l20-report-56-man", "G", {"masks": m([(530, 41, 598, 62)])}),
 (47, 526, "l20-report-card-56", "Q", {"masks": m([(530, 41, 598, 62)]), "crop": (440, 24, 898, 432)}),
 (51, 551, "l20-report-32-woman", "G", {"masks": m([(566, 106, 695, 125)])}),
 (51, 551, "l20-report-card-32", "Q", {"masks": m([(566, 106, 695, 125)]), "crop": (474, 88, 922, 468)}),
 (56, 582, "l20-report-15-boy", "QG", {"masks": m([(89, 42, 163, 65)])}),
]


def blobs():
    """{(slide, shape id): bytes}, read straight from the package (groups included) so nothing is retyped."""
    out = {}
    prs = Presentation(DECK)
    z = zipfile.ZipFile(DECK)
    for i, s in enumerate(prs.slides, 1):
        rels = {r.rId: r.target_ref for r in s.part.rels.values()}

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:
                    walk(sh.shapes)
                    continue
                emb = sh._element.xpath('.//a:blip/@r:embed')
                tgt = rels.get(emb[0]) if emb else None
                if tgt and "media/" in tgt:
                    out[(i, sh.shape_id)] = z.read("ppt/media/" + tgt.split("media/")[-1])
        walk(s.shapes)
    return out


def apply(im, ops):
    if "band" in ops:                            # a translucent yellow band with a solid edge
        x0, y0, x1, y1 = ops["band"]
        over = Image.new("RGBA", im.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(over)
        d.rectangle((x0, y0, x1, y1), fill=(255, 214, 0, 95), outline=(214, 160, 0, 255), width=4)
        im = Image.alpha_composite(im.convert("RGBA"), over).convert("RGB")
    return _apply(im, {k: v for k, v in ops.items() if k != "band"})


def main():
    os.makedirs(QUIZ, exist_ok=True)
    os.makedirs(GUIDE, exist_ok=True)
    B = blobs()
    # A quiz picture's file name must not name its answer: quiz copies are numbered in the order of their
    # descriptive stems; the guide keeps the descriptive names (as for Lecture 17).
    QUIZ_NAME = {st: "l20-fig-%02d" % i
                 for i, st in enumerate(sorted(st for _, _, st, d, _ in FIGS if "Q" in d), 1)}
    for slide, sid, stem, dest, ops in FIGS:
        assert (slide, sid) in B, "slide %d has no picture %d" % (slide, sid)
        im = Image.open(io.BytesIO(B[(slide, sid)]))
        if getattr(im, "is_animated", False):
            im.seek(0)
        if im.mode in ("P", "RGBA", "LA") or "transparency" in im.info:
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, WHITE)
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        im = apply(im, ops)
        if im.width < 600:                      # small tracings read badly at natural size
            f = min(3, -(-700 // im.width))
            im = im.resize((im.width * f, im.height * f), Image.LANCZOS)
        if im.width > MAXW:
            im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
        for dd in ([QUIZ] if "Q" in dest else []) + ([GUIDE] if "G" in dest else []):
            p = os.path.join(dd, (QUIZ_NAME[stem] if dd == QUIZ else stem) + ".jpg")
            im.save(p, quality=88)
            print("%-5s slide %2d  %-40s %-40s %4dx%-4d %3d KB" % (
                "quiz" if dd == QUIZ else "guide", slide, stem, os.path.basename(p), im.width, im.height,
                os.path.getsize(p) // 1024))


if __name__ == "__main__":
    main()
