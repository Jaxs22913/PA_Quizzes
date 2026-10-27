# -*- coding: utf-8 -*-
"""Section 14 of the Microbiology Exam 3 guide -- Gram-Negative Bacilli of Medical Importance (Dr. Webster).

Same shape as tools/micro_e2/_micro_e2_guide_l13.py (FIG, DECK, IMAGES, extract(), fig(), TOC, SECTION, TEST):
tools/build_micro_e3_guide.py imports it, calls extract() for the figures and splices TOC, SECTION and TEST.
Later Exam 3 lectures get their own _micro_e3_guide_lN.py beside this one.

Attribution: Dr. Webster, from the deck file name ("Lecture 14  Dr. Webster Gram Negative Bacilli of Medical
Importance.pptx") and the calendar row for 23 October 2026 (1-4 pm). The title slide names no lecturer.

Instructional objectives are VERBATIM from the syllabus (Micro.pdf, page 5): the two numbered objectives under
"Gram Negative Bacilli of Medical Importance", which points to Appendix 2's list of eighteen organisms for Lecture 14
(page 9). The deck's own objective slide (slide 2) matches them, and slides 3-4 reproduce the list. Objective 1 is
answered organism by organism (14.1-14.19); objective 2 is drawn together in 14.20.

Content is from the VISIBLE slides only, text and pictures (the deck carries no speaker notes). All 64 pictures
were read at full size, including the ones that carry content no text box has: slide 5 (lipid A "most responsible
for pathogenesis"), slide 6 (smooth to rough lipopolysaccharide), slide 12 (Table 20.1), slide 32 (the pertussis
course), slide 44 (toxigenic against invasive diarrhea), slide 47 (the enteric biochemical tree), slide 53 (the
H, K and O antigens), slide 74 (the typhoid route), slide 84 (the plague cycle) and slide 86 (the plague warning
sign). Case-count charts, maps and outbreak tallies (slides 33-36, 42, 61-63, 75-77, 85-87, 91-93, 99) are
summarized, not tabulated.

AUDIO: PHASE 1 IS SLIDES ONLY. The recording (a classmate's Notability export, three segments, 2:03:43, no
Notability transcript) is being transcribed; the MICROL14AUDIO block below is the placeholder the audio pass
replaces with what was stressed, what was said to skip, and any correction of a slide, as in the Exam 2 blocks.
"""
import io, os

FIG = "micro-exam-3-study-guide-images"
DECK = "Lecture 14  Dr. Webster Gram Negative Bacilli of Medical Importance.pptx"
INBOX_DECK = os.path.expanduser(
    "~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 3/" + DECK)

# (slide, picture index on that slide in shape-tree order counting group members and placeholder pictures,
# published file name). Multi-picture slides resolved by geometry: on slide 19 picture 3 is the corneal ulcer
# (picture 2 the skin rash, picture 4 the ear); on slide 48 picture 1 is the RIGHT plate under "MacConkey's Agar"
# and picture 2 the LEFT plate under "EMB agar"; on slide 68 picture 2 is the red colonies (picture 1 is a
# bathroom-tile photograph); on slide 71 picture 2 is the swarming plate (picture 1 the flagella stain).
# A .jpg name for a PNG blob (or a JPEG over 250 KB or over 900 px wide) means re-encoded as JPEG, quality 85, at most 900 px wide.
IMAGES = [
    (5, 1, "l14-s005-lipopolysaccharide.jpg"),
    (6, 1, "l14-s006-lps-smooth-rough.png"),
    (12, 1, "l14-s012-gram-negative-scheme.png"),
    (15, 1, "l14-s015-polar-flagellum.jpg"),
    (16, 1, "l14-s016-pseudomonas-pigment.jpg"),
    (18, 1, "l14-s018-pseudomonas-virulence.png"),
    (19, 1, "l14-s019-pseudomonas-infection-sites.jpg"),
    (19, 3, "l14-s019-corneal-ulcer.jpg"),
    (22, 1, "l14-s022-burkholderia-cepacia.png"),
    (24, 1, "l14-s024-melioidosis.jpg"),
    (29, 1, "l14-s029-tularemia-ulcer.jpg"),
    (32, 1, "l14-s032-pertussis-course.jpg"),
    (41, 1, "l14-s041-legionella.jpg"),
    (44, 1, "l14-s044-toxigenic-invasive.jpg"),
    (47, 1, "l14-s047-enteric-biochemistry.jpg"),
    (48, 2, "l14-s048-emb-agar.jpg"),
    (48, 1, "l14-s048-macconkey-agar.jpg"),
    (53, 1, "l14-s053-enteric-antigens.png"),
    (64, 1, "l14-s064-rainbow-agar.jpg"),
    (67, 1, "l14-s067-klebsiella-capsule.jpg"),
    (68, 2, "l14-s068-serratia-colonies.png"),
    (71, 2, "l14-s071-proteus-swarming.jpg"),
    (74, 1, "l14-s074-typhoid-route.jpg"),
    (79, 1, "l14-s079-shigella-colon.jpg"),
    (82, 1, "l14-s082-yersinia-pestis.jpg"),
    (84, 1, "l14-s084-plague-cycle.jpg"),
    (101, 1, "l14-s101-conjunctivitis.jpg"),
]
MAXW = 900


def _pictures(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pictures(sh.shapes)
            continue
        try:
            sh.image
        except Exception:
            continue
        yield sh


def extract(imgdir):
    """Write IMAGES into imgdir from the deck (pictures, picture placeholders and grouped pictures)."""
    from pptx import Presentation
    from PIL import Image
    prs = Presentation(INBOX_DECK)
    os.makedirs(imgdir, exist_ok=True)
    for slide, k, name in IMAGES:
        pics = list(_pictures(prs.slides[slide - 1].shapes))
        assert len(pics) >= k, "%s slide %d has %d picture(s)" % (DECK, slide, len(pics))
        img = pics[k - 1].image
        ext = img.ext.replace("jpeg", "jpg")
        out = os.path.join(imgdir, name)
        if name.endswith(".jpg") and (ext != "jpg" or len(img.blob) > 250 * 1024 or Image.open(io.BytesIO(img.blob)).width > MAXW):
            im = Image.open(io.BytesIO(img.blob)).convert("RGB")
            if im.width > MAXW:
                im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
            im.save(out, quality=85)
        else:
            assert name.endswith("." + ext), (name, ext)
            open(out, "wb").write(img.blob)
    return len(IMAGES)


def fig(name, w, h, alt, cap, slide, credit=""):
    note = (credit + " ") if credit else ""
    return ('<figure class="fig"><img width="%d" height="%d" loading="lazy" src="%s/%s" alt="%s">'
            '<figcaption>%s <span class="src">%sSource: %s, Slide %d.</span></figcaption></figure>'
            % (w, h, FIG, name, alt, cap, note, DECK, slide))


MH = "Copyright &copy; The McGraw-Hill Companies."

TOC = """
  <a class="top-link" href="#gram-negative-bacilli">14 &middot; Gram-Negative Bacilli of Medical Importance</a>
  <a class="sub-link" href="#gnb-audio">&#9733; What the recording adds</a>
  <a class="sub-link" href="#gnb-endotoxin">14.1 Objective 1 &mdash; Endotoxin, the shared weapon</a>
  <a class="sub-link" href="#gnb-sorting">14.2 Objective 1 &mdash; Sorting the Gram-negative bacilli</a>
  <a class="sub-link" href="#gnb-pseudomonas">14.3 Objective 1 &mdash; <em>Pseudomonas aeruginosa</em></a>
  <a class="sub-link" href="#gnb-relatives">14.4 Objective 1 &mdash; <em>Burkholderia</em>, <em>Acinetobacter</em>, <em>Stenotrophomonas</em></a>
  <a class="sub-link" href="#gnb-brucella">14.5 Objective 1 &mdash; <em>Brucella</em> and brucellosis</a>
  <a class="sub-link" href="#gnb-francisella">14.6 Objective 1 &mdash; <em>Francisella tularensis</em> and tularemia</a>
  <a class="sub-link" href="#gnb-bordetella">14.7 Objective 1 &mdash; <em>Bordetella pertussis</em> and whooping cough</a>
  <a class="sub-link" href="#gnb-alcaligenes">14.8 Objective 1 &mdash; <em>Alcaligenes</em></a>
  <a class="sub-link" href="#gnb-legionella">14.9 Objective 1 &mdash; <em>Legionella</em> and legionellosis</a>
  <a class="sub-link" href="#gnb-enterics">14.10 Objective 1 &mdash; The Enterobacteriaceae: traits, media, antigens</a>
  <a class="sub-link" href="#gnb-ecoli">14.11 Objective 1 &mdash; <em>Escherichia coli</em> and its four pathotypes</a>
  <a class="sub-link" href="#gnb-coliforms">14.12 Objective 1 &mdash; The other coliforms</a>
  <a class="sub-link" href="#gnb-proteus">14.13 Objective 1 &mdash; <em>Proteus</em> and its relatives</a>
  <a class="sub-link" href="#gnb-salmonella">14.14 Objective 1 &mdash; <em>Salmonella</em>: typhoid and the salmonelloses</a>
  <a class="sub-link" href="#gnb-shigella">14.15 Objective 1 &mdash; <em>Shigella</em> and bacillary dysentery</a>
  <a class="sub-link" href="#gnb-yersinia-enteric">14.16 Objective 1 &mdash; The enteric yersinias</a>
  <a class="sub-link" href="#gnb-plague">14.17 Objective 1 &mdash; <em>Yersinia pestis</em> and plague</a>
  <a class="sub-link" href="#gnb-pasteurella">14.18 Objective 1 &mdash; <em>Pasteurella multocida</em></a>
  <a class="sub-link" href="#gnb-haemophilus">14.19 Objective 1 &mdash; <em>Haemophilus</em></a>
  <a class="sub-link" href="#gnb-transmission">14.20 Objective 2 &mdash; Modes of transmission</a>
  <a class="sub-link" href="#gnb-notes">14.21 Where the slides disagree, and where current practice differs</a>
"""

SECTION_HEAD = """
<section class="deck" id="gram-negative-bacilli">
  <h2 class="deck-title">14 &middot; Gram-Negative Bacilli of Medical Importance</h2>
  <p class="lecturer">Dr. Webster &middot; Lecture 14 &middot; 23 October 2026</p>
  <div class="io-box">
    <h3>Instructional Objectives</h3>
    <ol>
      <li>Compare and contrast the morphology, medium, laboratory tests, biochemical reactions, diseases,
      clinical manifestations, methods of treatment, and prevention of Gram-negative bacilli of medical
      importance.</li>
      <li>Describe the modes of transmission by which humans acquire Gram-negative bacilli of medical
      importance.</li>
    </ol>
  </div>

  <div class="callout"><strong>One weapon, three questions.</strong> Every Gram-negative rod carries
  <strong>endotoxin</strong>, the lipopolysaccharide of its outer membrane (14.1). Then three questions sort them
  (14.2): <em>what oxygen does it need, does it ferment lactose, is it oxidase positive?</em> The
  <strong>aerobes that do not ferment</strong> are <em>Pseudomonas</em> and its relatives, <em>Brucella</em>,
  <em>Francisella</em>, <em>Bordetella</em> and <em>Alcaligenes</em>, with <em>Legionella</em> paired alongside <em>Bordetella</em>; the
  <strong>facultative anaerobes</strong> are the enterics, split into lactose-fermenting
  <strong>coliforms</strong> (<em>Escherichia</em>, <em>Klebsiella</em>, <em>Enterobacter</em>,
  <em>Citrobacter</em>) and non-lactose-fermenting <strong>non-coliforms</strong> (<em>Proteus</em>,
  <em>Salmonella</em>, <em>Shigella</em>, <em>Yersinia</em>), plus the oxidase-positive <em>Haemophilus</em> and
  <em>Pasteurella</em>. The syllabus points to its Appendix 2 list of eighteen organisms for this lecture
  (slides 3&ndash;4 reproduce it): <em>Pseudomonas aeruginosa</em> and the genera like it, <em>Brucella</em>,
  <em>Francisella tularensis</em>, <em>Bordetella pertussis</em>, <em>Legionella</em>, the coliform and
  non-coliform enterics of the normal flora, the four <em>Escherichia coli</em> pathotypes (enterotoxigenic,
  enteroinvasive, enteropathogenic, enterohemorrhagic O157:H7), <em>Proteus</em>, <em>Salmonella typhi</em> and
  the other salmonellae, <em>Shigella</em>, <em>Yersinia pestis</em> and <em>Haemophilus</em>. Objective 1 is
  answered organism by organism against its own headings (morphology, medium, tests, biochemistry, disease,
  manifestations, treatment, prevention); objective 2 is drawn together in 14.20.</div>

<!--MICROL14AUDIO-->
  <div class="prof-flag" id="gnb-audio"><span class="prof-flag-label">&#9733; From the lecture recording &mdash; 23 October 2026</span>
  <p><b>Not read yet.</b> A classmate&rsquo;s recording of this lecture (three parts, about 2 hours 4 minutes) has
  arrived and is being transcribed. It came without Notability&rsquo;s own transcript, so ours will be the only
  one, with no second opinion to compare against. Until it has been read, this section is built from the slides
  alone: the quizzes are weighted by slide count, and the only emphasis marked is the slide&rsquo;s own (the
  three asterisks after <em>cystic fibrosis</em> on slide 18, 14.3). When the recording has been read, this box
  will carry what Dr. Webster stressed, what she said to skip, and any correction of a slide.</p>
  </div>
<!--/MICROL14AUDIO-->
"""

BODY = """
  <h3 class="sub" id="gnb-endotoxin">14.1 &middot; Objective 1 &mdash; Endotoxin, the shared weapon</h3>
  <p><strong>Endotoxin is a structural component of the Gram-negative outer membrane.</strong> Chemically it is
  <strong>lipopolysaccharide</strong>; <strong>lipid A</strong> is the biologically active component of the
  complex, the part most responsible for pathogenesis. It is released <strong>on the lysis of Gram-negative
  cells, including through antibiotic therapy</strong>, and <strong>as bacterial cells grow and divide,
  particularly when they are rapidly proliferating</strong>.</p>
  """ + fig("l14-s005-lipopolysaccharide.jpg", 900, 321,
      "Left, Gram-negative rods; middle, a section of the cell wall with the lipopolysaccharide layer on the outer membrane above the peptidoglycan; right, one lipopolysaccharide unit with its polysaccharide chain attached to lipid A, beside an electron micrograph of a rod.",
      "<b>Where endotoxin sits.</b> The lipopolysaccharide layer is the outside of the outer membrane, above the "
      "peptidoglycan. In each unit a polysaccharide chain rides on <b>lipid A</b>, labeled here as the part most "
      "responsible for pathogenesis.", 5, MH) + """
  <table>
    <tr><th>Part of the lipopolysaccharide</th><th>What it is</th></tr>
    <tr><td><strong>Lipid A</strong></td><td>The lipid region to which the polysaccharide is bound; the <strong>biologically active</strong> component</td></tr>
    <tr><td><strong>Core polysaccharide</strong></td><td><strong>Identical for large groups of bacteria</strong>, with some unique sugars</td></tr>
    <tr><td><strong>O antigen polysaccharide</strong></td><td><strong>Unique for each organism</strong>; repeating units of 3&ndash;4 different hexoses. The <strong>number of repeating units</strong> designates the antigen as <strong>rough, semi-rough or smooth</strong></td></tr>
  </table>
  """ + fig("l14-s006-lps-smooth-rough.png", 377, 231,
      "Three drawings of lipopolysaccharide: smooth, with lipid A, the core polysaccharide and a long chain of O antigen units; semi-rough, with two O antigen units; rough, with lipid A and only part of the core.",
      "<b>Smooth, semi-rough, rough.</b> The more repeating O antigen units on the core, the smoother the "
      "lipopolysaccharide; a rough one has none.", 6) + """
  <h4 class="subsub">What endotoxin does</h4>
  <ul>
    <li><strong>Usually protective.</strong> In the gut and airways the first contact between epithelium and
    Gram-negative bacteria is the outer membrane and its lipopolysaccharide. These are antigens of the
    <strong>innate (nonspecific) immune system</strong>, signaling through a <strong>lipopolysaccharide binding
    complex and a Toll-like receptor</strong>. Under most circumstances this gives a <strong>moderated,
    protective inflammatory response</strong>, which keeps most mucosal infections mucosal and helps maintain our
    symbiosis with some Gram-negative bacilli.</li>
    <li><strong>Chronic infection.</strong> Some bacteria have adapted by <strong>modifying their
    lipopolysaccharide</strong> to evade the immune system.</li>
    <li><strong>Potentially damaging</strong>, through complement, coagulation and a variety of cells:
      <ul>
        <li><strong>Complement</strong> turns on by the classical pathway (antigen-antibody reaction) or the
        alternative pathway, and can result in <strong>hemorrhagic necrosis</strong>.</li>
        <li><strong>Coagulation</strong>: tissue injury through coagulation changes, <strong>disseminated
        intravascular coagulation</strong>; fibrin deposited in the lungs, liver and spleen; it may trigger
        <strong>bradykinin, a vasodilator in shock</strong>.</li>
        <li><strong>Platelets</strong>: altered activity causing aggregation (differs among species).
        <strong>Neutrophils</strong>: low levels may enhance phagocytosis, but <strong>higher levels may interfere
        with chemotaxis</strong> to the site of infection. <strong>Macrophages</strong> release pro-coagulants,
        collagenase, pyrogens, prostaglandins, colony stimulating factor, tumor necrosis factor, interleukin 1 and
        interleukin 6, nitric oxide and reactive oxygen species.</li>
      </ul></li>
  </ul>
  <div class="pearl"><strong>Low dose against flood.</strong> Low levels of endotoxin give a low-level
  inflammatory response: fever, malaise, aches and pains. A large flood releases the mediators and contributes to
  <strong>septic shock</strong>: tumor necrosis factor and interleukin 1 make neutrophils stick to endothelium and
  activate clotting, forming microthrombi throughout the body; arteries and arterioles dilate, peripheral arterial
  resistance falls and cardiac output at first <em>rises</em> (<strong>warm shock</strong>); later cardiac output
  falls and blood pressure drops, leading to tissue hypoxia, infarction and multisystem organ failure.</div>

  <h3 class="sub" id="gnb-sorting">14.2 &middot; Objective 1 &mdash; Sorting the Gram-negative bacilli</h3>
  <p>The Gram-negative bacilli are differentiated by three properties: <strong>oxygen requirement</strong>
  (aerobic, facultatively anaerobic, obligate anaerobic), the <strong>ability to ferment lactose</strong>, and the
  <strong>oxidase enzyme</strong>.</p>
  """ + fig("l14-s012-gram-negative-scheme.png", 449, 537,
      "Table 20.1, Gram-negative pathogens: Gram-negative rods branch into aerobes that do not ferment carbohydrates (Pseudomonas, Burkholderia, Bordetella, Francisella, Brucella, Alcaligenes), facultative anaerobes that ferment carbohydrates, and obligate anaerobes (Bacteroides). The facultative anaerobes split into lactose-fermenting (Escherichia, Klebsiella, Citrobacter, Enterobacter) and non-lactose-fermenting; the non-lactose fermenters split into oxidase-negative (Salmonella, Shigella, Proteus, Providencia, Morganella, Hafnia, Edwardsiella, Serratia, Yersinia) and oxidase-positive (Haemophilus, Pasteurella).",
      "<b>The scheme for the whole lecture (Table 20.1).</b> Oxygen first, then lactose, then oxidase. Note "
      "where <em>Serratia</em> and <em>Hafnia</em> sit here (non-lactose-fermenting) against slide 50, which "
      "calls them coliforms (14.21).", 12) + """
  <p><strong>Aerobic Gram-negative nonenteric bacilli</strong> are a large, diverse group of
  <strong>non-spore-forming</strong> bacteria, from habitats as varied as the large intestine, animals, the
  respiratory tract, soil and water. Most are not medically important; some are true pathogens and some
  opportunists. <strong>All use endotoxin as one of their offensive weapons.</strong> The medically important
  genera fall in three pairs:</p>
  <table>
    <tr><th>Pair</th><th>Kind of pathogen</th></tr>
    <tr><td><em>Pseudomonas</em> and <em>Burkholderia</em></td><td><strong>Opportunistic</strong> pathogens</td></tr>
    <tr><td><em>Brucella</em> and <em>Francisella</em></td><td><strong>Zoonotic</strong> pathogens</td></tr>
    <tr><td><em>Bordetella</em> and <em>Legionella</em></td><td>Mainly <strong>human</strong> pathogens</td></tr>
  </table>

  <h3 class="sub" id="gnb-pseudomonas">14.3 &middot; Objective 1 &mdash; <em>Pseudomonas aeruginosa</em></h3>
  <table>
    <tr><th>Heading</th><th>The pseudomonads, and <em>Pseudomonas aeruginosa</em></th></tr>
    <tr><td>Morphology</td><td>Small Gram-negative rods with a <strong>single polar flagellum</strong></td></tr>
    <tr><td>Habitat</td><td>Free living, primarily in soil, sea water and fresh water; also colonize plants and animals. Important decomposers and bioremediators (exoenzymes such as protease, amylase, pectinase and cellulase; oil-eating and pesticide-clearing strains); frequent contaminants in homes and clinical settings; food spoilage (dairy, meat, eggs). <em>P. aeruginosa</em>: soil and water, an intestinal resident in a minority of normal people, also on skin and in saliva</td></tr>
    <tr><td>Biochemistry</td><td><strong>Aerobic respiration; do not ferment carbohydrates.</strong> Produce <strong>oxidase and catalase</strong>. Produce pseudomycins (antifungal agents). Many make <strong>water-soluble pigments</strong> (blue-green, yellow-green, brown, red)</td></tr>
    <tr><td>Identifying signs</td><td><strong>Grapelike odor</strong>; <strong>greenish-blue pigment, pyocyanin</strong></td></tr>
    <tr><td>Hardiness</td><td>Resistant to soaps, dyes, quaternary ammonium disinfectants, drugs and drying; survives in nutrient-poor environments; a <strong>major agent of nosocomial infections</strong></td></tr>
    <tr><td>Why it contaminates equipment</td><td>Frequent contaminant of ventilators, intravenous solutions and anesthesia equipment because it <strong>produces biofilms</strong></td></tr>
  </table>
  """ + fig("l14-s015-polar-flagellum.jpg", 768, 387,
      "Electron micrograph of a single golden rod-shaped bacterium trailing one long flagellum from one end.",
      "<b>One polar flagellum</b>, the pseudomonad pattern. Compare the peritrichous flagella of "
      "<em>Proteus</em> (14.13).", 15) + """
  """ + fig("l14-s016-pseudomonas-pigment.jpg", 454, 800,
      "Two agar plates with antibiotic disks: the upper plate is stained green around the growth, the lower one blue-gray, with clear zones around several disks.",
      "<b>The pigment colors the medium.</b> <em>Pseudomonas aeruginosa</em> growing on two plates; its "
      "water-soluble pigment stains the agar greenish-blue.", 16, MH) + """
  <h4 class="subsub">Virulence and resistance</h4>
  <ul>
    <li><strong>Exotoxin A inhibits protein synthesis.</strong></li>
    <li><strong>Slime layer</strong>: resistance to phagocytosis and biofilm formation. Cells in biofilms adhere by
    flagella and pili and form microcolonies; three polysaccharides stabilize the biofilm and impair bacterial
    clearance.</li>
    <li><strong>Exoenzymes</strong>: <strong>exoenzyme S induces apoptosis of T cells</strong>; others degrade
    phospholipid.</li>
    <li><strong>Multidrug resistant</strong>, through restricted outer membrane permeability, which limits drug entry, antibiotic-inactivating
    enzymes, efflux pumps, horizontal transfer of resistance genes and biofilm formation. The slide notes it is one
    of the top three critical pathogens the World Health Organization prioritized for research into new treatments
    (with <em>Acinetobacter baumannii</em> and <em>Escherichia coli</em>; see 14.21).</li>
  </ul>
  """ + fig("l14-s018-pseudomonas-virulence.png", 929, 399,
      "Diagram of a Pseudomonas aeruginosa rod with a flagellum, outer membrane proteins, lipopolysaccharide, type 3 secretion system and type IV pili; secreted compounds listed as exotoxin A, quorum sensing molecules, siderophores, rhamnolipids and alginate; arrows to biofilm formation.",
      "<b>The virulence kit</b>: flagellum and type IV pili for attachment, secreted exotoxin A and alginate, and "
      "biofilm formation.", 18) + """
  <h4 class="subsub">Disease</h4>
  <ul>
    <li><mark class="prof-highlight">A common cause of <strong>nosocomial</strong> infection in hosts with
    <strong>burns, neoplastic disease and cystic fibrosis</strong></mark> (the slide stars cystic fibrosis with three
    asterisks). Complications include pneumonia, urinary tract infection, abscesses, otitis and corneal disease;
    <strong>septicemia</strong> leads to endocarditis, meningitis and bronchopneumonia.</li>
    <li><strong>Typical infections in healthy people</strong>: skin rashes, urinary tract infection,
    <strong>otitis externa (swimmer&rsquo;s ear)</strong>, and <strong>corneal ulcers in contact lens
    wearers</strong>.</li>
  </ul>
  """ + fig("l14-s019-pseudomonas-infection-sites.jpg", 785, 344,
      "Body diagram of Pseudomonas aeruginosa infections: central nervous system, eye, ear, nose and throat, respiratory infections including ventilator-associated and hospital-acquired pneumonia and chronic infection in cystic fibrosis, endocarditis, bacteremia, gastroenteritis, urinary tract infections including catheter-associated ones, osteomyelitis, burn wound and surgical site infections, hot tub folliculitis and other skin and soft tissue infections.",
      "<b>Everywhere it can reach.</b> Starred on the figure: contact lens keratitis, central line-associated "
      "bacteremia, ventilator-associated and hospital-acquired pneumonia, burn wound and surgical site infections, "
      "and catheter-associated urinary tract infections. The ear panel carries otitis externa, swimmer&rsquo;s ear.",
      19) + """
  """ + fig("l14-s019-corneal-ulcer.jpg", 500, 393,
      "Close photograph of an eye held open by a retractor, with a large gray-white opacity over the cornea and red, inflamed conjunctiva.",
      "<b>A corneal ulcer</b>, the <em>Pseudomonas aeruginosa</em> infection of contact lens wearers.", 19,
      "Medscape.") + """
  <p><strong>Treatment and prevention.</strong> The antibiotics listed are cephalosporins, aminoglycosides,
  carbenicillin, polymyxin, quinolones and monobactams. A <strong>vaccine is in development, particularly for
  cystic fibrosis patients</strong>: most candidates rely on antibodies against bacterial protein; one uses the
  PopB protein to stimulate T helper 17 cells, which release interleukin 17 and enhance antibacterial mucosal
  defenses; one in current testing uses the porin membrane protein OprH and has worked against at least two
  strains. But <strong>in cystic fibrosis, where the infection is often chronic, vaccination may induce increased
  pulmonary inflammation and tissue damage</strong>: it may worsen lung inflammation. Mucosal delivery by liposomes and nanoparticles is being
  investigated.</p>

  <h3 class="sub" id="gnb-relatives">14.4 &middot; Objective 1 &mdash; <em>Burkholderia</em>, <em>Acinetobacter</em>, <em>Stenotrophomonas</em></h3>
  <p>The related Gram-negative aerobic rods are <strong>similar to the pseudomonads</strong>: wide habitats in
  soil, water and related environments; <strong>obligate aerobes that do not ferment sugars</strong>;
  opportunistic pathogens. (The slide also calls them motile and oxidase positive; see 14.21.)</p>
  <table>
    <tr><th>Organism</th><th>Where it lives and how it gets in</th><th>Disease</th><th>Notes</th></tr>
    <tr><td><em>Burkholderia cepacia</em></td><td>Biodegrades many substances; a plant pathogen</td><td>Opportunist of the respiratory tract, urinary tract and skin</td><td><strong>Rapidly mutates; multidrug resistant; a particular problem in cystic fibrosis</strong>, as a copathogen with <em>Pseudomonas</em> (a high fatality)</td></tr>
    <tr><td><em>Burkholderia pseudomallei</em></td><td><strong>Soil and water of tropical climates</strong>; acquired through a <strong>penetrating injury or inhalation</strong></td><td><strong>Melioidosis</strong>: wound infection with localized nodules, fever and muscle aches; bronchitis and pneumonia; septicemia; endotoxic shock; liver and brain abscesses</td><td>The body map adds parotitis, bone, joint, genitourinary and skeletal-muscle infection</td></tr>
    <tr><td><em>Acinetobacter baumannii</em></td><td>Nosocomial and community acquired; <strong>survives months on fomites</strong></td><td>Wounds, lungs, urinary tract, burns, blood; <strong>infections after traumatic (combat) wounds</strong>; the immunocompromised</td><td>Also listed as needing urgent research; <strong>multidrug resistant, so combination treatment</strong> (carbapenems, polymyxin B, ampicillin)</td></tr>
    <tr><td><em>Stenotrophomonas maltophilia</em></td><td>Fresh water, soil, fecal flora of humans and domestic animals; <strong>forms biofilms</strong>; contaminates disinfectants, dialysis and respiratory equipment, water dispensers and catheters</td><td>Isolated from respiratory, soft tissue, blood, cerebrospinal fluid and eye specimens</td><td>Multiple drug resistance; combination therapy (cephalosporins, tetracyclines, fluoroquinolones, penicillins)</td></tr>
  </table>
  """ + fig("l14-s022-burkholderia-cepacia.png", 600, 742,
      "Diagram of Burkholderia cepacia: it produces antimicrobial compounds, colonizes plant roots and fixes nitrogen, and degrades groundwater contaminants; it causes sour skin onion rot in plants and pneumonia and septicemia in cystic fibrosis and immunocompromised patients.",
      "<b>Useful in soil, dangerous in cystic fibrosis.</b> The same organism protects crops and degrades "
      "pollutants, rots onions, and causes pneumonia and septicemia in cystic fibrosis and the "
      "immunocompromised.", 22, "Nature Reviews Microbiology.") + """
  """ + fig("l14-s024-melioidosis.jpg", 405, 345,
      "Body diagram of melioidosis: entry by inhalation and cutaneous inoculation, with central nervous system involvement, acute suppurative parotitis, pneumonia, bone and joint infections, liver and spleen abscesses, genitourinary and skin infections and skeletal muscle.",
      "<b>Melioidosis</b>, <em>Burkholderia pseudomallei</em>: in through the lungs or a skin wound, then almost "
      "any organ.", 24, "Copyright &copy; 2006 Nature Publishing Group, Nature Reviews Microbiology.") + """

  <h3 class="sub" id="gnb-brucella">14.5 &middot; Objective 1 &mdash; <em>Brucella</em> and brucellosis</h3>
  <table>
    <tr><th>Heading</th><th><em>Brucella</em></th></tr>
    <tr><td>Morphology</td><td><strong>Tiny Gram-negative coccobacilli</strong></td></tr>
    <tr><td>Species</td><td><strong><em>Brucella abortus</em> (cattle)</strong> and <strong><em>Brucella suis</em> (pigs)</strong></td></tr>
    <tr><td>Disease names</td><td>Brucellosis: <strong>Malta fever, undulant fever, Bang disease</strong></td></tr>
    <tr><td>Transmission</td><td>A <strong>zoonosis</strong> of animal handling and slaughtering: contact with the <strong>blood, urine or placentas of infected animals</strong>, or <strong>unpasteurized dairy products</strong>. <strong>Human-to-human transmission is rare</strong></td></tr>
    <tr><td>Pathogenesis</td><td>Enters through damaged skin or mucous membranes; phagocytes carry it to the blood, with lesions in the liver, spleen, bone marrow and kidneys</td></tr>
    <tr><td>Manifestations</td><td>A severe febrile illness with a <strong>fluctuating fever over weeks to a year (undulant fever)</strong>; chills, sweating, headache, muscle pain, weight loss</td></tr>
    <tr><td>Laboratory test</td><td>Serological: an <strong>agglutination titer test</strong></td></tr>
    <tr><td>Treatment</td><td>A combination of tetracycline and rifampin or streptomycin</td></tr>
    <tr><td>Prevention</td><td>An <strong>animal vaccine</strong> is available. A potential bioweapon</td></tr>
  </table>

  <h3 class="sub" id="gnb-francisella">14.6 &middot; Objective 1 &mdash; <em>Francisella tularensis</em> and tularemia</h3>
  <table>
    <tr><th>Heading</th><th><em>Francisella tularensis</em></th></tr>
    <tr><td>Disease</td><td><strong>Tularemia, rabbit fever</strong>: a zoonosis of mammals endemic to the northern hemisphere, particularly rabbits</td></tr>
    <tr><td>Transmission</td><td>Contact with <strong>infected animals</strong>, <strong>water and dust</strong>, or <strong>bites by vectors</strong>. The infective dose is very small (10&ndash;50 organisms)</td></tr>
    <tr><td>Course</td><td>Incubation of days to 3 weeks; headache, backache, fever, chills, malaise and weakness; <strong>ulcerative skin lesions and swollen lymph nodes</strong>; sore throat, intestinal involvement and respiratory infection. The systemic and pulmonic forms carry the highest death rate</td></tr>
    <tr><td>Relapse</td><td><strong>Intracellular persistence can lead to relapse</strong></td></tr>
    <tr><td>Treatment and prevention</td><td>Gentamicin or tetracycline; an <strong>attenuated vaccine</strong>. A potential bioterrorism agent</td></tr>
  </table>
  """ + fig("l14-s029-tularemia-ulcer.jpg", 220, 168,
      "Photograph of a finger with a raw, ulcerated red lesion on its side.",
      "<b>The ulcerative skin lesion of tularemia</b>, here on a finger, the kind of entry site left by handling "
      "an infected animal.", 29, "Centers for Disease Control and Prevention.") + """

  <h3 class="sub" id="gnb-bordetella">14.7 &middot; Objective 1 &mdash; <em>Bordetella pertussis</em> and whooping cough</h3>
  <table>
    <tr><th>Heading</th><th><em>Bordetella pertussis</em></th></tr>
    <tr><td>Morphology</td><td><strong>Minute, encapsulated coccobacillus</strong></td></tr>
    <tr><td>Disease</td><td><strong>Pertussis, whooping cough</strong>: a communicable childhood acute respiratory syndrome, often with severe, life-threatening complications in babies; about half of cases are in children from birth to 4 years</td></tr>
    <tr><td>Reservoir and transmission</td><td>Reservoir: <strong>apparently healthy carriers</strong>. Transmission by <strong>direct contact or inhalation of aerosols</strong></td></tr>
    <tr><td>Virulence</td><td>Receptors that recognize and <strong>bind ciliated respiratory epithelial cells</strong>; toxins that <strong>destroy and dislodge ciliated cells</strong>, destroy phagocyte activity, and cause hemolysis and leukolysis. <strong>Loss of the ciliary mechanism</strong> leads to a buildup of mucus and blockage of the airways</td></tr>
    <tr><td>Stages</td><td><strong>Catarrhal</strong>: nasal drainage, congestion, sneezing, occasional cough. <strong>Paroxysmal</strong>: recurrent persistent cough followed by a deep inhalation, the <strong>whoop</strong>. Complications come from compromised respiration</td></tr>
    <tr><td>Treatment</td><td>Two weeks of erythromycin, azithromycin or trimethoprim-sulfamethoxazole (the slide uses the brand name Septra)</td></tr>
    <tr><td>Prevention</td><td><strong>DTaP (diphtheria, tetanus and acellular pertussis): an acellular vaccine containing toxoid and other antigens</strong>, a five-dose series between 2 months and 6 years; Boostrix for children and adolescents aged 10&ndash;18</td></tr>
  </table>
  """ + fig("l14-s032-pertussis-course.jpg", 599, 416,
      "Outline of a body: 1, entry at the nose and mouth; 2, disease, with encephalopathy (infrequent), rhinorrhea, paroxysmal cough, lymphocytosis and bronchopneumonia (infrequent); 3, spread; 4, exit.",
      "<b>The course of pertussis</b>: in and out through the airway; the disease itself is rhinorrhea, the "
      "paroxysmal cough and a lymphocytosis, with encephalopathy and bronchopneumonia as the infrequent "
      "complications.", 32) + """
  <div class="pearl"><strong>Why pertussis is rising again.</strong> Case charts from 1922 to 2026 show the fall
  after the whole-cell vaccine and a rise since the acellular vaccine. The deck gives three threads:
  <strong>infants under 2 months are too young to be vaccinated</strong> (they cannot get the vaccine), so incidence is highest under 1 year old (in
  Florida in 2023, over half of reported cases were not up to date on vaccination); the age pattern in older
  children and adolescents <strong>reflects waning immunity from childhood vaccination and the need for booster
  doses</strong>; and, &ldquo;another possible connection&rdquo;, <strong>pertactin</strong>, the protein that helps
  the bacterium adhere to airway linings and a vaccine component, is missing from <strong>pertactin-negative
  strains now developing</strong>.</div>

  <h3 class="sub" id="gnb-alcaligenes">14.8 &middot; Objective 1 &mdash; <em>Alcaligenes</em></h3>
  <p>An aerobe of the non-fermenting branch (14.2), not on the syllabus list. Lives primarily in soil and water and
  may become normal flora. <em>Alcaligenes faecalis</em> is the most common clinical species, isolated from feces,
  sputum and urine; it is occasionally associated with opportunistic pneumonia, septicemia and meningitis.</p>

  <h3 class="sub" id="gnb-legionella">14.9 &middot; Objective 1 &mdash; <em>Legionella</em> and legionellosis</h3>
  <table>
    <tr><th>Heading</th><th><em>Legionella pneumophila</em></th></tr>
    <tr><td>Morphology and stain</td><td><strong>Weakly Gram negative</strong>; motile; rod shaped but pleomorphic, coccobacilli to filamentous</td></tr>
    <tr><td>Culture</td><td><strong>Cannot be easily cultured</strong>; grown on <strong>charcoal yeast extract agar</strong></td></tr>
    <tr><td>Habitat</td><td>Widely distributed in water: tap water, cooling towers, spas, ponds, air conditioning, supermarket vegetable sprayers. <strong>Lives in close association with amebas</strong></td></tr>
    <tr><td>Transmission</td><td><strong>By aerosols</strong>. Outbreaks have come from a cooling tower and from the spray of a decorative fountain</td></tr>
    <tr><td>Diseases</td><td><strong>Legionellosis</strong>: <strong>Legionnaires disease</strong> (named for the 1976 epidemic of pneumonia among American Legion members at a convention in Philadelphia) and <strong>Pontiac fever</strong></td></tr>
    <tr><td>Who gets it</td><td><strong>Males over 50; smokers</strong>; as a <strong>nosocomial</strong> disease, elderly patients with diabetes, cancer, transplants, chronic alcoholism or chronic lung disease</td></tr>
    <tr><td>Manifestations</td><td>Fever, cough, diarrhea, abdominal pain, pneumonia</td></tr>
    <tr><td>Treatment</td><td>Erythromycin with rifampin, azithromycin, quinolones</td></tr>
  </table>
  """ + fig("l14-s041-legionella.jpg", 304, 590,
      "Two panels: (a) a culture plate with gray-white Legionella growth on dark charcoal yeast extract agar; (b) an electron micrograph of ameba cells packed with Legionella bacteria.",
      "<b>Panel (a)</b>: <em>Legionella pneumophila</em> growing on charcoal yeast extract agar. <b>Panel (b)</b>: "
      "<em>Legionella</em> bacteria inside ameba cells, their home in water systems.", 41, MH) + """

  <h3 class="sub" id="gnb-enterics">14.10 &middot; Objective 1 &mdash; The Enterobacteriaceae: traits, media, antigens</h3>
  <p>The <strong>enterics</strong> are a large family of <strong>small, non-spore-forming Gram-negative
  rods</strong>. Many live in soil, water and decaying matter and are common occupants of the large bowel of
  animals, including humans. They are the <strong>most frequent cause of diarrheal disease</strong>, by
  <strong>enterotoxins</strong> or by <strong>invasive infection</strong>, and with <em>Pseudomonas</em> they
  account for almost half of nosocomial infections; they have been isolated from soap dishes, sinks, catheters
  and tracheal cannulas.</p>
  """ + fig("l14-s044-toxigenic-invasive.jpg", 860, 900,
      "Two drawings of intestinal epithelium: (a) bacteria lying on the surface with arrows showing oversecretion of fluids; (b) bacteria that have entered the lining, with blood pouring out.",
      "<b>Two kinds of diarrheal disease.</b> <b>(a) Toxigenic</b>: bacteria on the surface drive an "
      "oversecretion of fluids (example: enterotoxigenic <em>Escherichia coli</em>). <b>(b) Invasive</b>: "
      "bacteria enter the lining and it bleeds (example: <em>Shigella</em>).", 44, MH) + """
  <table>
    <tr><th>Every enteric</th><th>Splits them</th></tr>
    <tr><td><strong>Facultative anaerobes</strong>, growing best in air. <strong>All ferment glucose</strong>, reduce
    nitrates to nitrites, are <strong>oxidase negative</strong> and <strong>catalase positive</strong>. All carry the
    <strong>O antigen</strong></td>
    <td><strong>Lactose fermentation</strong>: <strong>coliforms</strong> ferment lactose, <strong>non-coliforms</strong> do not.
    Enrichment, selective and differential media screen samples for pathogens</td></tr>
  </table>
  """ + fig("l14-s047-enteric-biochemistry.jpg", 800, 549,
      "Flow chart: enrichment, then plating on eosin methylene blue, MacConkey or Hektoen enteric agar, then triple sugar iron. Lactose and glucose positive organisms are sorted by motility, indole, hydrogen sulfide, Voges-Proskauer, citrate and lysine decarboxylase into Escherichia, Klebsiella, Enterobacter, Citrobacter, Erwinia, Kluyvera and Moellerella; lactose negative, glucose positive organisms by phenylalanine, hydrogen sulfide, citrate, motility, a galactosidase test, gelatinase and urease into Proteus, Providencia, Morganella, Serratia, Salmonella, Hafnia, Edwardsiella, Yersinia and Shigella.",
      "<b>The biochemical tree.</b> Rapid lactose fermentation on triple sugar iron splits the enterics; then "
      "motility, indole, hydrogen sulfide, Voges-Proskauer, phenylalanine, citrate and urease. Worth reading off: "
      "<em>Klebsiella</em> is nonmotile and Voges-Proskauer positive; <em>Proteus</em> is phenylalanine and "
      "hydrogen sulfide positive; <em>Yersinia</em> and <em>Shigella</em> are both phenylalanine, citrate and "
      "motility negative, and urease splits them (<em>Yersinia</em> positive). Kits such as the Enterotube II run many of these tests at once (slide 49).", 47, MH) + """
  <table>
    <tr><th>Isolation medium</th><th>Selective for</th><th>Differential for</th></tr>
    <tr><td><strong>Eosin methylene blue (EMB) agar</strong></td><td>Gram-negative bacteria; the selective agents are the dyes <strong>eosin and methylene blue</strong></td><td>Lactose fermenters</td></tr>
    <tr><td><strong>MacConkey&rsquo;s agar</strong></td><td>Gram-negative bacteria; the selective agents are <strong>crystal violet and bile salts</strong></td><td>Lactose fermenters</td></tr>
  </table>
  """ + fig("l14-s048-emb-agar.jpg", 648, 643,
      "A dark red agar plate streaked with colonies that shine a metallic yellow-green.",
      "<b>Eosin methylene blue agar.</b>", 48) + """
  """ + fig("l14-s048-macconkey-agar.jpg", 528, 291,
      "A pink-orange agar plate with three streaks: pale growth on the left labeled lactose negative, and two deep pink streaks labeled lactose positive, the right-hand one E. coli.",
      "<b>MacConkey&rsquo;s agar.</b> Lactose fermenters such as <em>Escherichia coli</em> come up pink; "
      "non-fermenters stay pale.", 48) + """
  <table>
    <tr><th>In the normal flora: coliforms</th><th>In the normal flora: non-coliforms</th><th>True pathogens</th></tr>
    <tr><td><em>Escherichia coli</em>, <em>Enterobacter</em>, <em>Klebsiella</em>, <em>Hafnia</em>, <em>Citrobacter</em>, <em>Serratia</em> (slide 50; see 14.21)</td>
    <td><em>Proteus</em>, <em>Providencia</em>, <em>Morganella</em>, <em>Edwardsiella</em></td>
    <td>Enteric: <em>Salmonella typhi</em>, <em>S. cholerae-suis</em>, <em>S. enteritidis</em>; <em>Shigella dysenteriae</em>, <em>S. flexneri</em>, <em>S. boydii</em>, <em>S. sonnei</em>; <em>Yersinia enterocolitica</em>, <em>Y. pseudotuberculosis</em>. Nonenteric: <em>Yersinia pestis</em></td></tr>
  </table>
  <h4 class="subsub">Antigens and virulence factors</h4>
  <p>Complex surface antigens contribute to pathogenicity and trigger the immune response:</p>
  <ul>
    <li><strong>H</strong> &mdash; the <strong>flagellar</strong> antigen.</li>
    <li><strong>K</strong> &mdash; the <strong>capsule and/or fimbrial</strong> antigen (<strong>Vi</strong> in <em>Salmonella</em>).</li>
    <li><strong>O</strong> &mdash; the <strong>somatic or cell wall</strong> antigen, which <strong>all have</strong>.</li>
  </ul>
  <p>Other virulence factors: endotoxin, exotoxins, capsules, fimbriae, attachment receptors for host cells, and
  the ability to transmit virulence genes (toxins, capsules, hemolysins, fimbriae).</p>
  """ + fig("l14-s053-enteric-antigens.png", 542, 483,
      "Drawing of a rod with peritrichous flagella and a capsule: the capsule is labeled K antigen, or Vi in Salmonella; the cell wall somatic O antigen; the flagella H antigen.",
      "<b>Where the H, K and O antigens are.</b> The serotype names in 14.11 (O157:H7) are read off these.", 53) + """

  <h3 class="sub" id="gnb-ecoli">14.11 &middot; Objective 1 &mdash; <em>Escherichia coli</em> and its four pathotypes</h3>
  <p><em>Escherichia coli</em> is the common <strong>aerobic, non-fastidious</strong> bacterium of the gut, but
  not the most prevalent: the <strong>anaerobic <em>Bacteroides</em> and <em>Bifidobacterium</em></strong> are. Of
  about 150 strains, <strong>some have developed virulence through plasmid transfer</strong>; others are
  opportunists.</p>
  <p>The diarrheal strains are acquired from <strong>fecally contaminated food and water</strong> (the entry-and-exit
  diagram of slide 56). Each of the four pathotypes, row by row:</p>
  <table>
    <tr><th>Pathotype</th><th>How it causes disease</th><th>Disease</th><th>Source and diagnosis</th><th>Treatment and prevention</th></tr>
    <tr><td><strong>Enterotoxigenic <em>Escherichia coli</em> (ETEC)</strong></td>
      <td>Enterotoxigenic <em>Escherichia coli</em>: <strong>heat-labile and heat-stable toxins</strong> promote secretion of fluid and electrolytes by the intestinal epithelium; <strong>fimbriae</strong> adhere it to the <strong>small intestine</strong></td>
      <td><strong>Traveler&rsquo;s diarrhea</strong> (most cases): diarrhea, cramping, low-grade to no fever</td>
      <td>Fecally contaminated food and water</td>
      <td><strong>Fluid and electrolyte therapy</strong>; kaolin or Lomotil slow gut motility; <strong>bismuth salicylate counteracts the enterotoxin</strong> and is antimicrobial; fluoroquinolones</td></tr>
    <tr><td><strong>Enteroinvasive <em>Escherichia coli</em> (EIEC)</strong></td>
      <td>Enteroinvasive <em>Escherichia coli</em>: adhesin proteins bind the intestinal wall; <strong>no toxins</strong>, <strong>mechanical cell damage</strong>; an inflammatory disease of the <strong>large intestine</strong>, <strong>similar to shigellosis</strong></td>
      <td><strong>Profuse bloody diarrhea and high fever</strong>; the very young and very old most susceptible; incubation within 24 hours</td>
      <td>Fecally contaminated food and water; diagnosed by immunochemical tests, tissue culture and gene probes</td>
      <td>Fluid and electrolyte therapy; trimethoprim-sulfamethoxazole</td></tr>
    <tr><td><strong>Enteropathogenic <em>Escherichia coli</em> (EPEC)</strong></td>
      <td>Enteropathogenic <em>Escherichia coli</em>: <strong>adheres to the intestinal lining and destroys the microvilli</strong>; inhibits uptake of vitamin B1 (thiamin)</td>
      <td>The leading cause of the <strong>wasting form of infantile diarrhea</strong>; the greatest single cause of infant mortality worldwide</td>
      <td>Fecally contaminated food and water</td>
      <td>Fluid and electrolyte therapy; mecillinam, trimethoprim-sulfamethoxazole. <strong>Prevention: breast feeding</strong></td></tr>
    <tr><td><strong>Enterohemorrhagic <em>Escherichia coli</em> (EHEC), O157:H7</strong></td>
      <td>Enterohemorrhagic <em>Escherichia coli</em>, also called Shiga toxin-producing <em>Escherichia coli</em>: a cell wall receptor fuses with the host cell membrane, a portal for toxins including <strong>Shiga toxin</strong> (from <em>Shigella</em>), which <strong>binds ribosomes and disrupts protein synthesis</strong>; endothelial cells detach, platelets meet the subendothelium and coagulation starts; stress pathways release high tumor necrosis factor and interleukin 8</td>
      <td>Diarrhea, <strong>hemorrhagic colitis</strong> and <strong>hemolytic uremic syndrome</strong> (in a minority of cases: toxins cause apoptosis of renal tubular epithelial cells, and coagulation); gastroenteritis, dysentery, fever, acute abdominal pain</td>
      <td>Reservoir: <strong>the intestines of cattle</strong>; contaminated <strong>beef, water and vegetables</strong>; a small infective dose (about 100 cells). Diagnosed on <strong>Rainbow agar</strong>, where O157:H7 colonies are <strong>black</strong></td>
      <td>Dialysis; corticosteroids &ldquo;if diarrhea is not involved&rdquo; (the slide&rsquo;s wording); transfusions and plasmapheresis. <strong>Antibiotic therapy may exacerbate hemolytic uremic syndrome</strong> (it may worsen it)</td></tr>
  </table>
  <p><strong>Enterohemorrhagic outbreaks</strong> have come from a fast-food beef patty and slivered onions (2024),
  romaine lettuce, raw milk and cheese, and chicken wraps (2026). <strong>Rainbow agar</strong> carries chromogenic
  substrates for beta-galactosidase (blue/black) and beta-glucuronidase (red); <strong>O157:H7 is glucuronidase
  negative</strong>, so its colonies come up black.</p>
  <p>The name <strong>O157:H7</strong> reads the antigens of 14.10: <strong>O</strong> (somatic) antigen type 157 and
  <strong>H</strong> (flagellar) antigen type 7.</p>
  """ + fig("l14-s064-rainbow-agar.jpg", 458, 435,
      "Close view of an agar plate scattered with small black colonies among pink-purple ones.",
      "<b>Rainbow agar.</b> The black colonies are <em>Escherichia coli</em> O157:H7: galactosidase gives the "
      "blue-black, and O157:H7 lacks the glucuronidase that would add red.", 64) + """
  <p><strong><em>Escherichia coli</em> outside the gut.</strong> It causes most <strong>urinary tract
  infections</strong>, more frequently in <strong>women</strong> and in <strong>catheterized patients</strong>;
  also neonatal meningitis, pneumonia, septicemia, wound infections and complications of invasive procedures. The
  <strong>coliform count is the indicator of fecal contamination in water</strong>.</p>

  <h3 class="sub" id="gnb-coliforms">14.12 &middot; Objective 1 &mdash; The other coliforms</h3>
  <p>Clinically important mainly as <strong>opportunists</strong> and for <strong>antibiotic resistance</strong>.</p>
  <table>
    <tr><th>Organism</th><th>Distinguishing feature</th><th>Infections</th></tr>
    <tr><td><em>Klebsiella pneumoniae</em></td><td><strong>A normal inhabitant of the respiratory tract; a large capsule</strong></td><td><strong>Nosocomial pneumonia</strong>, meningitis, bacteremia, wound infections, urinary tract infections</td></tr>
    <tr><td><em>Enterobacter</em> species</td><td>Serious if in the bloodstream</td><td>Urinary tract and surgical wound infections; spinal fluid, sputum, blood</td></tr>
    <tr><td><em>Citrobacter</em> species</td><td>&mdash;</td><td>Opportunistic urinary tract infections and bacteremia</td></tr>
    <tr><td><em>Serratia marcescens</em></td><td><strong>Red pigment, prodigiosin</strong></td><td><strong>Pneumonia in chronic alcoholics</strong>; burn and wound infections; urinary tract infections (often asymptomatic, after surgery or instrumentation); septicemia and meningitis, particularly in the immunocompromised</td></tr>
  </table>
  """ + fig("l14-s067-klebsiella-capsule.jpg", 641, 409,
      "Capsule stain on a dark magenta background: pale rods, each ringed by a wide clear capsule, labeled cell body and capsule.",
      "<b>The <em>Klebsiella pneumoniae</em> capsule</b>: each cell body sits in a wide clear halo.", 67) + """
  """ + fig("l14-s068-serratia-colonies.png", 205, 136,
      "Close view of round colonies with fringed edges, deep red in the center, on a pale blue background.",
      "<b><em>Serratia marcescens</em></b>: the red of prodigiosin.", 68) + """

  <h3 class="sub" id="gnb-proteus">14.13 &middot; Objective 1 &mdash; <em>Proteus</em> and its relatives</h3>
  <p>The <strong>noncoliform, lactose-negative enterics</strong> fall into <strong>opportunists</strong>
  (<em>Proteus</em>, <em>Morganella</em>, <em>Providencia</em>) and <strong>true pathogens with true virulence
  factors</strong>, not normal flora (<em>Salmonella</em>, <em>Shigella</em>). The opportunists are ordinarily
  harmless saprobes in soil, manure, sewage and polluted water and commensals of humans and animals, often
  resistant to antibiotics.</p>
  <table>
    <tr><th>Organism</th><th>Features</th><th>Infections</th></tr>
    <tr><td><em>Proteus</em></td><td><strong>Highly motile</strong>, with <strong>peritrichous flagella</strong>; diagnostic <strong>wavelike, swarming growth</strong> (<em>Proteus vulgaris</em>); most respond to ampicillin and cephalosporins</td><td>Urinary tract and wound infections, pneumonia, septicemia, infant diarrhea; nosocomial infections; related to <strong>bladder and kidney stones</strong>; associated with <strong>Crohn&rsquo;s disease</strong></td></tr>
    <tr><td><em>Providencia</em> species</td><td><strong>Multiple animal reservoirs</strong>; almost always resistant to tetracyclines, older penicillins and cephalosporins</td><td>Nosocomial infections; <strong>urinary tract infections with catheters</strong>; gastroenteritis</td></tr>
    <tr><td><em>Morganella morganii</em></td><td>Multidrug resistant, though most strains are naturally susceptible to several drug classes</td><td>Nosocomial; similar to <em>Proteus</em>: urinary tract and wound infections</td></tr>
  </table>
  """ + fig("l14-s071-proteus-swarming.jpg", 325, 568,
      "An agar surface covered by concentric rings of growth spreading outward from two central points.",
      "<b>Swarming</b>: <em>Proteus vulgaris</em> moves out across the plate in waves, the rings its diagnostic "
      "growth pattern.", 71) + """

  <h3 class="sub" id="gnb-salmonella">14.14 &middot; Objective 1 &mdash; <em>Salmonella</em>: typhoid and the salmonelloses</h3>
  <p><em>Salmonella typhi</em> causes <strong>typhoid fever</strong> and <em>Salmonella paratyphi</em>
  paratyphoid fever, the most serious diseases of the genus.</p>
  <table>
    <tr><th>Heading</th><th><em>Salmonella typhi</em> (<em>S. enterica</em> serotype typhi) and <em>S. paratyphi</em></th></tr>
    <tr><td>Significance</td><td>The <strong>most serious pathogens of the genus</strong>: typhoid fever and paratyphoid fever</td></tr>
    <tr><td>Host</td><td><strong>Humans are the exclusive hosts</strong></td></tr>
    <tr><td>Traits</td><td>Flagellated; survive outside the host; <strong>resistant to chemicals: bile and dyes</strong>; do not lose virulence on cultivation</td></tr>
    <tr><td>Transmission</td><td><strong>Ingestion of fecally contaminated food or water</strong>; occasionally close personal contact. &ldquo;<strong>Boil it, cook it, peel it, or forget it.</strong>&rdquo;</td></tr>
    <tr><td>Carriers</td><td>Asymptomatic carriers; some <strong>chronic carriers shed bacilli from the gallbladder</strong></td></tr>
    <tr><td>Pathogenesis</td><td>Bacilli <strong>adhere to the small intestine</strong> and cause an invasive diarrhea that <strong>leads to septicemia</strong>; carried to the liver and spleen; the intestinal wall can ulcerate and perforate</td></tr>
    <tr><td>Manifestations</td><td><strong>Sustained fever</strong> of 103&ndash;104 degrees; a rash of <strong>rose-colored spots</strong>; weakness, stomach pains, loss of appetite</td></tr>
    <tr><td>Treatment</td><td>Chronic infections: chloramphenicol or trimethoprim-sulfamethoxazole</td></tr>
    <tr><td>Prevention</td><td><strong>Two vaccines for temporary protection: oral (4 doses) and injected (1 dose)</strong></td></tr>
  </table>
  """ + fig("l14-s074-typhoid-route.jpg", 511, 687,
      "Drawing: typhoid bacilli enter an intestinal villus, pass into its blood vessels and are carried to the liver and spleen; below, a section of intestinal wall with ulcers and perforations.",
      "<b>The typhoid route</b>: into a villus of the small intestine, into the blood, on to the liver and "
      "spleen; the wall can ulcerate and perforate. Compare <em>Shigella</em>, which stays in the large-bowel "
      "lining (14.15).", 74) + """
  <p><strong>The animal salmonelloses.</strong> Salmonelloses other than typhoid fever are called enteric fevers,
  <em>Salmonella</em> food poisoning and gastroenteritis. They are <strong>usually less severe than typhoid fever
  but more prevalent</strong> (some 40,000&ndash;50,000 cases a year). <em>Salmonella</em> is normal intestinal flora
  in animals: <strong><em>S. cholerae-suis</em> is a zoonosis of swine</strong>; <em>S. enteritidis</em> covers
  about 1,700 serotypes, all zoonotic in origin, though humans can become carriers. They are <strong>acquired by
  ingestion</strong>, from cattle, <strong>poultry and eggs</strong>, rodents, <strong>reptiles (pet turtles,
  bearded dragons)</strong>, animal and dairy products, backyard flocks, and fomites contaminated with animal
  intestinal flora; food-borne outbreaks have come from cucumbers and charcuterie meats. Enteric fevers bring fever
  and septicemia; gastroenteritis brings vomiting, diarrhea and fluid loss. Treatment: chloramphenicol and
  trimethoprim-sulfamethoxazole.</p>

  <h3 class="sub" id="gnb-shigella">14.15 &middot; Objective 1 &mdash; <em>Shigella</em> and bacillary dysentery</h3>
  <table>
    <tr><th>Heading</th><th><em>Shigella</em></th></tr>
    <tr><td>Species</td><td><em>S. dysenteriae</em>, <em>S. sonnei</em>, <em>S. flexneri</em>, <em>S. boydii</em></td></tr>
    <tr><td>Disease</td><td><strong>Shigellosis, an incapacitating dysentery</strong>: abdominal cramps and watery stool with blood; a few cases also develop hemolytic uremic syndrome. Half of cases are in children under 5</td></tr>
    <tr><td>Transmission</td><td><strong>Ingestion of fecally contaminated food</strong></td></tr>
    <tr><td>Pathogenesis</td><td><strong>Invades the villi of the large intestine</strong> (it invades the large-bowel villi); <strong>does not perforate the intestine or invade the blood</strong>; enters Peyer&rsquo;s patches and sets off inflammation</td></tr>
    <tr><td>Toxins</td><td><strong>Endotoxin initiates the fever</strong>; <strong>enterotoxin damages the mucosa and villi</strong>, causing bleeding and mucus secretion</td></tr>
    <tr><td>Treatment</td><td>Fluid replacement; ciprofloxacin and trimethoprim-sulfamethoxazole</td></tr>
  </table>
  """ + fig("l14-s079-shigella-colon.jpg", 831, 290,
      "Drawing of a length of large intestine opened to show inflamed red mucosa with scattered dark bleeding spots.",
      "<b>The large-intestinal mucosa in shigellosis</b>: inflamed and bleeding in patches, the invasive "
      "pattern of 14.10.", 79) + """

  <h3 class="sub" id="gnb-yersinia-enteric">14.16 &middot; Objective 1 &mdash; The enteric yersinias</h3>
  <ul>
    <li><strong><em>Yersinia enterocolitica</em></strong>: from domestic and wild animals, fish, fruits, vegetables
    and water. The bacteria enter the small intestinal mucosa; some enter the lymphatics and survive in phagocytes.
    <strong>Inflammation of the ileum can mimic appendicitis.</strong></li>
    <li><strong><em>Yersinia pseudotuberculosis</em></strong>: an infection similar to <em>Y. enterocolitica</em>,
    with <strong>more lymph node inflammation</strong>.</li>
    <li>Treatment is mostly supportive (aminoglycosides and trimethoprim-sulfamethoxazole).</li>
  </ul>

  <h3 class="sub" id="gnb-plague">14.17 &middot; Objective 1 &mdash; <em>Yersinia pestis</em> and plague</h3>
  <table>
    <tr><th>Heading</th><th><em>Yersinia pestis</em> (nonenteric)</th></tr>
    <tr><td>Morphology</td><td><strong>Tiny Gram-negative rod with unusual bipolar staining and capsules</strong></td></tr>
    <tr><td>Virulence</td><td>Capsular and envelope proteins protect against phagocytosis and foster intracellular growth; <strong>coagulase</strong>, endotoxin, murine toxin</td></tr>
    <tr><td>Reservoir</td><td>About 200 species of mammals, rodents, without causing them disease. In the United States, <strong>cases are west of the Mississippi</strong> (a handful a year)</td></tr>
    <tr><td>How humans get it</td><td>Contact with <strong>wild animals (sylvatic plague)</strong>, <strong>domestic or semidomestic animals (urban plague)</strong>, or infected humans (rare)</td></tr>
    <tr><td>The flea</td><td>The bacteria replicate in the flea&rsquo;s gut; <strong>coagulase causes blood clotting that blocks the esophagus</strong>; the flea becomes <strong>ravenous</strong> and tries to bite more hosts; <strong>regurgitated material transmits the bacterium</strong></td></tr>
    <tr><td>Infective dose</td><td>Very small: 3&ndash;50 bacilli</td></tr>
  </table>
  """ + fig("l14-s082-yersinia-pestis.jpg", 791, 484,
      "Stained smear with several large red white blood cells and many tiny pink rods, labeled Y. pestis; an inset circle magnifies four rods, darker at both ends.",
      "<b>Bipolar staining.</b> Tiny <em>Yersinia pestis</em> rods among white blood cells; the inset shows the "
      "stain gathered at both ends of each rod.", 82) + """
  """ + fig("l14-s084-plague-cycle.jpg", 800, 397,
      "Cycle diagram: mice and fleas as endemic reservoir hosts; chipmunks and ground squirrels, with fleas, as amplifying hosts; humans as accidental hosts with bubonic plague from flea bites and pneumonic plague passed person to person by droplets; human-to-human transmission is less common.",
      "<b>The plague cycle.</b> Fleas keep it going among endemic reservoir rodents and amplifying hosts "
      "(chipmunks, ground squirrels); humans are accidental hosts, bubonic from a flea bite, pneumonic by aerosol "
      "and droplets.", 84, MH) + """
  <table>
    <tr><th>Form of plague</th><th>What happens</th></tr>
    <tr><td><strong>Bubonic</strong></td><td>The bacillus multiplies in the flea bite, enters the lymph, and causes necrosis and swelling of a lymph node: the <strong>bubo</strong>, a swollen, necrotic lymph node in the groin or axilla (the classic sign)</td></tr>
    <tr><td><strong>Septicemic</strong></td><td>Progression to massive bacterial growth; virulence factors cause intravascular coagulation, subcutaneous hemorrhage and <strong>purpura that degenerates into necrosis and gangrene</strong> (the purpura turns to gangrene): <strong>the black plague</strong></td></tr>
    <tr><td><strong>Pneumonic</strong></td><td>Infection localized to the lungs, <strong>highly contagious by aerosol</strong>, fatal without treatment</td></tr>
  </table>
  <p><strong>Diagnosis, treatment, prevention.</strong> It <strong>must be diagnosed quickly</strong>: death can come
  within 2 days of symptoms. Diagnosis rests on <strong>history, symptoms and laboratory findings from aspiration of
  the buboes</strong>. Treatment: streptomycin, tetracycline or chloramphenicol. A <strong>killed or attenuated
  vaccine</strong> is available for military personnel, veterinarians and laboratory workers. Prevention by
  <strong>quarantine and control of the rodent population in human habitats</strong>. A park warning sign (slide 86)
  spells out the same advice: chipmunks, ground squirrels and other wild rodents may be infected; plague reaches
  humans by the bite of an infected flea or by handling an infected animal; protect pets from fleas, avoid
  animal burrows, and never touch sick or dead animals.</p>
  <div class="pearl"><strong>Plague in the world, and in our genes.</strong> Plague is endemic in Madagascar
  (seasonal, October to March, when rats move into homes in the wet season), with urban outbreaks of the pneumonic
  form in 2017 and 2021; in Russia, a laboratory worker at a Siberian anti-plague institute was hospitalized with a
  respiratory illness and died. The
  Black Death of 1346&ndash;1350 killed a large share of a population that was immunologically naive, putting
  strong selective pressure on genetic variants that conferred immunity, so later outbreaks killed far fewer.
  Deoxyribonucleic acid from that era shows that <strong>those plague-immunity variants overlap with alleles now
  associated with increased susceptibility to autoimmune diseases</strong>.</div>

  <h3 class="sub" id="gnb-pasteurella">14.18 &middot; Objective 1 &mdash; <em>Pasteurella multocida</em></h3>
  <p>With <em>Haemophilus</em>, one of the <strong>oxidase-positive nonenteric</strong> pathogens. A
  <strong>zoonotic genus, normal flora in animals</strong>. Opportunistic infections, mostly in
  <strong>veterinarians</strong>, but they can come from domestic animals: <strong>animal bites or scratches cause a
  local abscess</strong> that can spread to joints, bones and lymph nodes. <strong>The immunocompromised are at risk
  of septicemia</strong> and complications of the central nervous system and heart. Treatment: penicillin and
  tetracycline.</p>

  <h3 class="sub" id="gnb-haemophilus">14.19 &middot; Objective 1 &mdash; <em>Haemophilus</em></h3>
  <p><em>Haemophilus</em> means <strong>blood loving</strong>: tiny Gram-negative pleomorphic rods,
  <strong>fastidious</strong>, sensitive to drying, temperature extremes and disinfectants. <strong>None can grow on
  blood agar without special techniques: chocolate agar (cooked blood)</strong> is used, because they
  <strong>require hemin, nicotinamide adenine dinucleotide (NAD) or its phosphate (NADP)</strong>. Some species are
  normal colonists of the upper respiratory tract or vagina; others are virulent.</p>
  <table>
    <tr><th>Species</th><th>Disease</th><th>Spread and who</th><th>Treatment and prevention</th></tr>
    <tr><td><strong><em>H. influenzae</em></strong></td><td><strong>Misnamed</strong> as the cause of the flu. <strong>Acute bacterial meningitis</strong> (fever, vomiting, stiff neck, neurological impairment), <strong>epiglottitis</strong>, otitis media, sinusitis, pneumonia, bronchitis</td><td>&mdash;</td><td>Chloramphenicol and ampicillin; <strong>vaccine: the subunit vaccine Hib</strong> (<em>Haemophilus influenzae</em> type b)</td></tr>
    <tr><td><em>H. aegyptius</em></td><td><strong>Conjunctivitis, pink eye</strong></td><td>Primarily children; spread by <strong>contaminated fingers and insects</strong></td><td>Antibiotic eyedrops</td></tr>
    <tr><td><em>H. ducreyi</em></td><td><strong>Chancroid</strong>, a sexually transmitted disease</td><td>Primarily in the tropics; primarily males</td><td>Erythromycin</td></tr>
    <tr><td><em>H. parainfluenzae</em> and <em>H. aphrophilus</em></td><td><strong>Infective endocarditis</strong> in congenital and rheumatic heart disease</td><td><strong>Normal oral and nasopharyngeal flora</strong>; from <strong>routine dental procedures</strong></td><td>Ciprofloxacin</td></tr>
  </table>
  <p>Invasive <em>H. influenzae</em> disease in the United States (1994&ndash;2020) is now mostly nontypeable
  strains; type b, the vaccine&rsquo;s target, has become the rarest (slide 99).</p>
  """ + fig("l14-s101-conjunctivitis.jpg", 721, 600,
      "Close photograph of an eye with intensely red, inflamed conjunctiva around the iris.",
      "<b>Acute conjunctivitis (pink eye)</b>, the <em>Haemophilus aegyptius</em> infection of children.", 101,
      MH) + """

  <h3 class="sub" id="gnb-transmission">14.20 &middot; Objective 2 &mdash; Modes of transmission</h3>
  <table>
    <tr><th>Organism</th><th>How humans acquire it</th></tr>
    <tr><td><em>Pseudomonas aeruginosa</em></td><td>Soil and water; <strong>nosocomial</strong>, from contaminated ventilators, intravenous solutions and anesthesia equipment (biofilms); contact lenses (corneal ulcers), water (swimmer&rsquo;s ear)</td></tr>
    <tr><td><em>Burkholderia pseudomallei</em></td><td>Tropical soil and water, by <strong>penetrating injury or inhalation</strong></td></tr>
    <tr><td><em>Acinetobacter baumannii</em></td><td>Nosocomial and community; <strong>months on fomites</strong>; traumatic and combat wounds</td></tr>
    <tr><td><em>Stenotrophomonas maltophilia</em></td><td>Biofilms on disinfectants, dialysis and respiratory equipment, water dispensers and catheters</td></tr>
    <tr><td><em>Brucella</em></td><td>Zoonosis: <strong>animal handling and slaughtering</strong> (blood, urine, placentas) or <strong>unpasteurized dairy</strong>; through damaged skin or mucous membranes; human-to-human rare</td></tr>
    <tr><td><em>Francisella tularensis</em></td><td><strong>Infected animals (rabbits), water and dust, or vector bites</strong></td></tr>
    <tr><td><em>Bordetella pertussis</em></td><td>From apparently healthy carriers by <strong>direct contact or inhaled aerosols</strong></td></tr>
    <tr><td><em>Legionella</em></td><td><strong>Aerosols</strong> from water systems (cooling towers, spas, fountains, air conditioning, vegetable sprayers)</td></tr>
    <tr><td>Enterics in hospitals</td><td>Soap dishes, sinks, catheters, tracheal cannulas</td></tr>
    <tr><td><em>Escherichia coli</em> diarrheal strains</td><td><strong>Fecally contaminated food and water</strong>; O157:H7 from <strong>cattle</strong> via beef, water and vegetables</td></tr>
    <tr><td><em>Escherichia coli</em> urinary tract infection</td><td>Women; catheterized patients</td></tr>
    <tr><td><em>Providencia</em></td><td>Multiple animal reservoirs; catheters</td></tr>
    <tr><td><em>Salmonella typhi</em></td><td><strong>Humans only</strong>: fecally contaminated food or water, occasionally close contact; chronic gallbladder carriers</td></tr>
    <tr><td>Animal salmonelloses</td><td><strong>Ingestion</strong> from cattle, poultry and eggs, rodents, reptiles, dairy, backyard flocks, contaminated fomites</td></tr>
    <tr><td><em>Shigella</em></td><td><strong>Ingestion of fecally contaminated food</strong></td></tr>
    <tr><td><em>Yersinia enterocolitica</em></td><td>Animals, fish, fruits, vegetables, water</td></tr>
    <tr><td><em>Yersinia pestis</em></td><td><strong>Flea bite</strong> (regurgitation by a blocked flea); handling infected animals; <strong>aerosol</strong> in pneumonic plague; sylvatic (wild animals) or urban (domestic animals)</td></tr>
    <tr><td><em>Pasteurella multocida</em></td><td><strong>Animal bites or scratches</strong></td></tr>
    <tr><td><em>Haemophilus</em></td><td><em>H. aegyptius</em>: fingers and insects. <em>H. ducreyi</em>: sexual. <em>H. parainfluenzae</em> and <em>H. aphrophilus</em>: own oral flora, during dental procedures</td></tr>
  </table>

  <h3 class="sub" id="gnb-notes">14.21 &middot; Where the slides disagree, and where current practice differs</h3>
  <p>The slides are what is taught and tested; where they disagree with each other, the point is not keyed in the
  quizzes. Notes:</p>
  <ul>
    <li><strong><em>Serratia</em> and <em>Hafnia</em>: coliform or not?</strong> Slide 50 lists both among the
    coliforms (lactose fermenters), and slide 68 calls <em>Serratia</em> a coliform; Table 20.1 (slide 12) and the
    biochemical tree (slide 47) put both on the non-lactose-fermenting branch. (<em>Serratia</em> is in fact a slow,
    late lactose fermenter.) No question asks their lactose status.</li>
    <li><strong>Which <em>Haemophilus</em> causes childhood meningitis?</strong> Slide 97 names <em>H.
    parainfluenzae</em> and <em>H. ducreyi</em> as the virulent species &ldquo;responsible for childhood meningitis,
    and chancroid&rdquo;; slide 98 gives meningitis to <em>H. influenzae</em>, and slide 100 calls <em>H.
    parainfluenzae</em> normal oral flora that causes endocarditis. Go by slides 98 and 100.</li>
    <li><strong>&ldquo;Motile, oxidase positive&rdquo;</strong> is given to <em>Burkholderia</em>,
    <em>Acinetobacter</em> and <em>Stenotrophomonas</em> together (slide 21). In fact <em>Acinetobacter</em> is
    nonmotile and oxidase negative, and Lecture 9 taught <em>Acinetobacter baumannii</em> among the Gram-negative
    cocci (it is a coccobacillus). Not keyed; their shared traits are keyed as obligate aerobes that do not ferment
    sugars.</li>
    <li><strong>The World Health Organization list (slide 17).</strong> The 2017 list put carbapenem-resistant
    <em>Pseudomonas aeruginosa</em> among the critical pathogens with <em>Acinetobacter baumannii</em> and the
    resistant Enterobacteriaceae; the 2024 update moved it to the high-priority group. Not keyed.</li>
    <li><strong>Boostrix (slide 31)</strong> is given for ages 10&ndash;18 on the slide; it is licensed from age 10
    through adulthood, including pregnancy. Not keyed.</li>
    <li><strong>Drug lists.</strong> Several treatment lines reflect older practice (chloramphenicol for typhoid and
    for <em>H. influenzae</em> meningitis, erythromycin with rifampin for <em>Legionella</em>, streptomycin for
    plague); current regimens differ. Doses on slide 56 are not reproduced. Which antibiotic treats what is not asked
    in the quizzes; the non-drug methods (fluids and electrolytes, dialysis and plasmapheresis, vaccines, breast
    feeding, quarantine, combination therapy for multidrug resistance) are.</li>
    <li><strong>Spellings.</strong> The deck&rsquo;s &ldquo;Acinetobacter baumanii&rdquo; is <em>Acinetobacter
    baumannii</em>; &ldquo;Salmonella typhhi&rdquo; <em>Salmonella typhi</em>; &ldquo;Mecillianam&rdquo; mecillinam;
    &ldquo;sulfatrimethoprim&rdquo; and &ldquo;Sulfamethoprim&rdquo; are trimethoprim-sulfamethoxazole.</li>
  </ul>

  <button type="button" class="test-yourself-btn" onclick="window.openTestYourself('Test yourself &mdash; Gram-Negative Bacilli', TEST_YOURSELF.gramNegativeBacilli)">Test yourself! &rarr;</button>
  <p class="src">Source: <em>Gram Negative Bacilli of Medical Importance</em> (Lecture 14  Dr. Webster Gram Negative
  Bacilli of Medical Importance.pptx), Slides 1&ndash;101, and the PAJ 5200 syllabus instructional objectives.</p>
</section>
"""

SECTION = SECTION_HEAD + BODY

TEST = """
    gramNegativeBacilli: [
      {q:"Which part of the lipopolysaccharide is biologically active?",
       choices:["The O antigen","The core polysaccharide","Lipid A","The peptidoglycan"],correct:2,
       explain:"Lipid A is the biologically active component of endotoxin. The O antigen is unique to each organism; the core polysaccharide is shared by large groups."},
      {q:"Which three properties sort the Gram-negative bacilli?",
       choices:["Spores, acid-fastness, capsule","Oxygen, lactose fermentation, oxidase","Catalase, coagulase, hemolysis","Shape, size, pigment"],correct:1,
       explain:"Oxygen requirement, lactose fermentation and the oxidase enzyme. Coliforms are the lactose-fermenting enterics; Haemophilus and Pasteurella are oxidase positive."},
      {q:"Which Escherichia coli strain is the leading cause of traveler's diarrhea?",
       choices:["Enterotoxigenic","Enteroinvasive","Enteropathogenic","Enterohemorrhagic"],correct:0,
       explain:"Enterotoxigenic Escherichia coli, with its heat-labile and heat-stable toxins. Enteropathogenic causes wasting infantile diarrhea; enterohemorrhagic O157:H7 causes hemolytic uremic syndrome."},
      {q:"How does an infected flea transmit plague?",
       choices:["Through its feces","By regurgitating while biting","Through its eggs","Only by being swallowed"],correct:1,
       explain:"Coagulase clots blood that blocks the flea's esophagus; the ravenous flea bites again and again and regurgitates the bacteria into the bite."},
      {q:"Why is Haemophilus grown on chocolate agar?",
       choices:["It needs low oxygen","It needs hemin and similar blood factors","It ferments lactose there","Bile salts kill it"],correct:1,
       explain:"Haemophilus is fastidious and needs hemin or nicotinamide adenine dinucleotide, which chocolate agar (cooked blood) supplies; none grows on blood agar without special techniques."}
    ],
"""
