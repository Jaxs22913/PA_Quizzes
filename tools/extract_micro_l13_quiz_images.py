#!/usr/bin/env python3
"""Write the Lecture 13 (Misc. Bacterial Agents of Disease) QUIZ pictures into
"Microbiology Exam 2/micro-exam-2-quiz-images/".

QUIZ images carry NO label that names the answer, so each is a crop of the slide's picture:
  slide 13  the RIGHT photograph only (the palms); the trunk photograph, the banner and the credit cropped off
  slide 14  the LEFT picture by geometry (the gumma; the right one is the pupil); banner and credit cropped
  slide 15  panel (b) only (the teeth); the infant's face, "(a)/(b)", banner and credit cropped off
  slide 30  the photograph; banner and credit cropped
  slide 50  the photograph; banner and credit cropped
Each picture is chosen by its position on the slide (python-pptx left/top), never by an extraction index, and the
crop boxes were read off the original pixels (content bands between the white strips). Each result was viewed.

    python3 tools/extract_micro_l13_quiz_images.py
"""
import io, os
from pptx import Presentation
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DECK = os.path.expanduser("~/Desktop/PA Quizzes/Semester 2/Microbiology Inbox/Exam 2/PAJ5200.Misc inf agents-2.pptx")
OUT = os.path.join(ROOT, "Microbiology Exam 2", "micro-exam-2-quiz-images")

# (slide, which picture by position: "only" | "left" | "right", crop box on the original pixels, output name)
JOBS = [
    (13, "only", (610, 33, 1193, 495), "l13-s13-palm-rash.jpg"),
    (14, "left", (2, 23, 797, 545), "l13-s14-gumma.jpg"),
    (15, "only", (17, 487, 783, 857), "l13-s15-hutchinson-teeth.jpg"),
    (30, "only", (0, 16, 800, 566), "l13-s30-erythema-migrans.jpg"),
    (50, "only", (3, 23, 797, 533), "l13-s50-spotted-rash.jpg"),
]
MAXW = 600


def pick(slide, which):
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if which == "only":
        assert len(pics) == 1, "expected one picture, found %d" % len(pics)
        return pics[0]
    pics.sort(key=lambda sh: sh.left)
    assert len(pics) == 2, "expected two pictures side by side"
    return pics[0] if which == "left" else pics[1]


def main():
    prs = Presentation(DECK)
    os.makedirs(OUT, exist_ok=True)
    for n, which, box, name in JOBS:
        sh = pick(prs.slides[n - 1], which)
        im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB").crop(box)
        if im.width > MAXW:
            im = im.resize((MAXW, round(im.height * MAXW / im.width)), Image.LANCZOS)
        im.save(os.path.join(OUT, name), quality=85)
        print("wrote %-32s %dx%d  (slide %d, %s picture)" % (name, im.width, im.height, n, which))


if __name__ == "__main__":
    main()
